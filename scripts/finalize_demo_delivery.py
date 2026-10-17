"""Validate the recorded job and package a short operation tour and screenshots."""
from __future__ import annotations

import hashlib
import json
import os
import subprocess
import sys
import zipfile
from pathlib import Path

import cv2
import fitz
import numpy as np
from playwright.sync_api import sync_playwright
from docx import Document
from PIL import Image, ImageDraw, ImageFont, ImageOps

from demo_runtime import PROJECT_DIR, configuration, job_root, read_json, write_json

OUTPUT = PROJECT_DIR / configuration()["delivery_dir"]
FONT = Path("C:/Windows/Fonts/msyh.ttc")
GENERATION_SPEED = int(configuration()["video_idle_speed"])


def title_card(size: tuple[int, int], title: str, subtitle: str):
    image = Image.new("RGB", size, "#f2f6f1")
    draw = ImageDraw.Draw(image)
    draw.text((100, size[1] // 2 - 110), title, font=ImageFont.truetype(str(FONT), 48), fill="#23483b")
    draw.text((104, size[1] // 2 - 15), subtitle, font=ImageFont.truetype(str(FONT), 24), fill="#728474")
    draw.line((104, size[1] // 2 + 55, size[0] - 104, size[1] // 2 + 55), fill="#c6d8ca", width=2)
    return cv2.cvtColor(np.asarray(image), cv2.COLOR_RGB2BGR)


def build_short_video(manifest: dict):
    assert manifest["pace_marker"] == "black_corner_idle_white_corner_active"
    capture = cv2.VideoCapture(str(OUTPUT / "operation-full.webm"))
    fps = capture.get(cv2.CAP_PROP_FPS)
    size = (int(capture.get(cv2.CAP_PROP_FRAME_WIDTH)), int(capture.get(cv2.CAP_PROP_FRAME_HEIGHT)))
    target = OUTPUT / "operation-demo.mp4"
    writer = cv2.VideoWriter(str(target), cv2.VideoWriter_fourcc(*"mp4v"), fps, size)
    assert writer.isOpened()
    written = 0
    banner = Image.new("RGB", (410, 42), "#234b3b")
    ImageDraw.Draw(banner).text((16, 7), f"等待中 · {GENERATION_SPEED} 倍速", font=ImageFont.truetype(str(FONT), 19), fill="white")
    fast_banner = cv2.cvtColor(np.asarray(banner), cv2.COLOR_RGB2BGR)
    normal_banner = Image.new("RGB", (410, 42), "#365d4d")
    ImageDraw.Draw(normal_banner).text((16, 7), "操作与结果查看 · 原速", font=ImageFont.truetype(str(FONT), 19), fill="white")
    normal_banner = cv2.cvtColor(np.asarray(normal_banner), cv2.COLOR_RGB2BGR)
    frame_index, idle_frames, active_frames, active_kept = 0, 0, 0, 0
    accumulator = 0.0
    mapping, idle_flags = [], []
    while True:
        ok, frame = capture.read()
        if not ok:
            break
        is_idle = float(frame[2:8, 2:8].mean()) < 100
        idle_flags.append(is_idle)
        index = frame_index
        frame_index += 1
        if is_idle:
            idle_frames += 1
            accumulator += 1 / GENERATION_SPEED
            if accumulator < 1 - 1e-8:
                continue
            accumulator -= 1
            frame[0:42, size[0] - 410:size[0]] = fast_banner
        else:
            active_frames += 1
            active_kept += 1
            accumulator = 0.0
            frame[0:42, size[0] - 410:size[0]] = normal_banner
        frame[:10, :10] = frame[12, 12]
        writer.write(frame)
        mapping.append(index)
        written += 1
    capture.release()
    writer.release()
    assert idle_frames > 0 and active_frames == active_kept
    assert all(left < right for left, right in zip(mapping, mapping[1:]))
    kept = set(mapping)
    assert all(is_idle or index in kept for index, is_idle in enumerate(idle_flags))
    np.savez_compressed(OUTPUT / "video_frame_mapping.npz", source_frames=np.array(mapping), idle_flags=np.array(idle_flags))
    return {"path": str(target), "seconds": written / fps, "fps": fps, "generation_speed": GENERATION_SPEED,
            "full_source": manifest["full_webm"], "source_frames": frame_index, "output_frames": written,
            "idle_frames": idle_frames, "active_frames": active_frames, "all_active_frames_kept": True,
            "source_order_monotonic": True, "hard_cuts": 0, "timing_basis": "frame_marker"}


def capture_node_screenshots(job_id: str):
    """Use a separate, unrecorded context; full-page capture can resize screencasts."""
    config = configuration()
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(channel="msedge", headless=True)
        page = browser.new_page(viewport={"width": 1600, "height": 1000})
        page.goto(f"http://{config['host']}:{config['port']}", wait_until="networkidle")
        page.wait_for_function("typeof bootstrap !== 'undefined' && document.querySelectorAll('.node').length === 6")
        page.evaluate("""async id => {
            job = (await (await fetch('/api/jobs/' + id)).json()).job;
            renderJob();
            if (job.report) await showCurrent(job.report);
        }""", job_id)
        for stage in config["stages"]:
            page.locator(f".node[data-stage='{stage['id']}']").click()
            page.wait_for_timeout(200)
            page.screenshot(path=str(OUTPUT / f"node-{stage['id']}-result.png"), full_page=True)
        page.locator("#follow-stage").click()
        page.screenshot(path=str(OUTPUT / "02-completed-workflow.png"), full_page=True)
        browser.close()


def package_delivery():
    manifest = read_json(OUTPUT / "delivery_manifest.json")
    verification = read_json(OUTPUT / "video_continuity_check.json")
    assert verification["status"] == "pass"
    assert verification["decoded_frames"] == manifest["short_video"]["output_frames"]
    manifest["video_verification"] = verification
    write_json(OUTPUT / "delivery_manifest.json", manifest)
    package = OUTPUT / "report-demo-delivery.zip"
    names = {"README.txt", "delivery_manifest.json", "video_continuity_check.json",
             "operation_timeline.json", "word_openability.json", "generated-environment-report.docx",
             "01-dashboard.png", "02-completed-workflow.png", "03-word-cover.png", "04-word-ranking.png",
             "05-esi-and-conclusions.png", "06-chart-preview.png", "07-random-report-review.png"}
    with zipfile.ZipFile(package, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(OUTPUT.iterdir()):
            if path.is_file() and (path.suffix in {".mp4", ".npz"} or path.name.startswith("node-") or path.name in names):
                archive.write(path, path.name)
    with zipfile.ZipFile(package) as archive:
        assert archive.testzip() is None
    return package


def main():
    manifest = read_json(OUTPUT / "delivery_manifest.json")
    assert manifest["status"] == "pass"
    state = read_json(job_root() / manifest["job_id"] / "state.json")
    assert state["status"] == "success" and all(node["status"] == "success" for node in state["stages"])
    assert all(node.get("result") and not node["result"].get("preview_error") for node in state["stages"])
    assert [event["sequence"] for event in state["events"]] == list(range(1, len(state["events"]) + 1))
    report = state["report"]
    document = Document(report["docx"])
    heading = next(index for index, paragraph in enumerate(document.paragraphs) if paragraph.style.name == "Heading 1" and "发展潜力" in paragraph.text)
    recommendations = [paragraph.text for paragraph in document.paragraphs[heading + 1:] if paragraph.text.strip()]
    assert len(recommendations) == 3 and sum(map(len, recommendations)) <= 180
    assert [text[:2] for text in recommendations] == ["1.", "2.", "3."]
    assert report["has_esi"] and report["code"] == "0830" and len(report["sources"]) >= 3
    qa = read_json(Path(report["docx"]).parent / "qa/word_qa_report.json")
    assert qa["overall_status"] == "pass" and not qa["render"]["blank_pages"] and not qa["render"]["sparse_pages"]
    with fitz.open(report["pdf"]) as pdf:
        tail = "\n".join(page.get_text() for page in list(pdf)[-2:])
        assert "六个学科" in tail and "0.231" in tail and "八个学科" not in tail
        assert all(page.get_text().strip() for page in pdf)
    assert hashlib.sha256(Path(report["docx"]).read_bytes()).hexdigest() == hashlib.sha256(Path(manifest["downloaded_docx"]).read_bytes()).hexdigest()
    environment = dict(os.environ, WORD_OPEN_DOCX_PATH=report["docx"], PYTHONIOENCODING="utf-8")
    opened = subprocess.run([sys.executable, "-B", str(PROJECT_DIR / "scripts/word_openability_worker.py")], cwd=PROJECT_DIR,
                            env=environment, capture_output=True, timeout=50, check=True)
    word = json.loads(opened.stdout.decode("utf-8"))
    assert word["word_com_status"] == "opened"
    write_json(OUTPUT / "word_openability.json", word)
    # Export every rendered page in review sheets. The source PNGs remain available for full-size inspection.
    for start in range(0, report["page_count"], 6):
        group = report["pages"][start:start + 6]
        sheet = Image.new("RGB", (1950, 1840), "#e8eee6")
        draw = ImageDraw.Draw(sheet)
        for index, path in enumerate(group):
            with Image.open(path) as page:
                thumb = ImageOps.contain(page.convert("RGB"), (630, 880))
            x, y = index % 3 * 650 + 10, index // 3 * 920 + 30
            sheet.paste(thumb, (x, y))
            draw.text((x, y - 23), f"Page {start + index + 1}", fill="#23483b")
        sheet.save(OUTPUT / f"word-review-{start + 1:02d}.jpg", quality=94)
    short = build_short_video(manifest)
    capture_node_screenshots(manifest["job_id"])
    manifest.update(short_video=short, final_checks={"status": "pass", "page_count": report["page_count"], "blank_pages": [], "sparse_pages": [],
                    "recommendation_points": len(recommendations), "recommendation_characters": sum(map(len, recommendations)),
                    "six_subject_esi_wording": True, "download_hash_matches": True, "word_openability": word["word_com_status"]})
    write_json(OUTPUT / "delivery_manifest.json", manifest)
    (OUTPUT / "README.txt").write_text(f"学科报告操作演示 第二版\n\noperation-demo.mp4：连续操作视频，纯等待阶段为{GENERATION_SPEED}倍速，点击、鼠标轨迹、翻页和节点结果审阅均为原速。\noperation-full.mp4 / operation-full.webm：完整录屏。\nnode-开头的PNG：六个节点的真实结果。\ngenerated-environment-report.docx：通过页面下载的本次报告。\nvideo_frame_mapping.npz：成片帧与原始录屏的逐帧对应。\n\n启动界面：双击项目根目录 start_report_demo.cmd。\n", encoding="utf-8", newline="\n")
    from check_demo_video import main as check_video

    check_video()
    package = package_delivery()
    print(json.dumps({"status": "pass", "package": str(package), "short_video": short, "checks": manifest["final_checks"]}, ensure_ascii=True), flush=True)


if __name__ == "__main__":
    main()
