from __future__ import annotations

import csv
import json
import hashlib
import re
import shutil
import sqlite3
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Emu, Pt, RGBColor
from lxml import etree
from PIL import Image

import sanitize_docx_ooxml_namespaces


PROJECT_DIR = Path(__file__).resolve().parent.parent

TEMPLATE_DOCX = PROJECT_DIR / "data" / "同济大学XX学科的国际排名发展分析-2026年5月.docx"
SLOT_MAPPING_CSV = PROJECT_DIR / "configs" / "template_slot_inventory" / "logical_slot_artifact_mapping.csv"
NARRATIVE_MAPPING_JSON = PROJECT_DIR / "configs" / "template_slot_inventory" / "narrative_paragraph_mapping.json"
DIM_SUBJECT_ALIAS_CSV = PROJECT_DIR / "processed" / "dim_subject_alias.csv"
SQLITE_PATH = PROJECT_DIR / "processed" / "ranking.sqlite"
RANKING_SOURCES = ["软科", "QS", "US News", "THE"]

# 这里显式指定本次试跑使用哪一次学科产物。正常报告请通过 build_subject_word_report.py 注入。
SUBJECT_RUN_DIR = PROJECT_DIR / "outputs" / "subject_agent_runs" / "_set_subject_run_dir_via_build_subject_word_report"

OUTPUT_DIR = PROJECT_DIR / "outputs" / "word_replacement_trial"
OUTPUT_DOCX = OUTPUT_DIR / "同济大学目标学科_模板式图形修正版.docx"
REPLACEMENT_LOG_JSON = OUTPUT_DIR / "replacement_log.json"
REPLACEMENT_AUDIT_JSON = OUTPUT_DIR / "replacement_audit.json"
REPLACEMENT_AUDIT_MD = OUTPUT_DIR / "replacement_audit.md"

# 第一版已验证“图表产物能否进入 Word 原槽位”。
# 现在默认使用原生 Word 表格，保留可编辑性。
TABLE_REPLACEMENT_MODE = "native"  # native / image
INSERT_CAPTION_ONLY_FIGURES = True

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
}


def read_all_slot_mapping() -> list[dict[str, str]]:
    with SLOT_MAPPING_CSV.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def read_word_assembly_manifest() -> dict[str, Any]:
    manifest_path = SUBJECT_RUN_DIR / "word_assembly" / "word_assembly_manifest.json"
    if not manifest_path.exists():
        return {}
    return json.loads(manifest_path.read_text(encoding="utf-8"))


def split_slot_rows_by_manifest(rows: list[dict[str, str]]) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    manifest = read_word_assembly_manifest()
    if not manifest:
        return rows, []
    included = {
        str(item.get("logical_slot_id", "")).strip()
        for item in manifest.get("slot_decisions", [])
        if item.get("include_in_word") is True
    }
    excluded = {
        str(item.get("logical_slot_id", "")).strip()
        for item in manifest.get("slot_decisions", [])
        if item.get("include_in_word") is not True
    }
    if not included:
        return rows, []
    included_rows = [row for row in rows if row.get("logical_slot_id", "").strip() in included]
    excluded_rows = [row for row in rows if row.get("logical_slot_id", "").strip() in excluded]
    return included_rows, excluded_rows


def read_slot_mapping() -> list[dict[str, str]]:
    rows, _ = split_slot_rows_by_manifest(read_all_slot_mapping())
    return rows


def read_narrative_mapping() -> dict[str, dict[str, Any]]:
    payload = json.loads(NARRATIVE_MAPPING_JSON.read_text(encoding="utf-8"))
    items = payload.get("paragraph_mappings", [])
    return {str(item.get("logical_slot_id", "")): item for item in items if item.get("logical_slot_id")}


def read_narrative_content() -> dict[tuple[str, str], str]:
    path = SUBJECT_RUN_DIR / "narrative_generation" / "narrative_generation_result.json"
    if path.exists():
        payload = json.loads(path.read_text(encoding="utf-8"))
        items = payload.get("segment_texts", [])
        result: dict[tuple[str, str], str] = {}
        for item in items:
            slot_id = str(item.get("logical_slot_id", "")).strip()
            segment_key = str(item.get("segment_key", "")).strip()
            text = str(item.get("text", "")).strip()
            if slot_id and segment_key and text:
                result[(slot_id, segment_key)] = text
        if result:
            return result
    return {}


def read_target_subject_name() -> str:
    summary_path = SUBJECT_RUN_DIR / "agent_run_summary.json"
    if summary_path.exists():
        payload = json.loads(summary_path.read_text(encoding="utf-8"))
        request = payload.get("request", {})
        subject_name = str(request.get("target_subject", "")).strip()
        if subject_name:
            return subject_name

    name_parts = SUBJECT_RUN_DIR.name.split("_")
    if name_parts:
        subject_name = name_parts[-1].strip()
        if subject_name:
            return subject_name
    return "未命名学科"


def read_target_subject_payload() -> dict[str, Any]:
    summary_path = SUBJECT_RUN_DIR / "agent_run_summary.json"
    if summary_path.exists():
        return json.loads(summary_path.read_text(encoding="utf-8"))
    return {}


def canonical_the_area(value: Any) -> str:
    area = clean_text(value)
    crosswalk_path = PROJECT_DIR / "configs" / "runtime" / "the_weight_area_crosswalk.json"
    if area and crosswalk_path.exists():
        crosswalk = json.loads(crosswalk_path.read_text(encoding="utf-8"))
        return clean_text(crosswalk.get(area, area))
    return area


def read_target_subject_area() -> str:
    payload = read_target_subject_payload()
    resolved = payload.get("resolved_subject", {})
    value = str(resolved.get("the_subject_area", "")).strip()
    overrides = resolved.get("source_runtime_overrides", {})
    if not value and isinstance(overrides, dict):
        the_override = overrides.get("THE", {})
        if isinstance(the_override, dict):
            value = clean_text(the_override.get("weight_subject_or_area")) or clean_text(the_override.get("ranking_subject_or_area"))
    return clean_text(value) or "自然科学"


def artifact_png_path(artifact_id: str) -> Path | None:
    artifact_dir = SUBJECT_RUN_DIR / "artifacts" / artifact_id
    for filename in ("chart.png", "table.png"):
        path = artifact_dir / filename
        if path.exists():
            return path
    return None


def artifact_svg_path(artifact_id: str) -> Path | None:
    path = SUBJECT_RUN_DIR / "artifacts" / artifact_id / "chart.svg"
    return path if path.exists() else None


def artifact_matrix_path(artifact_id: str) -> Path | None:
    path = SUBJECT_RUN_DIR / "artifacts" / artifact_id / "table_matrix.csv"
    return path if path.exists() else None


def read_csv_matrix(path: Path) -> list[list[str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = [[cell.strip() for cell in row] for row in csv.reader(handle)]
    return [row for row in rows if any(cell for cell in row)]


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() == "nan" else text


def parse_slot_index(reference_object: str, prefix: str) -> int | None:
    for part in reference_object.split("+"):
        part = part.strip()
        if part.startswith(prefix):
            try:
                return int(part.split("_", 1)[1])
            except (IndexError, ValueError):
                return None
    return None


def extract_template_image_targets(docx_path: Path) -> dict[int, str]:
    with ZipFile(docx_path) as zip_file:
        document_xml = zip_file.read("word/document.xml")
        rels_xml = zip_file.read("word/_rels/document.xml.rels")

    root = ET.fromstring(document_xml)
    rel_root = ET.fromstring(rels_xml)
    rels = {rel.attrib.get("Id", ""): rel.attrib.get("Target", "") for rel in rel_root}

    image_targets: dict[int, str] = {}
    image_index = 0
    for drawing in root.findall(".//w:drawing", NS):
        blip = drawing.find(".//a:blip", NS)
        if blip is None:
            continue
        rel_id = blip.attrib.get(f"{{{NS['r']}}}embed", "")
        target = rels.get(rel_id, "")
        if target.startswith("media/"):
            image_index += 1
            image_targets[image_index] = f"word/{target}"
    return image_targets


def extract_template_image_refs(docx_path: Path) -> dict[int, dict[str, str]]:
    with ZipFile(docx_path) as zip_file:
        document_xml = zip_file.read("word/document.xml")
        rels_xml = zip_file.read("word/_rels/document.xml.rels")

    root = ET.fromstring(document_xml)
    rel_root = ET.fromstring(rels_xml)
    rels = {rel.attrib.get("Id", ""): rel.attrib.get("Target", "") for rel in rel_root}

    image_refs: dict[int, dict[str, str]] = {}
    image_index = 0
    for drawing in root.findall(".//w:drawing", NS):
        blip = drawing.find(".//a:blip", NS)
        if blip is None:
            continue
        rel_id = blip.attrib.get(f"{{{NS['r']}}}embed", "")
        target = rels.get(rel_id, "")
        if target.startswith("media/"):
            image_index += 1
            image_refs[image_index] = {
                "rel_id": rel_id,
                "target": target,
                "target_member": f"word/{target}",
            }
    return image_refs


def iter_visible_figure_rows(rows: list[dict[str, str]]) -> list[tuple[int, dict[str, str]]]:
    visible_rows: list[tuple[int, dict[str, str]]] = []
    visible_image_index = 0
    for row in rows:
        is_embedded_figure = row.get("slot_kind") == "figure" and bool(row.get("reference_media"))
        is_virtual_figure = row.get("slot_kind") == "figure" and "virtual_image" in row.get("reference_object", "")
        is_caption_only_figure = row.get("slot_kind") == "figure_caption_without_embedded_object"
        if not is_embedded_figure and not is_virtual_figure and not is_caption_only_figure:
            continue
        visible_image_index += 1
        visible_rows.append((visible_image_index, row))
    return visible_rows


def visible_figure_index_map(rows: list[dict[str, str]]) -> dict[str, int]:
    return {
        row.get("logical_slot_id", ""): visible_index
        for visible_index, row in iter_visible_figure_rows(rows)
        if row.get("logical_slot_id")
    }


def replace_zip_members(source_docx: Path, target_docx: Path, replacements: dict[str, Path]) -> None:
    temp_docx = target_docx.with_suffix(".tmp.docx")
    with ZipFile(source_docx, "r") as zin, ZipFile(temp_docx, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            replacement = replacements.get(item.filename)
            if replacement is None:
                zout.writestr(item, zin.read(item.filename))
            else:
                zout.writestr(item, replacement.read_bytes())
    temp_docx.replace(target_docx)


def replace_docx_xml_member(docx_path: Path, member_name: str, xml_bytes: bytes) -> None:
    temp_docx = docx_path.with_suffix(".tmp.docx")
    with ZipFile(docx_path, "r") as zin, ZipFile(temp_docx, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename == member_name:
                zout.writestr(item, xml_bytes)
            else:
                zout.writestr(item, zin.read(item.filename))
    temp_docx.replace(docx_path)


def inject_svg_image_relationships(docx_path: Path, svg_replacements: list[dict[str, str]]) -> int:
    if not svg_replacements:
        return 0

    new_members = {
        item["new_member"]: Path(item["svg_path"]).read_bytes()
        for item in svg_replacements
        if Path(item["svg_path"]).exists()
    }
    if not new_members:
        return 0

    temp_docx = docx_path.with_suffix(".tmp.docx")
    changed = 0
    with ZipFile(docx_path, "r") as zin, ZipFile(temp_docx, "w", ZIP_DEFLATED) as zout:
        existing_members = set(zin.namelist())
        old_members = {
            f"word/{item['old_target']}"
            for item in svg_replacements
            if item.get("old_target") and not item["old_target"].lower().endswith(".svg")
        }
        for zip_item in zin.infolist():
            if zip_item.filename in new_members:
                continue
            if zip_item.filename in old_members:
                changed += 1
                continue
            data = zin.read(zip_item.filename)
            if zip_item.filename == "word/_rels/document.xml.rels":
                text = data.decode("utf-8")
                for item in svg_replacements:
                    old_target = item["old_target"]
                    new_target = item["new_target"]
                    old_text = f'Target="{old_target}"'
                    new_text = f'Target="{new_target}"'
                    if old_text in text:
                        text = text.replace(old_text, new_text, 1)
                        changed += 1
                data = text.encode("utf-8")
            elif zip_item.filename == "[Content_Types].xml":
                text = data.decode("utf-8")
                if 'Extension="svg"' not in text:
                    text = text.replace(
                        "</Types>",
                        '<Default Extension="svg" ContentType="image/svg+xml"/></Types>',
                    )
                    changed += 1
                data = text.encode("utf-8")
            zout.writestr(zip_item, data)
        for member_name, payload in new_members.items():
            if member_name not in existing_members:
                changed += 1
            zout.writestr(member_name, payload)
    temp_docx.replace(docx_path)
    return changed


def content_width_emu(document: Document) -> int:
    section = document.sections[0]
    return int(section.page_width - section.left_margin - section.right_margin)


def max_section_content_width_emu(document: Document) -> int:
    widths = [int(section.page_width - section.left_margin - section.right_margin) for section in document.sections]
    return max(widths) if widths else content_width_emu(document)


def insert_image_before_table(document: Document, table_index: int, image_path: Path) -> None:
    table = document.tables[table_index - 1]
    paragraph_element = OxmlElement("w:p")
    table._element.addprevious(paragraph_element)
    paragraph = Paragraph(paragraph_element, table._parent)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = 0
    paragraph.paragraph_format.space_after = 0
    paragraph.add_run().add_picture(str(image_path), width=Emu(content_width_emu(document)))
    table._element.getparent().remove(table._element)


def cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell: Any, text: str, *, bold: bool, font_size: float, header: bool) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if header or len(text) <= 12 else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = 0
    paragraph.paragraph_format.space_after = 0
    paragraph.paragraph_format.line_spacing = 0.86 if font_size <= 5.4 else 1
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(font_size)
    run.font.name = "Microsoft YaHei" if header else "SimSun"
    run.font.color.rgb = RGBColor(31, 41, 51)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑" if header else "宋体")


def set_cell_margins(cell: Any, margin_dxa: int = 35) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "left", "bottom", "right"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_dxa))
        node.set(qn("w:type"), "dxa")


def mark_repeat_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def native_table_font_size(column_count: int) -> float:
    if column_count >= 18:
        return 5.2
    if column_count >= 12:
        return 6.4
    if column_count >= 8:
        return 7.0
    if column_count >= 5:
        return 7.5
    return 8.5


def apply_table_widths(table: Any, column_count: int, total_width_emu: int) -> None:
    if column_count <= 0:
        return
    column_width = int(total_width_emu / column_count)
    table.autofit = False
    for row in table.rows:
        for cell in row.cells:
            cell.width = Emu(column_width)


def table_widths_by_header(matrix: list[list[str]], total_width_emu: int) -> list[int]:
    column_count = max(len(row) for row in matrix)
    if not matrix or column_count <= 0:
        return []
    headers = [str(value) for value in matrix[0]]
    if column_count < 12:
        return [int(total_width_emu / column_count)] * column_count

    fixed: list[int | None] = []
    for header in headers:
        if header == "序号":
            fixed.append(int(0.42 * 914400))
        elif "学校" in header:
            fixed.append(int(1.15 * 914400))
        elif header == "省份":
            fixed.append(int(0.48 * 914400))
        elif "平均" in header:
            fixed.append(int(0.62 * 914400))
        else:
            fixed.append(None)
    used = sum(int(value) for value in fixed if value is not None)
    flexible_count = sum(1 for value in fixed if value is None)
    flexible_width = max(int(0.44 * 914400), int((total_width_emu - used) / max(1, flexible_count)))
    widths = [int(value) if value is not None else flexible_width for value in fixed]
    scale = total_width_emu / max(1, sum(widths))
    return [max(int(0.34 * 914400), int(width * scale)) for width in widths]


def apply_table_fixed_widths(table: Any, widths_emu: list[int]) -> None:
    if not widths_emu:
        return
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths_emu) // 635))
    for row in table.rows:
        for column_index, cell in enumerate(row.cells):
            width = widths_emu[min(column_index, len(widths_emu) - 1)]
            cell.width = Emu(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width // 635))


def build_native_table(document: Document, matrix: list[list[str]], total_width_emu: int | None = None) -> Any:
    header = matrix[0] if matrix else []
    filtered = [header]
    for row in matrix[1:]:
        if len(row) >= 2 and row[-1] == "0":
            continue
        filtered.append(row)
    matrix = filtered
    row_count = len(matrix)
    column_count = max(len(row) for row in matrix)
    table = document.add_table(rows=row_count, cols=column_count)
    table.style = "Table Grid"
    font_size = native_table_font_size(column_count)
    target_width = total_width_emu or content_width_emu(document)
    widths = table_widths_by_header(matrix, target_width)
    if widths:
        apply_table_fixed_widths(table, widths)
    else:
        apply_table_widths(table, column_count, target_width)

    decorate_body_rows = column_count <= 10
    for row_index, row in enumerate(matrix):
        cells = table.rows[row_index].cells
        for column_index, cell in enumerate(cells):
            value = row[column_index] if column_index < len(row) else ""
            if decorate_body_rows or row_index == 0:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 10 if column_count >= 18 else 25 if column_count >= 12 else 55)
            set_cell_text(cell, value, bold=row_index == 0, font_size=font_size, header=row_index == 0)
            if row_index == 0:
                cell_shading(cell, "E9EEF5")
            elif decorate_body_rows and row_index % 2 == 0:
                cell_shading(cell, "F7F9FC")
    if table.rows:
        mark_repeat_header(table.rows[0])
    return table


def normalize_existing_table_geometry(document: Document) -> int:
    changed = 0
    for table in document.tables:
        if not table.rows or not table.columns:
            continue
        table_xml = table._tbl.xml
        column_count = len(table.columns)
        if 'w:type="pct"' not in table_xml and 'w:type="fixed"' in table_xml:
            continue
        total_width = max_section_content_width_emu(document) if column_count >= 10 else content_width_emu(document)
        widths = [int(total_width / column_count)] * column_count
        apply_table_fixed_widths(table, widths)
        changed += 1
    return changed


def replace_table_with_native_table(document: Document, table_index: int, matrix: list[list[str]], total_width_emu: int | None = None) -> None:
    old_table = document.tables[table_index - 1]
    new_table = build_native_table(document, matrix, total_width_emu)
    old_table._element.addprevious(new_table._element)
    old_table._element.getparent().remove(old_table._element)


def paragraph_text(paragraph: Paragraph) -> str:
    return "".join(run.text for run in paragraph.runs).strip()


def paragraph_run_style(paragraph: Paragraph) -> dict[str, Any]:
    for run in paragraph.runs:
        if not run.text:
            continue
        rpr = run._element.rPr
        east_asia = None
        if rpr is not None and rpr.rFonts is not None:
            east_asia = rpr.rFonts.get(qn("w:eastAsia"))
        return {
            "font_name": run.font.name,
            "east_asia": east_asia,
            "font_size": run.font.size,
            "bold": run.bold,
            "italic": run.italic,
            "underline": run.underline,
        }
    return {
        "font_name": "仿宋_GB2312",
        "east_asia": "仿宋_GB2312",
        "font_size": Pt(14),
        "bold": False,
        "italic": False,
        "underline": False,
    }


def set_paragraph_text_with_template_style(paragraph: Paragraph, text: str, template_style: dict[str, Any]) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._element)
    run = paragraph.add_run(text)
    run.font.name = template_style.get("font_name") or "仿宋_GB2312"
    run.font.size = template_style.get("font_size") or Pt(14)
    run.bold = bool(template_style.get("bold"))
    run.italic = bool(template_style.get("italic"))
    run.underline = template_style.get("underline")
    east_asia = template_style.get("east_asia") or run.font.name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def replace_subject_placeholder_text(text: str, subject_name: str, the_subject_area: str = "") -> str:
    subject_label = f"{subject_name}学科"
    title = f"同济大学{subject_label}的国际排名发展分析"
    the_area = the_subject_area.strip() or "自然科学"
    replacements = {
        "同济大学数学学科的国际排名发展分析": title,
        "同济大学XX学科的国际排名发展分析": title,
        "同济大学 XX 学科的国际排名发展分析": title,
        "同济大学xx学科的国际排名发展分析": title,
        "同济大学 xx 学科的国际排名发展分析": title,
        "XX 学科": subject_label,
        "XX学科": subject_label,
        "xx 学科": subject_label,
        "xx学科": subject_label,
        "XX 学": subject_label,
        "XX学": subject_label,
        "xx 学": subject_label,
        "xx学": subject_label,
        "数学学科": subject_label,
        "数学学": subject_label,
        "数学学科指标权重": f"{subject_label}指标权重",
        "数学指标权重": f"{subject_name}指标权重",
        "计算机科学于技术": subject_name,
        "THE世界学科排名自然科学领域指标权重": f"THE世界学科排名{the_area}领域指标权重",
        "（待更新）": "",
        "(待更新)": "",
        "待更新": "",
    }
    result = text
    for old, new in replacements.items():
        result = result.replace(old, new)
    result = result.replace(f"图4、软科世界一流学科{subject_name}指标权重", f"图4、QS学科排名{subject_name}指标权重")
    result = result.replace(f"图3、U.S.News学科排名{subject_label}指标权重", f"图5、U.S.News学科排名{subject_label}指标权重")
    result = result.replace(f"图4、THE世界学科排名{the_area}领域指标权重", f"图6、THE世界学科排名{the_area}领域指标权重")
    result = result.replace(
        f"图5、同济{subject_label}在世界学科排名中的历史（排名区间的取值中间位次）",
        f"图7、同济大学{subject_label}在世界学科排名中的历史（排名区间的取值中间位次）",
    )
    result = result.replace(
        f"图6、{subject_label}在U.S.News学科排名中的历史（趋势图分析）",
        f"图8、{subject_label}在U.S.News学科排名中的历史（趋势图分析）",
    )
    result = result.replace("图8、软科排名前200位内地高校分项指标相对排名", "图9、软科排名前200位内地高校分项指标相对排名")
    result = result.replace("图7、U.S.News排名前200位内地高校分项指标相对排名", "图10、U.S.News排名前200位内地高校分项指标相对排名")
    result = result.replace("图9、THE排名前200位内地高校分项指标相对排名", "图11、THE排名前200位内地高校分项指标相对排名")
    result = result.replace("图10、37所内地高校论文及被引总量及相对排名", f"图12、37所高校{subject_label}学科论文及被引总量相对排名")
    result = result.replace("图11、37所内地高校论文CNCI及篇均被引数及相对排名", f"图13、37所高校{subject_label}学科论文CNCI及篇均被引相对排名")
    result = result.replace("图12、37所内地高校高被引TOP论文量及相对排名", f"图14、37所高校{subject_label}学科高被引论文相对排名")
    result = result.replace("图13、37所内地高校国际合著论文占比及相对排名", f"图15、37所高校{subject_label}学科国际合作论文相对排名")
    result = result.replace(f"表3、{subject_label}在软科排名中的细分指标排名历史", f"表1、{subject_label}在软科排名中的细分指标排名历史")
    result = result.replace(f"表4、{subject_label}在THE排名中的细分指标排名历史", f"表3、{subject_label}在THE排名中的细分指标排名历史")
    result = result.replace("2025年发布的四大国际排名中内地高校的位次分布如下", "当前最新四大国际排名中内地高校的位次分布如下")
    result = result.replace("表5、2025年发布的四大国际排名中内地高校位次分布", "表4、当前最新四大国际排名中内地高校位次分布")
    result = result.replace("表7、软科排名中内地高校位于前列的比较", "表5、软科排名中内地高校位于前列的比较")
    result = result.replace("表8、THE排名中内地高校位于前列的比较", "表7、THE排名中内地高校位于前列的比较")
    result = re.sub(r"同济大学\s*(?:XX|xx|数学)\s*学科的国际排名发展分析", title, result)
    result = re.sub(r"(?i)xx\s*学科", subject_label, result)
    result = re.sub(r"(?i)xx\s*学", subject_label, result)
    return result


def iter_container_paragraphs(container: Any) -> list[Paragraph]:
    paragraphs = list(getattr(container, "paragraphs", []))
    for table in getattr(container, "tables", []):
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(iter_container_paragraphs(cell))
    return paragraphs


def replace_subject_placeholders(document: Document, subject_name: str, the_subject_area: str = "") -> int:
    changed = 0
    containers: list[Any] = [document]
    for section in document.sections:
        containers.extend(
            [
                section.header,
                section.first_page_header,
                section.even_page_header,
                section.footer,
                section.first_page_footer,
                section.even_page_footer,
            ]
        )

    seen: set[int] = set()
    for container in containers:
        for paragraph in iter_container_paragraphs(container):
            paragraph_id = id(paragraph._p)
            if paragraph_id in seen:
                continue
            seen.add(paragraph_id)
            original_text = paragraph_text(paragraph)
            replaced_text = replace_subject_placeholder_text(original_text, subject_name, the_subject_area)
            if replaced_text != original_text:
                set_paragraph_text_with_template_style(paragraph, replaced_text, paragraph_run_style(paragraph))
                changed += 1
    return changed


def is_heading_paragraph(text: str) -> bool:
    value = text.strip()
    return bool(
        re.match(r"^[一二三四五六七八九十]+、", value)
        or re.match(r"^\d+、", value)
        or re.match(r"^附表[:：]", value)
        or re.match(r"^第[一二三四五六七八九十]+", value)
    )


def is_caption_paragraph(text: str) -> bool:
    value = text.strip()
    return value.startswith("图") or value.startswith("表") or value.startswith("附表")


def cleaned_reference_caption(row: dict[str, str]) -> str:
    return replace_subject_placeholder_text(
        row.get("reference_caption", "").strip(),
        read_target_subject_name(),
        read_target_subject_area(),
    )


def paragraph_element_has_visual(paragraph_element: Any) -> bool:
    return (
        paragraph_element.find(".//" + qn("w:drawing")) is not None
        or paragraph_element.find(".//" + qn("w:pict")) is not None
    )


def paragraph_element_is_blank(paragraph_element: Any) -> bool:
    return paragraph_element.tag == qn("w:p") and not paragraph_element_text(paragraph_element) and not paragraph_element_has_visual(paragraph_element)


def remove_nearby_visual_paragraphs(paragraph: Paragraph) -> int:
    parent = paragraph._p.getparent()
    removed = 0
    for direction in ("previous", "next"):
        current = paragraph._p.getprevious() if direction == "previous" else paragraph._p.getnext()
        checked = 0
        while current is not None and checked < 5:
            checked += 1
            next_current = current.getprevious() if direction == "previous" else current.getnext()
            if current.tag != qn("w:p"):
                break
            has_visual = paragraph_element_has_visual(current)
            is_blank = paragraph_element_is_blank(current)
            if has_visual or is_blank:
                parent.remove(current)
                removed += 1
                current = next_current
                if has_visual:
                    break
                continue
            break
    return removed


def remove_nearby_table(paragraph: Paragraph) -> int:
    parent = paragraph._p.getparent()
    removed = 0
    for direction in ("previous", "next"):
        current = paragraph._p.getprevious() if direction == "previous" else paragraph._p.getnext()
        checked = 0
        while current is not None and checked < 6:
            checked += 1
            next_current = current.getprevious() if direction == "previous" else current.getnext()
            if current.tag == qn("w:tbl"):
                parent.remove(current)
                removed += 1
                break
            if current.tag == qn("w:p") and paragraph_element_is_blank(current):
                current = next_current
                continue
            break
    return removed


def excluded_slot_note(row: dict[str, str]) -> str:
    subject_name = read_target_subject_name()
    artifact_id = clean_text(row.get("artifact_id", ""))
    section = clean_text(row.get("section", ""))
    artifact_key = artifact_id.lower()
    if "mainland_rank_band_distribution" in artifact_key:
        return (
            f"\u5f53\u524d{subject_name}\u5b66\u79d1\u6682\u65e0\u53ef\u6838\u9a8c\u7684"
            "\u591a\u699c\u5355\u524d200\u5185\u5730\u9ad8\u6821\u4f4d\u6b21\u5206\u5e03\u8868\uff1b"
            "\u672c\u7ae0\u4ec5\u4fdd\u7559\u5df2\u5f62\u6210\u6837\u672c\u7684\u699c\u5355\u5bf9\u6807\u8868\uff0c"
            "\u7f3a\u5931\u699c\u5355\u4e0d\u4ee5\u7a7a\u8868\u6216\u4ee3\u7406\u6837\u672c\u8865\u9f50\u3002"
        )
    if "the_indicator_weight" in artifact_key:
        return f"当前{subject_name}学科暂无可核验的 THE 对应口径同济最新综合排名记录，本报告不生成 THE 权重图；待补充 THE 对应排名和分项数据后再纳入指标权重分析。"
    if "the_tongji_indicator_history" in artifact_key:
        the_area = read_target_subject_area()
        if source_has_tongji_latest("THE"):
            history_summary = summarize_source_history_for_residual_text("THE", subject_name)
            if history_summary:
                return f"{history_summary}但当前暂无可核验的 THE 分项指标历史表，本小节仅呈现综合位次历史，不借用其他榜单或近邻口径展开分项趋势归因。"
            return f"THE 已有{the_area}口径下的同济综合排名历史记录，并已纳入综合排名历史趋势图；但当前暂无可核验的 THE 分项指标历史表，本报告不借用其他榜单或近邻口径替代，待补充 THE 分项数据后再展开分项趋势分析。"
        return f"当前{subject_name}学科暂无可核验的 THE 分项指标历史表，本报告不借用其他榜单或近邻口径替代；待补充 THE 对应分项数据后再展开趋势分析。"
    if "the_mainland_peer_ranking" in artifact_key:
        return f"当前{subject_name}学科暂无可核验的 THE 前200内地高校对标样本，本报告不生成 THE 对标表；相关同行分析以已生成榜单样本为准。"
    if "the_peer_indicator_relative_rank" in artifact_key:
        return f"当前{subject_name}学科暂无可核验的 THE 分项指标相对排名图，本报告不借用其他榜单图表；待补充 THE 对应分项数据后再纳入对标分析。"
    if "the" in artifact_key or "THE" in section:
        return f"当前{subject_name}学科暂无可直接用于本报告该位置的 THE 学科排名或分项指标数据，待补充 THE 对应口径后再生成对应图表。"
    if "qs_indicator_weight" in artifact_key:
        qs_area = (
            source_runtime_override_value("QS", "ranking_subject_or_area")
            or source_runtime_override_value("QS", "weight_subject_or_area")
            or resolved_subject_value("qs_subject_name")
            or subject_name
        )
        if source_has_tongji_latest("QS"):
            return f"当前{subject_name}学科暂无可核验的 QS {qs_area}口径指标权重图，本报告不借用其他 QS 学科或近邻口径替代；已有 QS 综合排名和同行分项数据仍在对应章节呈现。"
        return f"当前{subject_name}学科暂无可核验的 QS {qs_area}口径指标权重图，也暂无可核验的 QS 同济最新综合排名记录；本报告不借用其他 QS 学科或近邻口径替代，待补充对应口径数据后再纳入分析。"
    if "qs_mainland_peer_ranking" in artifact_key:
        return f"当前{subject_name}学科暂无可核验的 QS 前200内地高校对标样本，本报告不生成该榜单对标表；相关同行分析以已生成图表为准。"
    if "qs_peer_indicator_relative_rank" in artifact_key:
        return ""
    if "qs" in artifact_key or "QS" in section:
        return f"当前{subject_name}学科暂无可直接用于本报告该位置的 QS 学科排名或分项指标数据，待补充 QS 对应口径后再生成对应图表。"
    if "usnews_indicator_weight" in artifact_id:
        if source_has_tongji_latest("US News"):
            return f"当前{subject_name}学科暂无 U.S. News 官方对应口径权重数据，本报告仅保留已有 U.S. News 排名与分项历史，不套用其他学科权重图。"
        return f"当前{subject_name}学科暂无 U.S. News 官方对应口径权重数据，也暂无可核验的 U.S. News 同济最新综合排名记录；本报告不套用其他学科权重图，待补充对应口径数据后再纳入分析。"
    if "arwu_tongji_indicator_history" in artifact_key:
        subject_label = arwu_public_subject_label(subject_name)
        return f"当前{subject_name}学科暂无可核验的软科多年分项指标历史表，本报告仅保留软科{subject_label}公开口径的权重和内地高校样本，不生成同济软科分项历史明细表。"
    if "arwu_indicator_weight" in artifact_key:
        return f"当前{subject_name}学科暂无软科世界一流学科排名的直接对应权重数据，本报告不套用其他学科权重图；待补充软科对应口径后再纳入指标权重分析。"
    if "arwu_mainland_peer_ranking" in artifact_key:
        return f"当前{subject_name}学科暂无软科前200内地高校对标样本，本报告不生成软科对标表；相关同行分析以可核验榜单样本为主。"
    if "arwu_peer_indicator_relative_rank" in artifact_key:
        return f"当前{subject_name}学科暂无软科前200内地高校分项指标对标样本，本报告不生成软科分项指标相对排名图；待补充软科排名和分项数据后再纳入对标分析。"
    if "arwu" in artifact_key or "软科" in section:
        return f"当前{subject_name}学科暂无可直接用于本报告该位置的软科排名或分项指标数据，待补充软科对应口径后再生成对应图表。"
    if "ranking_overview" in artifact_id:
        return f"当前{subject_name}学科仅覆盖部分国际榜单最新排名，概览图按可用榜单生成；缺失榜单不作强行补齐。"
    if "commonality" in artifact_id:
        return f"当前{subject_name}学科仅覆盖部分榜单权重或分项指标，本报告按可用来源说明指标共性，不强行补齐缺失榜单。"
    if "paper_" in artifact_key:
        return ""
    return ""


def replace_excluded_narrative_paragraphs(document: Document, row: dict[str, str], note: str) -> int:
    if not note:
        return 0
    heading = find_heading_for_logical_section(document, row.get("section", ""))
    if heading is None:
        return 0
    next_element = heading._p.getnext()
    if next_element is not None and next_element.tag == qn("w:p"):
        next_text = clean_text(paragraph_element_text(next_element))
        if next_text == note:
            return 0
    insert_note_after_heading(heading, note)
    return 1


def find_heading_for_logical_section(document: Document, section: str) -> Paragraph | None:
    text = re.sub(r"\s+", "", clean_text(section))
    major_match = re.match(r"^([一二三四五六七八九十]+)、", text)
    sub_match = re.search(r"、(\d+)、(.+)$", text)
    if not major_match or not sub_match:
        return None
    major_prefix = f"{major_match.group(1)}、"
    sub_prefix = f"{sub_match.group(1)}、"
    sub_tail = sub_match.group(2)
    source_markers = [marker for marker in ["THE", "QS", "U.S.News", "U.S.News", "USNews", "软科"] if marker in sub_tail]
    major_index: int | None = None
    for index, paragraph in enumerate(document.paragraphs):
        paragraph_text_compact = re.sub(r"\s+", "", paragraph_text(paragraph))
        if getattr(getattr(paragraph, "style", None), "name", "") == "Heading 1" and paragraph_text_compact.startswith(major_prefix):
            major_index = index
            break
    if major_index is None:
        return None
    for paragraph in document.paragraphs[major_index + 1 :]:
        style_name = getattr(getattr(paragraph, "style", None), "name", "") or ""
        paragraph_text_compact = re.sub(r"\s+", "", paragraph_text(paragraph))
        if style_name == "Heading 1":
            break
        if style_name != "Heading 2":
            continue
        if not paragraph_text_compact.startswith(sub_prefix):
            continue
        if source_markers and not any(marker in paragraph_text_compact.replace(" ", "") for marker in source_markers):
            continue
        return paragraph
    return None


def subsection_heading_from_section(section: str) -> str:
    parts = [part.strip() for part in clean_text(section).split("、") if part.strip()]
    if len(parts) >= 3:
        return f"{parts[1]}、{parts[2]}"
    if len(parts) >= 2:
        return "、".join(parts[1:])
    return clean_text(section)


def insert_note_after_heading(paragraph: Paragraph, note: str) -> None:
    new_element = OxmlElement("w:p")
    paragraph._element.addnext(new_element)
    inserted = Paragraph(new_element, paragraph._parent)
    inserted.style = paragraph.part.document.styles["Normal"]
    inserted.paragraph_format.space_before = Pt(0)
    inserted.paragraph_format.space_after = Pt(6)
    inserted.paragraph_format.first_line_indent = Pt(28)
    run = inserted.add_run(note)
    run.font.name = "仿宋_GB2312"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋_GB2312")
    run.font.size = Pt(14)


def remove_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


STALE_TEMPLATE_PARAGRAPH_PATTERNS = (
    # Old ARWU template narrative for a specific source subject. These are
    # not valid cross-subject conclusions and must not survive replacement.
    "\u6821\u53cb\u83b7\u5956",
    "\u6559\u5e08\u83b7\u5956",
    "N&S\u8bba\u6587",
    "\u77ed\u671f\u770b\u540c\u6d4e\u96be\u6709\u9876\u5c16\u6821\u53cb",
    "\u6240\u6709\u76f8\u5bf9\u6392\u540d\u4e3a200\u591a\u4f4d",
    "\u56db\u5927\u6392\u540d\u4e2dU.S.News\u6392\u540d\u6700\u770b\u91cd\u79d1\u7814\u4ea7\u51fa",
    "THE\u57282023\u5e74\u8c03\u6574\u4e86\u6392\u540d\u6307\u6807\u4f53\u7cfb",
)


def remove_stale_template_paragraphs(document: Document) -> int:
    changed = 0
    for paragraph in list(document.paragraphs):
        text = paragraph_text(paragraph)
        if not text:
            continue
        normalized = re.sub(r"\s+", "", text)
        if any(pattern in normalized for pattern in STALE_TEMPLATE_PARAGRAPH_PATTERNS):
            remove_paragraph(paragraph)
            changed += 1
    return changed


def keep_table_captions_with_tables(document: Document) -> int:
    body = document.element.body
    changed = 0
    for caption_element in list(body):
        if caption_element.tag != qn("w:p"):
            continue
        caption_text = re.sub(r"\s+", "", paragraph_element_text(caption_element))
        if not (caption_text.startswith("\u8868") or caption_text.startswith("\u9644\u8868")):
            continue

        current = caption_element.getnext()
        intervening_paragraphs = 0
        has_intervening_text = False
        target_table = None
        while current is not None and intervening_paragraphs <= 4:
            if current.tag == qn("w:tbl"):
                target_table = current
                break
            if current.tag != qn("w:p"):
                break
            text = re.sub(r"\s+", "", paragraph_element_text(current))
            if text:
                if text.startswith(("\u8868", "\u9644\u8868", "\u56fe")):
                    break
                has_intervening_text = True
            intervening_paragraphs += 1
            current = current.getnext()

        if target_table is not None and has_intervening_text:
            caption_element.addnext(target_table)
            changed += 1
    return changed


def remove_template_narrative_texts(document: Document, row: dict[str, str], anchor_index: int) -> int:
    mapping = read_narrative_mapping().get(clean_text(row.get("logical_slot_id", "")))
    if not mapping:
        return 0
    template_texts = [
        clean_text(paragraph_spec.get("template_text", ""))
        for paragraph_spec in mapping.get("paragraphs", [])
        if clean_text(paragraph_spec.get("template_text", ""))
    ]
    if not template_texts:
        return 0

    removed = 0
    idx = anchor_index + 1
    while idx < len(document.paragraphs):
        paragraph = document.paragraphs[idx]
        text = clean_text(paragraph_text(paragraph))
        if idx > anchor_index + 1 and is_word_heading1(paragraph):
            break
        if idx > anchor_index + 1 and getattr(getattr(paragraph, "style", None), "name", "") == "Heading 2":
            break
        if any(template_text and template_text in text for template_text in template_texts):
            remove_paragraph(paragraph)
            removed += 1
            continue
        idx += 1
    return removed


def normalize_excluded_section_notes(document: Document, excluded_rows: list[dict[str, str]]) -> int:
    changed = 0
    for row in excluded_rows:
        note = excluded_slot_note(row)
        if not note:
            continue
        heading = subsection_heading_from_section(row.get("section", ""))
        if not heading:
            continue
        for index, paragraph in enumerate(list(document.paragraphs)):
            text = clean_text(paragraph_text(paragraph))
            style_name = getattr(getattr(paragraph, "style", None), "name", "") or ""
            if style_name != "Heading 2" or text != note:
                continue
            set_paragraph_text_with_template_style(paragraph, heading, paragraph_run_style(paragraph))
            insert_note_after_heading(paragraph, note)
            changed += 1
            changed += remove_template_narrative_texts(document, row, index)
            break
    return changed


def find_heading1_containing(document: Document, needle: str) -> Paragraph | None:
    for paragraph in document.paragraphs:
        style_name = getattr(getattr(paragraph, "style", None), "name", "") or ""
        if style_name == "Heading 1" and needle in paragraph_text(paragraph):
            return paragraph
    return None


def remove_rank_distribution_intro_block(document: Document, note: str) -> dict[str, int]:
    heading = find_heading1_containing(document, "\u5bf9\u6807\u9ad8\u6821\u6574\u4f53\u5206\u6790")
    if heading is None:
        return {}

    removed_paragraphs = 0
    removed_tables = 0
    current = heading._p.getnext()
    while current is not None:
        next_current = current.getnext()
        if current.tag == qn("w:p"):
            paragraph = Paragraph(current, heading._parent)
            style_name = getattr(getattr(paragraph, "style", None), "name", "") or ""
            if style_name in {"Heading 1", "Heading 2"}:
                break
            parent = current.getparent()
            if parent is not None:
                parent.remove(current)
                removed_paragraphs += 1
        elif current.tag == qn("w:tbl"):
            parent = current.getparent()
            if parent is not None:
                parent.remove(current)
                removed_tables += 1
        current = next_current

    insert_note_after_heading(heading, note)
    return {"removed_paragraphs": removed_paragraphs, "removed_tables": removed_tables}


def remove_excluded_slot_content(document: Document, excluded_rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    log_rows: list[dict[str, Any]] = []
    for row in excluded_rows:
        slot_id = clean_text(row.get("logical_slot_id", ""))
        artifact_id = clean_text(row.get("artifact_id", ""))
        note = excluded_slot_note(row)
        if artifact_id == "table_04_mainland_rank_band_distribution":
            changed = remove_rank_distribution_intro_block(document, note)
            log_rows.append(
                {
                    "slot": slot_id,
                    "artifact_id": artifact_id,
                    "operation": "remove_excluded_slot_content",
                    "status": "ok" if changed else "skipped",
                    "caption_found": False,
                    "removed_visual_paragraphs": 0,
                    "removed_tables": changed.get("removed_tables", 0) if changed else 0,
                    "caption_replaced": changed.get("removed_paragraphs", 0) if changed else 0,
                    "narrative_replaced": 1 if changed else 0,
                    "note": note,
                }
            )
            continue
        caption = cleaned_reference_caption(row)
        caption_index = find_caption_paragraph_index(document, caption)
        caption_paragraph = document.paragraphs[caption_index] if caption_index is not None else None
        removed_visual_paragraphs = 0
        removed_tables = 0
        caption_replaced = 0
        if caption_paragraph is not None:
            if row.get("slot_kind") == "table":
                removed_tables = remove_nearby_table(caption_paragraph)
            else:
                removed_visual_paragraphs = remove_nearby_visual_paragraphs(caption_paragraph)
            remove_paragraph(caption_paragraph)
            caption_replaced = 1
        narrative_replaced = replace_excluded_narrative_paragraphs(document, row, note)
        log_rows.append(
            {
                "slot": slot_id,
                "artifact_id": artifact_id,
                "operation": "remove_excluded_slot_content",
                "status": "ok" if caption_index is not None or narrative_replaced else "skipped",
                "caption_found": caption_index is not None,
                "removed_visual_paragraphs": removed_visual_paragraphs,
                "removed_tables": removed_tables,
                "caption_replaced": caption_replaced,
                "narrative_replaced": narrative_replaced,
                "note": note,
            }
        )
    return log_rows


def read_artifact_rows(artifact_id: str, filename: str) -> list[dict[str, str]]:
    path = SUBJECT_RUN_DIR / "artifacts" / artifact_id / filename
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [
            {str(key): clean_text(value) for key, value in row.items()}
            for row in csv.DictReader(handle)
        ]


def read_source_decision_by_source() -> dict[str, dict[str, str]]:
    path = SUBJECT_RUN_DIR / "source_decision_matrix.csv"
    if not path.exists():
        return {}
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return {
            clean_text(row.get("ranking_source", "")): {str(key): clean_text(value) for key, value in row.items()}
            for row in csv.DictReader(handle)
        }


def source_has_tongji_latest(source: str) -> bool:
    row = read_source_decision_by_source().get(source, {})
    return clean_text(row.get("has_tongji_latest_rank", "")).lower() in {"true", "1", "yes"}


def source_runtime_status(source: str) -> str:
    return clean_text(read_source_decision_by_source().get(source, {}).get("overall_status", ""))


def subject_mapping_key(value: Any) -> str:
    text = clean_text(value).lower().replace("&", " and ")
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "", text)


def read_mapping_table_aliases(source: str) -> list[str]:
    payload = read_target_subject_payload()
    subject_id = clean_text(payload.get("resolved_subject", {}).get("subject_id", ""))
    if not subject_id or not DIM_SUBJECT_ALIAS_CSV.exists():
        return []
    source_key = "US News" if source in {"U.S. News", "US News"} else source
    aliases: list[str] = []
    with DIM_SUBJECT_ALIAS_CSV.open("r", encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            if clean_text(row.get("subject_id", "")) != subject_id:
                continue
            if clean_text(row.get("ranking_source", "")) != source_key:
                continue
            alias = clean_text(row.get("ranking_subject_alias", ""))
            if alias and not is_missing_subject_label(alias):
                aliases.append(alias)
    return ordered_unique(aliases)


def latest_tongji_rank_for_mapping_alias(source: str, alias: str) -> str:
    if not SQLITE_PATH.exists():
        return ""
    source_key = "US News" if source in {"U.S. News", "US News"} else source
    alias_key = subject_mapping_key(alias)
    if not alias_key:
        return ""
    query = """
    SELECT ranking_year, release_year, ranking_subject_raw, ranking_subject_norm, rank_raw, rank_midpoint
    FROM fact_ranking
    WHERE ranking_source = ?
      AND (is_tongji = 1 OR lower(is_tongji) = 'true' OR university_name_cn LIKE '%同济%' OR university_name_en LIKE '%Tongji%')
    """
    try:
        with sqlite3.connect(SQLITE_PATH) as conn:
            rows = conn.execute(query, [source_key]).fetchall()
    except Exception:
        return ""
    matched = []
    for ranking_year, release_year, raw, norm, rank_raw, rank_midpoint in rows:
        raw_key = subject_mapping_key(raw)
        norm_key = subject_mapping_key(norm)
        if alias_key not in {raw_key, norm_key}:
            continue
        year = to_number(ranking_year) or to_number(release_year) or 0
        rank = clean_rank_display(rank_raw) or format_number(rank_midpoint)
        if rank:
            matched.append((year, int(year) if year else 0, rank))
    if not matched:
        return ""
    matched.sort(key=lambda item: item[0])
    _, year, rank = matched[-1]
    return f"{year}年第{rank}" if year else f"最新记录第{rank}"


def mapping_table_scope_note_for_source(source: str, selected_label: str) -> str:
    aliases = read_mapping_table_aliases(source)
    if len(aliases) <= 1:
        return ""
    selected_key = subject_mapping_key(selected_label)
    alternatives = [alias for alias in aliases if subject_mapping_key(alias) != selected_key]
    if not alternatives:
        return ""
    selected_text = selected_label if selected_label and not is_missing_subject_label(selected_label) else aliases[0]
    details = []
    for alias in alternatives:
        latest = latest_tongji_rank_for_mapping_alias(source, alias)
        if latest:
            details.append(f"“{alias}”口径有同济{latest}")
        else:
            details.append(f"“{alias}”口径暂未匹配到同济最新综合位次")
    source_label = display_source_name(source)
    return (
        f"映射表中，{source_label}还对应{'、'.join(f'“{alias}”' for alias in alternatives)}。"
        f"本报告图表以“{selected_text}”作为主呈现口径；"
        + "；".join(details)
        + "，因此不另行制图，只在文字中作为同榜单映射背景提示。"
    )


def arwu_public_subject_label(default: str = "") -> str:
    rows = read_artifact_rows("figure_03_arwu_indicator_weight", "chart_data.csv")
    if rows:
        value = clean_text(rows[0].get("ranking_subject_or_area", ""))
        if value:
            return value
    return default


def artifact_file_exists(artifact_id: str, filename: str) -> bool:
    return (SUBJECT_RUN_DIR / "artifacts" / artifact_id / filename).exists()


def first_present(row: dict[str, str], keys: list[str]) -> str:
    for key in keys:
        value = clean_text(row.get(key, ""))
        if value:
            return value
    return ""


def ordered_unique(values: list[Any]) -> list[str]:
    seen = set()
    rows: list[str] = []
    for value in values:
        text = clean_text(value)
        if not text or text in seen:
            continue
        seen.add(text)
        rows.append(text)
    return rows


def to_number(value: Any) -> float | None:
    text = clean_rank_display(value).replace(",", "")
    if not text or text in {"-", "未上榜", "未发布"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def clean_rank_display(value: Any) -> str:
    text = clean_text(value)
    if not text:
        return ""
    text = re.sub(r"#+\s*(?=\d)", "", text)
    text = re.sub(r"=+\s*(?=\d)", "", text)
    return text.strip()


def format_number(value: Any, digits: int = 1) -> str:
    number = to_number(value)
    if number is None:
        return clean_rank_display(value)
    if abs(number - round(number)) < 1e-8:
        return str(int(round(number)))
    return f"{number:.{digits}f}"


def normalize_brand_text(text: str) -> str:
    replacement = re.sub(r"U\s*\\?\.?\s*S\s*\\?\.?\s*News", "U.S. News", text)
    replacement = re.sub(r"US\s+News", "U.S. News", replacement)
    replacement = re.sub(r"U\.S\.\s*News", "U.S. News", replacement)
    replacement = re.sub(r"U\.S\. News\s+", "U.S. News ", replacement)
    return replacement


def display_source_name(value: Any) -> str:
    source = clean_text(value)
    if normalize_brand_text(source) == "U.S. News":
        return "U.S. News"
    return source


def resolved_subject_value(key: str, default: str = "") -> str:
    payload = read_target_subject_payload()
    resolved = payload.get("resolved_subject", {})
    return clean_text(resolved.get(key, "")) or default


def source_runtime_override_value(source: str, key: str, default: str = "") -> str:
    payload = read_target_subject_payload()
    resolved = payload.get("resolved_subject", {})
    overrides = resolved.get("source_runtime_overrides", {})
    if not isinstance(overrides, dict):
        return default
    source_override = overrides.get(source, {})
    if not isinstance(source_override, dict):
        return default
    return clean_text(source_override.get(key, "")) or default


def report_subject_label_for_source(source: str, default: str = "") -> str:
    if source == "软科":
        return arwu_public_subject_label(default)
    if source == "QS":
        return (
            report_selected_subject_label("QS", "")
            or source_runtime_override_value("QS", "ranking_subject_or_area")
            or source_runtime_override_value("QS", "weight_subject_or_area")
            or resolved_subject_value("qs_subject_name")
            or default
        )
    if source in {"US News", "U.S. News"}:
        return (
            report_selected_subject_label("US News", "")
            or source_runtime_override_value("US News", "ranking_subject_or_area")
            or source_runtime_override_value("US News", "weight_subject_or_area")
            or resolved_subject_value("usnews_subject_name")
            or default
        )
    if source == "THE":
        return (
            report_selected_subject_label("THE", "")
            or source_runtime_override_value("THE", "ranking_subject_or_area")
            or resolved_subject_value("the_subject_area")
            or source_runtime_override_value("THE", "weight_subject_or_area")
            or default
        )
    return default


def is_missing_subject_label(value: str) -> bool:
    normalized = re.sub(r"\s+", "", clean_text(value))
    return normalized in {"", "-", "None", "none", "nan", "NaN", "\u65e0", "\u6682\u65e0", "\u65e0\u76f4\u63a5\u53e3\u5f84"}


def tongji_data_row(rows: list[dict[str, str]]) -> dict[str, str]:
    for row in rows:
        if clean_text(row.get("is_tongji", "")).lower() in {"1", "true"}:
            return row
        if "同济" in clean_text(row.get("university_name_cn", "")):
            return row
    return {}


def rank_text(row: dict[str, str]) -> str:
    raw = first_present(row, ["rank_raw", "world_rank_raw"])
    return clean_rank_display(raw) or format_number(row.get("rank_midpoint", ""))


def summarize_overview(artifact_id: str) -> str:
    rows = read_artifact_rows(artifact_id, "chart_data.csv")
    if not rows:
        return ""
    parts = []
    scope_parts = []
    mapping_notes = []
    for row in rows:
        source = clean_text(row.get("ranking_source", ""))
        subject = clean_text(row.get("ranking_subject_raw", ""))
        year = first_present(row, ["ranking_year", "release_year"])
        rank = rank_text(row)
        if source and rank:
            subject_part = f"{subject}" if subject else "对应学科"
            parts.append(f"{display_source_name(source)} {year}年{subject_part}第{rank}")
        if source and subject:
            scope_parts.append(f"{display_source_name(source)}：{subject}")
            note = mapping_table_scope_note_for_source(source, subject)
            if note:
                mapping_notes.append(note)
    if not parts:
        return ""
    source_count = len(parts)
    scope = "四类国际学科排名" if source_count >= 4 else f"{source_count}类可用国际学科排名"
    scope_text = f"当前实际采用口径分别为：{'；'.join(ordered_unique(scope_parts))}。" if scope_parts else ""
    mapping_text = "".join(ordered_unique(mapping_notes))
    return (
        f"同济大学{read_target_subject_name()}学科在{scope}中形成最新记录："
        + "，".join(parts)
        + f"。{scope_text}{mapping_text}未覆盖榜单不作强行补齐，解读时应同时说明榜单口径差异和数据缺口。"
    )


def summarize_weight(artifact_id: str) -> str:
    rows = read_artifact_rows(artifact_id, "chart_data.csv")
    if not rows:
        return ""
    source = clean_text(rows[0].get("ranking_source", ""))
    subject = clean_text(rows[0].get("ranking_subject_or_area", ""))
    primary_rows = [
        row
        for row in rows
        if clean_text(row.get("indicator_level", "")) == "一级指标"
        and clean_text(row.get("indicator_name", "")) not in {"合计", "总计"}
        and clean_text(row.get("indicator_group", "")) not in {"合计", "总计"}
    ]
    if primary_rows:
        rows = primary_rows
    group_totals: dict[str, float] = {}
    for row in rows:
        group = clean_text(row.get("indicator_name", "")) if primary_rows else first_present(row, ["indicator_group", "parent_indicator", "raw_group"]) or clean_text(row.get("indicator_name", ""))
        if group in {"合计", "总计"}:
            continue
        weight = to_number(row.get("weight_percent", ""))
        if group and weight is not None:
            group_totals[group] = group_totals.get(group, 0.0) + weight
    top_groups = sorted(group_totals.items(), key=lambda item: item[1], reverse=True)[:3]
    if not top_groups:
        return ""
    group_text = "、".join(f"{name}{value:.1f}%" for name, value in top_groups)
    source_name = display_source_name(source)
    subject_label = subject or read_target_subject_name()
    if source_name == "THE":
        subject_label = read_target_subject_area() or subject_label
    source_status_key = "US News" if source_name == "U.S. News" else source_name
    mapping_note = mapping_table_scope_note_for_source(source_status_key, subject_label)
    if not source_has_tongji_latest(source_status_key):
        return (
            f"{source_name} 在{subject_label}口径下的权重前三为{group_text}。"
            f"当前事实层尚未匹配到同济大学在该{source_name}口径下的最新综合位次，"
            "因此本图仅用于理解该榜单公开评价方法，不用于判断同济综合排名变化。"
            + mapping_note
        )
    if source_name == "QS":
        return f"QS 在{subject_label}口径下，学术声誉和雇主声誉合计占比较高，权重前三为{group_text}。因此 QS 位次提升不能只解释为论文指标改善，还需要同步观察国际同行声誉和雇主评价的积累。" + mapping_note
    if source_name == "U.S. News":
        return f"U.S. News 在{subject_label}口径下，权重前三为{group_text}，科研产出和引用影响构成主要评价基础。后续分析应重点看高被引占比、标准化引文影响和国际合作是否支撑综合位次。" + mapping_note
    if source_name == "THE":
        return f"THE 在{subject_label}口径下权重前三为{group_text}，评价重心较均衡地分布在研究环境、教学和研究质量。该口径下，单纯提升论文规模不足以解释排名变化，还需要看研究质量和国际化维度。" + mapping_note
    if source_name == "软科":
        return f"软科在{subject_label}口径下的权重前三为{group_text}，更偏向科研产出、国际合作和高水平成果。对同济而言，后续应把综合位次变化与高水平论文、论文影响力和一流师资等分项共同判断。" + mapping_note
    return f"{source_name}在{subject_label}口径下的权重主要集中在{group_text}，应与同济大学对应分项表现联动解读。" + mapping_note


def summarize_commonality(artifact_id: str) -> str:
    rows = read_artifact_rows(artifact_id, "chart_data.csv")
    if not rows:
        return ""
    totals: dict[str, float] = {}
    source_count = len({clean_text(row.get("ranking_source", "")) for row in rows if clean_text(row.get("ranking_source", ""))})
    for row in rows:
        group = first_present(row, ["indicator_group", "common_indicator_group", "indicator_category", "display_group", "indicator_name"])
        weight = to_number(first_present(row, ["weight_percent", "weight", "total_weight_percent"]))
        if group and weight is not None:
            totals[group] = totals.get(group, 0.0) + weight
    top_groups = sorted(totals.items(), key=lambda item: item[1], reverse=True)[:3]
    source_scope = "四类排名" if source_count >= 4 else f"{source_count or '多'}类可用排名"
    if top_groups:
        group_text = "、".join(f"{name}{value:.1f}%" for name, value in top_groups)
        return f"{source_scope}的指标并非简单重复，共同重心主要集中在{group_text}。这意味着后续建设不能只追逐单一综合名次，而应区分科研影响、声誉评价和国际化等不同路径。"
    return f"{source_scope}在科研影响、声誉评价、国际化和教学环境等维度上存在交叉，但各榜单权重差异明显。同济的优势和短板需要放回具体榜单口径中判断，避免把综合名次变化解释成单一能力变化。"


def summarize_history(artifact_id: str) -> str:
    rows = read_artifact_rows(artifact_id, "chart_data.csv")
    if not rows:
        return ""
    by_source: dict[str, list[dict[str, str]]] = {}
    for row in rows:
        source = clean_text(row.get("ranking_source", ""))
        if source:
            by_source.setdefault(source, []).append(row)
    parts = []
    for source, source_rows in by_source.items():
        source_rows = sorted(source_rows, key=lambda row: to_number(row.get("ranking_year", "")) or 0)
        first = source_rows[0]
        last = source_rows[-1]
        first_mid = to_number(first.get("rank_midpoint", ""))
        last_mid = to_number(last.get("rank_midpoint", ""))
        if first_mid is not None and last_mid is not None:
            direction = "提升" if last_mid < first_mid else "回落" if last_mid > first_mid else "保持稳定"
            parts.append(f"{display_source_name(source)}由{rank_text(first)}变为{rank_text(last)}，整体{direction}")
        else:
            parts.append(f"{display_source_name(source)}最新位次为{rank_text(last)}")
    if not parts:
        return ""
    return f"从历史序列看，{read_target_subject_name()}学科在不同榜单中的变化并不同步：" + "；".join(parts) + "。因此本节更适合观察长期区间和榜单口径，而不是只比较单一年份。"


def summarize_usnews_indicator_trend(artifact_id: str) -> str:
    rows = read_artifact_rows(artifact_id, "chart_data.csv")
    if not rows:
        return ""
    latest_year = max((to_number(row.get("ranking_year", "")) or 0 for row in rows), default=0)
    latest_rows = [row for row in rows if to_number(row.get("ranking_year", "")) == latest_year]
    rank_rows = [row for row in latest_rows if "排名" in clean_text(row.get("indicator_name", ""))]
    ranked = []
    for row in rank_rows:
        value = to_number(row.get("indicator_value_num", ""))
        if value is not None:
            ranked.append((value, clean_text(row.get("indicator_name", ""))))
    ranked.sort(key=lambda item: item[0])
    strong = "、".join(name for _, name in ranked[:2])
    weak = "、".join(name for _, name in ranked[-2:])
    year_text = str(int(latest_year)) if latest_year else "最新年份"
    return f"U.S. News 分项趋势显示，{year_text}年相对靠前的指标主要是{strong or '部分产出规模指标'}，相对靠后的指标主要是{weak or '部分质量和合作指标'}。这说明综合位次改善之外，还需要继续跟踪论文质量、国际合作和高被引占比等细分项。"


def summarize_peer_indicator(artifact_id: str) -> str:
    rows = read_artifact_rows(artifact_id, "chart_data.csv")
    if not rows:
        return ""
    ranked = []
    for row in rows:
        rel = to_number(row.get("tongji_relative_rank_with_peer_reference", ""))
        if rel is not None:
            name = first_present(row, ["display_indicator_name", "indicator_short_name", "indicator_name"])
            ranked.append((rel, name, row))
    if not ranked:
        return ""
    ranked.sort(key=lambda item: item[0])
    strong = "、".join(name for _, name, _ in ranked[:2] if name)
    weak = "、".join(name for _, name, _ in ranked[-2:] if name)
    peer_size = first_present(ranked[0][2], ["peer_scope_size", "comparison_scope_size"])
    source = "软科" if "figure_09" in artifact_id else "QS" if "figure_10_qs" in artifact_id else "U.S. News" if "figure_11" in artifact_id else "THE"
    subject_label = subject_label_from_artifact_rows(rows, report_subject_label_for_source(source.replace("U.S. News", "US News"), read_target_subject_name()))
    return f"{source}{subject_label}口径分项对标共覆盖{peer_size or '若干'}所内地高校。同济大学相对靠前的指标主要是{strong or '部分优势项'}，相对靠后的指标主要是{weak or '部分短板项'}，建设工作应优先把这些短板与榜单权重结合判断。"


def summarize_peer_table(artifact_id: str) -> str:
    rows = read_artifact_rows(artifact_id, "table_matrix.csv")
    if not rows:
        return ""
    source = "软科" if "table_05" in artifact_id else "U.S. News" if "table_06" in artifact_id else "QS" if "table_07" in artifact_id else "THE" if "table_08" in artifact_id else ""
    full_rows = read_artifact_rows(artifact_id, "latest_peer_records_full_top200.csv")
    subject_label = subject_label_from_artifact_rows(full_rows, read_target_subject_name())
    source_scope = f"{source}{subject_label}口径" if source and subject_label else f"{source}最新口径"
    displayed_count = len(rows)
    full_count = len(full_rows) or displayed_count
    tongji = next((row for row in rows if "同济" in "".join(row.values())), {})
    display_note = f"受篇幅限制，正文列示排名靠前的{displayed_count}所内地高校"
    if tongji:
        values = [value for key, value in tongji.items() if key not in {"序号", "学校名称", "省份"} and value]
        latest_value = values[-1] if values else ""
        return f"{source_scope}下，{display_note}。同济大学历年记录为{'、'.join(values[:4])}，最新位次为{latest_value}；与头部高校相比仍有梯队差距，但可通过连续位次变化判断追赶速度。"
    if source == "软科" and not source_has_tongji_latest("软科"):
        subject_label = arwu_public_subject_label(read_target_subject_name())
        return f"软科{subject_label}口径下，{display_note}。当前事实层尚未匹配到同济大学在该软科口径下的最新综合位次，因此该表仅用于观察头部高校格局，不用于判断同济软科位次或竞争梯队。"
    return f"{source_scope}下，{display_note}。该表的重点不是罗列名单，而是识别头部高校集中度和同济所处的竞争梯队。"


def summarize_distribution(artifact_id: str) -> str:
    rows = read_artifact_rows(artifact_id, "table_data.csv")
    if not rows:
        return ""
    totals = [row for row in rows if clean_text(row.get("位次区间", "")) == "1-200合计"]
    if totals:
        source_scope = "四类榜单" if len(totals) >= 4 else f"{len(totals)}类可用榜单"
        sorted_totals = sorted(
            totals,
            key=lambda row: to_number(row.get("高校数量", "")) or 0,
            reverse=True,
        )
        positive_totals = [row for row in sorted_totals if (to_number(row.get("高校数量", "")) or 0) > 0]
        missing_sources = [clean_text(row.get("榜单", "")) for row in sorted_totals if (to_number(row.get("高校数量", "")) or 0) <= 0]
        missing_text = f"{'、'.join(missing_sources)}暂未形成前200内地高校样本。" if missing_sources else ""
        if len(positive_totals) >= 2:
            largest = clean_text(positive_totals[0].get("榜单", ""))
            smallest = clean_text(positive_totals[-1].get("榜单", ""))
            return f"按学校去重后，{source_scope}前200位内地高校样本规模存在差异。{largest}样本更适合观察完整梯队，{smallest}样本规模较小，主要用于观察可核验头部样本；{missing_text}正文表格受篇幅限制仅展示排名靠前高校，后续对标分析需避免把不同榜单的样本规模直接等同。"
        if positive_totals:
            source = clean_text(positive_totals[0].get("榜单", ""))
            return f"按学校去重后，当前主要由{source}提供可核验前200内地高校样本；{missing_text}正文表格受篇幅限制仅展示排名靠前高校，后续对标分析应把已有样本和缺失榜单分开说明。"
        return f"按学校去重后，{source_scope}暂未形成可用的前200位内地高校样本。后续需先补齐对应榜单口径，再开展同行梯队比较。"
    return ""


def summarize_indicator_history_table(artifact_id: str) -> str:
    rows = read_artifact_rows(artifact_id, "table_matrix.csv")
    if not rows:
        return ""
    columns = list(rows[0].keys())
    year_cols = [col for col in columns if re.search(r"\d{4}", col)]
    indicator_col = next((col for col in columns if "指标" in col or "名称" in col), columns[0])
    indicators = [clean_text(row.get(indicator_col, "")) for row in rows[:4] if clean_text(row.get(indicator_col, ""))]
    source = "软科" if "table_01" in artifact_id else "U.S. News" if "table_02" in artifact_id else "THE"
    if len(year_cols) >= 3 and all(str(year) == str(int(year_cols[0]) + idx) for idx, year in enumerate(year_cols)):
        year_text = f"{year_cols[0]}至{year_cols[-1]}"
    elif len(year_cols) >= 2:
        year_text = "、".join(year_cols)
    else:
        year_text = "最新年份"
    indicator_text = "、".join(indicators[:3]) if indicators else "主要分项指标"
    if "table_01" in artifact_id and len(year_cols) >= 2:
        start_year = year_cols[0]
        end_year = year_cols[-1]
        changes: list[tuple[float, str, str, str]] = []
        stable_items: list[str] = []
        for row in rows:
            indicator = clean_text(row.get(indicator_col, ""))
            numeric_values = [to_number(row.get(year)) for year in year_cols]
            available_values = [value for value in numeric_values if value is not None]
            start_value = to_number(row.get(start_year))
            end_value = to_number(row.get(end_year))
            if not indicator or start_value is None or end_value is None:
                continue
            delta = end_value - start_value
            changes.append((delta, indicator, format_number(start_value), format_number(end_value)))
            if available_values and max(available_values) - min(available_values) <= 2:
                stable_items.append(indicator)
        drops = sorted([item for item in changes if item[0] < -2], key=lambda item: item[0])[:2]
        rises = sorted([item for item in changes if item[0] > 2], key=lambda item: item[0], reverse=True)[:1]
        parts = [f"软科分项历史表覆盖{year_text}，共列出{len(rows)}项指标。"]
        if stable_items:
            parts.append(f"{'、'.join(stable_items[:2])}整体较稳定。")
        if drops:
            drop_text = "；".join(f"{name}由{start}降至{end}" for _, name, start, end in drops)
            parts.append(f"{drop_text}，是需要重点关注的拖累项。")
        if rises:
            rise_text = "；".join(f"{name}由{start}升至{end}" for _, name, start, end in rises)
            parts.append(f"{rise_text}，但仍需结合综合位次判断持续性。")
        return "".join(parts)
    return f"{source}分项历史表覆盖{year_text}，共列出{len(rows)}项指标，重点包括{indicator_text}等。该表的分析价值在于识别综合位次背后的优势项和拖累项，而不是重复说明综合排名变化。"


def summarize_paper(artifact_id: str, segment_key: str) -> str:
    rows = read_artifact_rows(artifact_id, "chart_data.csv")
    tongji = tongji_data_row(rows)
    if not tongji:
        return ""
    paper_field = first_present(tongji, ["paper_field", "Field"])
    field_text = f"在{paper_field}口径下，" if paper_field else "在当前教育部一级学科口径下，"
    if "figure_13" in artifact_id:
        docs = format_number(first_present(tongji, ["WOS_documents", "P"]))
        docs_rank = format_number(first_present(tongji, ["WOS_documents_relative_rank", "P_relative_rank"]))
        cited = format_number(first_present(tongji, ["times_cited", "TCS"]))
        cited_rank = format_number(first_present(tongji, ["times_cited_relative_rank", "TCS_relative_rank"]))
        source_label = "Web of Science 当前学科论文数据" if first_present(tongji, ["WOS_documents"]) else "当前学科论文数据"
        if docs and docs_rank and cited and cited_rank:
            return f"{source_label}中，{field_text}同济大学论文量为{docs}篇，在37校同学科样本中排第{docs_rank}；总被引为{cited}次，排第{cited_rank}。这说明规模和影响力需要结合观察，不能只用论文数量判断学科竞争力。"
        return f"{source_label}已生成当前学科论文规模与引用影响对标图；该口径用于观察同一学科下的相对位置，具体数值以图表数据为准。"
    if "figure_14" in artifact_id:
        if segment_key == "paper_quality_intro":
            return "论文质量部分重点观察 CNCI、篇均被引以及高被引论文相关指标。这些指标相较论文总量更能反映成果影响力，可用于判断学科论文产出的质量结构。"
        cnci = format_number(first_present(tongji, ["CNCI", "MNCS"]), 2)
        cnci_rank = format_number(first_present(tongji, ["CNCI_relative_rank", "MNCS_relative_rank"]))
        impact = format_number(first_present(tongji, ["citation_impact", "MCS"]), 2)
        impact_rank = format_number(first_present(tongji, ["citation_impact_relative_rank", "MCS_relative_rank"]))
        cnci_label = "CNCI" if first_present(tongji, ["CNCI"]) else "MNCS"
        impact_label = "篇均被引" if first_present(tongji, ["citation_impact"]) else "MCS"
        if cnci and cnci_rank and impact and impact_rank:
            return f"{field_text}同济大学 {cnci_label} 为{cnci}，在37校同学科样本中排第{cnci_rank}；{impact_label} 为{impact}，排第{impact_rank}。该组指标更能反映论文质量和规范化影响力，是解释论文短板的重要依据。"
        return "论文质量对标图已按当前学科口径生成，可用于识别规范化影响力和篇均影响力的相对位置。"
    if "figure_15" in artifact_id:
        top1 = format_number(first_present(tongji, ["documents_top_1", "P_top_1"]))
        top10 = format_number(first_present(tongji, ["documents_top_10", "P_top_10"]))
        highly_cited = format_number(first_present(tongji, ["highly_cited_papers"]))
        r1 = format_number(first_present(tongji, ["documents_top_1_relative_rank", "P_top_1_relative_rank"]))
        r10 = format_number(first_present(tongji, ["documents_top_10_relative_rank", "P_top_10_relative_rank"]))
        highly_cited_rank = format_number(first_present(tongji, ["highly_cited_papers_relative_rank"]))
        if top10 and r10:
            return f"高被引论文方面，{field_text}同济大学 Top 1%论文数为{top1}、37校同学科样本中相对排名第{r1}，Top 10%论文数为{top10}、相对排名第{r10}。高影响力成果表现应与 CNCI 和篇均被引共同分析。"
        if highly_cited and highly_cited_rank:
            return f"高被引论文方面，{field_text}同济大学 Top 1%论文数为{top1}、37校同学科样本中相对排名第{r1}，高被引论文数为{highly_cited}、相对排名第{highly_cited_rank}。高影响力成果表现应与 CNCI 和篇均被引共同分析。"
        return f"高被引论文方面，{field_text}同济大学 Top 1%论文数为{top1}、37校同学科样本中相对排名第{r1}。高影响力成果表现应与 CNCI 和篇均被引共同分析。"
    if "figure_16" in artifact_id:
        intl = format_number(first_present(tongji, ["international_documents", "P_int_collab"]))
        pct = format_number(first_present(tongji, ["pct_international_documents", "PP_int_collab"]), 1)
        rank = format_number(first_present(tongji, ["international_documents_relative_rank", "P_int_collab_relative_rank"]))
        return f"国际合作方面，{field_text}同济大学国际合作论文为{intl}篇，占比{pct}%，在37校同学科样本中排第{rank}。该指标体现当前学科国际合作网络活跃度，也会影响部分国际榜单的国际化评价。"
    return ""


def artifact_data_summary_text(row: dict[str, str], segment_key: str) -> str:
    artifact_id = row.get("artifact_id", "")
    if artifact_id == "figure_01_tongji_subject_ranking_overview":
        return summarize_overview(artifact_id)
    if artifact_id == "figure_02_ranking_indicator_commonality":
        return summarize_commonality(artifact_id)
    if "indicator_weight" in artifact_id:
        return summarize_weight(artifact_id)
    if artifact_id == "figure_07_all_rankings_history":
        return summarize_history(artifact_id)
    if artifact_id == "figure_10_usnews_indicator_history_trend":
        return summarize_usnews_indicator_trend(artifact_id)
    if "peer_indicator" in artifact_id:
        return summarize_peer_indicator(artifact_id)
    if artifact_id == "table_04_mainland_rank_band_distribution":
        return summarize_distribution(artifact_id)
    if artifact_id.startswith(("table_01", "table_02", "table_03")):
        return summarize_indicator_history_table(artifact_id)
    if artifact_id.startswith(("table_05", "table_06", "table_07", "table_08")):
        return summarize_peer_table(artifact_id)
    if artifact_id.startswith(("figure_13", "figure_14", "figure_15", "figure_16")):
        return summarize_paper(artifact_id, segment_key)
    return ""


def fallback_narrative_text(row: dict[str, str], segment_key: str = "") -> str:
    subject_name = read_target_subject_name()
    subject_label = f"{subject_name}学科"
    artifact_id = row.get("artifact_id", "")
    slot_kind = row.get("slot_kind", "")
    data_summary = artifact_data_summary_text(row, segment_key)
    if data_summary:
        return data_summary

    if "indicator_weight" in artifact_id:
        return (
            f"该排名体系的指标权重展示了{subject_label}评价中不同维度的相对重要性。"
            "正文应优先关注权重较高的评价维度，并结合学校在对应分项中的表现判断主要优势和短板。"
        )
    if "commonality" in artifact_id:
        return (
            "该图用于对比不同国际排名的指标口径。"
            "四类排名在学术影响、科研产出、声誉评价和国际化等维度上存在交叉，也保留了各自的评价偏重。"
        )
    if "history" in artifact_id and slot_kind != "table":
        return (
            f"排名历史部分呈现了{subject_label}在主要国际榜单中的历年变化。"
            "阅读该图时应同时观察最新位次、长期趋势和榜单口径变化，避免只依据单一年份判断学科发展状态。"
        )
    if "peer_indicator" in artifact_id:
        return (
            "分项指标图把同济大学与内地对标高校放在同一指标口径下比较。"
            "红色标记代表同济大学，可据此识别相对靠前的分项以及需要进一步提升的评价维度。"
        )
    if artifact_id.startswith("figure_13") or artifact_id.startswith("figure_14") or artifact_id.startswith("figure_15") or artifact_id.startswith("figure_16"):
        if segment_key == "paper_quality_intro":
            return (
                "论文质量部分重点观察 CNCI、篇均被引以及高被引论文相关指标。"
                "这些指标相较论文总量更能反映成果影响力，可用于判断学科论文产出的质量结构。"
            )
        if segment_key == "paper_quality_values":
            return (
                f"从相对排名看，{subject_label}需要同时关注规范化影响力和篇均被引表现。"
                "若两项指标与论文规模排名存在差距，建设工作应更强调高质量论文产出和代表性成果培育。"
            )
        if segment_key == "paper_highly_cited_summary":
            return (
                "高被引论文指标反映学科在高影响力成果中的占位。"
                "该指标应与论文总量、总被引和 CNCI 共同判断，避免单独用规模指标替代质量判断。"
            )
        if segment_key == "paper_collab_summary":
            return (
                "国际合著指标反映学科国际合作网络的活跃程度。"
                "若国际合著占比或相对排名偏低，可围绕稳定合作团队、联合发表和国际项目组织能力进行提升。"
            )
        return (
            f"论文指标部分从产出规模、影响力或合作结构补充观察{subject_label}的科研表现。"
            "这部分指标应与排名结果交叉阅读，用于解释排名背后的论文基础和国际合作特征。"
        )
    if slot_kind == "table":
        return (
            "该表列示了相关排名或指标的明细数据。"
            "表格可作为图形解读的依据，重点关注同济大学所处位次、同类高校分布和连续年份中的变化。"
        )
    return (
        f"该图表提供了{subject_label}国际排名分析中的关键证据。"
        "正文应围绕图表中的位次、趋势和对标差距展开具体解释，并形成可落地的建设建议。"
    )


def narrative_placeholder_text(row: dict[str, str], position: str) -> str:
    return fallback_narrative_text(row, position)


def find_caption_paragraph_index(document: Document, caption: str) -> int | None:
    if not caption:
        return None
    for idx, paragraph in enumerate(document.paragraphs):
        text = paragraph_text(paragraph)
        if caption in text:
            return idx
    return None


def find_narrative_paragraph_index(document: Document, caption_index: int, direction: int) -> int | None:
    paragraphs = document.paragraphs
    idx = caption_index + direction
    while 0 <= idx < len(paragraphs):
        text = paragraph_text(paragraphs[idx]).strip()
        if not text:
            idx += direction
            continue
        if is_heading_paragraph(text) or is_caption_paragraph(text):
            return None
        return idx
    return None


def ordered_caption_paragraph_indices(document: Document) -> list[int]:
    return [
        idx
        for idx, paragraph in enumerate(document.paragraphs)
        if is_caption_paragraph(paragraph_text(paragraph))
    ]


def looks_like_bridge_paragraph(text: str) -> bool:
    value = text.strip()
    bridge_markers = [
        "如下",
        "如下：",
        "位次分布如下",
        "比较37所高校的数据",
    ]
    return any(marker in value for marker in bridge_markers)


def choose_narrative_target(
    row: dict[str, str],
    document: Document,
    caption_index: int,
) -> tuple[int | None, str]:
    caption = row.get("reference_caption", "").strip()
    if caption.startswith("附表"):
        return None, "appendix_skip"

    before_index = find_narrative_paragraph_index(document, caption_index, -1)
    after_index = find_narrative_paragraph_index(document, caption_index, 1)
    before_text = paragraph_text(document.paragraphs[before_index]).strip() if before_index is not None else ""
    after_text = paragraph_text(document.paragraphs[after_index]).strip() if after_index is not None else ""

    if row.get("slot_kind") == "table" and before_text and looks_like_bridge_paragraph(before_text) and after_index is not None:
        return after_index, "after_caption"
    if before_index is not None and not looks_like_bridge_paragraph(before_text):
        return before_index, "before_caption"
    if after_index is not None:
        return after_index, "after_caption"
    if before_index is not None:
        return before_index, "before_caption"
    return None, "narrative_paragraph_not_found"


def replace_narrative_placeholders(document: Document, rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    log_rows: list[dict[str, Any]] = []
    caption_indices_in_order = ordered_caption_paragraph_indices(document)
    for row_index, row in enumerate(rows):
        caption = row.get("reference_caption", "").strip()
        caption_index = find_caption_paragraph_index(document, caption)
        caption_match_mode = "exact"
        if caption_index is None and row_index < len(caption_indices_in_order):
            caption_index = caption_indices_in_order[row_index]
            caption_match_mode = "ordered_fallback"
        if caption_index is None:
            log_rows.append(
                {
                    "slot": row["logical_slot_id"],
                    "artifact_id": row["artifact_id"],
                    "operation": "replace_narrative_placeholder",
                    "status": "skipped",
                    "reason": "caption_not_found",
                }
            )
            continue
        target_index, position = choose_narrative_target(row, document, caption_index)
        if target_index is None:
            log_rows.append(
                {
                    "slot": row["logical_slot_id"],
                    "artifact_id": row["artifact_id"],
                    "operation": "replace_narrative_placeholder",
                    "status": "skipped",
                    "reason": position,
                }
            )
            continue
        paragraph = document.paragraphs[target_index]
        original_text = paragraph_text(paragraph)
        replacement_text = narrative_placeholder_text(row, position)
        paragraph.text = replacement_text
        log_rows.append(
            {
                "slot": row["logical_slot_id"],
                "artifact_id": row["artifact_id"],
                "operation": "replace_narrative_placeholder",
                "status": "ok",
                "target_paragraph": f"P{target_index + 1}",
                "position": position,
                "caption_match_mode": caption_match_mode,
                "original_text": original_text,
                "replacement_text": replacement_text,
            }
        )
    return log_rows


def strict_narrative_placeholder_text(row: dict[str, str], paragraph_spec: dict[str, Any]) -> str:
    segment = str(paragraph_spec.get("segment_key", "main"))
    return fallback_narrative_text(row, segment)


def resolve_narrative_text(
    row: dict[str, str],
    paragraph_spec: dict[str, Any],
    narrative_texts: dict[tuple[str, str], str],
) -> tuple[str, str]:
    slot_id = row.get("logical_slot_id", "").strip()
    segment_key = str(paragraph_spec.get("segment_key", "")).strip()
    generated_text = narrative_texts.get((slot_id, segment_key), "").strip()
    if generated_text:
        return generated_text, "generated_text"
    return strict_narrative_placeholder_text(row, paragraph_spec), "placeholder_fallback"


def replace_narrative_placeholders_strict(document: Document, rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    log_rows: list[dict[str, Any]] = []
    mapping_by_slot = read_narrative_mapping()
    narrative_texts = read_narrative_content()
    for row in rows:
        slot_id = row["logical_slot_id"]
        mapping = mapping_by_slot.get(slot_id)
        if not mapping:
            log_rows.append(
                {
                    "slot": slot_id,
                    "artifact_id": row["artifact_id"],
                    "operation": "replace_narrative_placeholder",
                    "status": "skipped",
                    "reason": "mapping_not_found",
                }
            )
            continue
        if str(mapping.get("replace_decision", "replace")) != "replace":
            log_rows.append(
                {
                    "slot": slot_id,
                    "artifact_id": row["artifact_id"],
                    "operation": "replace_narrative_placeholder",
                    "status": "skipped",
                    "reason": str(mapping.get("replace_decision", "skip")),
                    "notes": mapping.get("notes", ""),
                }
            )
            continue
        for paragraph_spec in mapping.get("paragraphs", []):
            paragraph_index = int(paragraph_spec["template_paragraph_index"])
            target_index = paragraph_index - 1
            if target_index < 0 or target_index >= len(document.paragraphs):
                log_rows.append(
                    {
                        "slot": slot_id,
                        "artifact_id": row["artifact_id"],
                        "operation": "replace_narrative_placeholder",
                        "status": "skipped",
                        "reason": "paragraph_out_of_range",
                        "target_paragraph": f"P{paragraph_index}",
                    }
                )
                continue
            paragraph = document.paragraphs[target_index]
            original_text = paragraph_text(paragraph)
            expected_text = str(paragraph_spec.get("template_text", "")).strip()
            template_style = paragraph_run_style(paragraph)
            replacement_text, content_mode = resolve_narrative_text(row, paragraph_spec, narrative_texts)
            set_paragraph_text_with_template_style(paragraph, replacement_text, template_style)
            log_rows.append(
                {
                    "slot": slot_id,
                    "artifact_id": row["artifact_id"],
                    "operation": "replace_narrative_placeholder",
                    "status": "ok",
                    "target_paragraph": f"P{paragraph_index}",
                    "segment_key": paragraph_spec.get("segment_key", ""),
                    "position": paragraph_spec.get("position_relative_to_caption", ""),
                    "content_mode": content_mode,
                    "text_match": "exact" if expected_text == original_text else "mismatch",
                    "original_text": original_text,
                    "replacement_text": replacement_text,
                }
            )
    return log_rows


def insert_image_after_paragraph(document: Document, paragraph: Paragraph, image_path: Path) -> None:
    new_element = OxmlElement("w:p")
    paragraph._element.addnext(new_element)
    inserted = Paragraph(new_element, paragraph._parent)
    inserted.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inserted.paragraph_format.space_before = 0
    inserted.paragraph_format.space_after = 0
    inserted.add_run().add_picture(str(image_path), width=Emu(content_width_emu(document)))


def insert_image_before_paragraph(document: Document, paragraph: Paragraph, image_path: Path) -> None:
    new_element = OxmlElement("w:p")
    paragraph._element.addprevious(new_element)
    inserted = Paragraph(new_element, paragraph._parent)
    inserted.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inserted.paragraph_format.space_before = 0
    inserted.paragraph_format.space_after = 0
    inserted.add_run().add_picture(str(image_path), width=Emu(content_width_emu(document)))


def insert_caption_after_paragraph(paragraph: Paragraph, caption_text: str) -> Paragraph:
    new_element = OxmlElement("w:p")
    paragraph._element.addnext(new_element)
    inserted = Paragraph(new_element, paragraph._parent)
    inserted.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inserted.paragraph_format.space_before = Pt(2)
    inserted.paragraph_format.space_after = Pt(6)
    run = inserted.add_run(caption_text)
    run.font.name = "SimSun"
    run.font.size = Pt(10.5)
    return inserted


def insert_caption_before_paragraph(paragraph: Paragraph, caption_text: str) -> Paragraph:
    new_element = OxmlElement("w:p")
    paragraph._element.addprevious(new_element)
    inserted = Paragraph(new_element, paragraph._parent)
    inserted.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inserted.paragraph_format.space_before = Pt(6)
    inserted.paragraph_format.space_after = Pt(2)
    run = inserted.add_run(caption_text)
    run.font.name = "SimSun"
    run.font.size = Pt(10.5)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return inserted


def insert_native_table_after_paragraph(
    document: Document,
    paragraph: Paragraph,
    matrix: list[list[str]],
    total_width_emu: int | None = None,
) -> Any:
    table = build_native_table(document, matrix, total_width_emu)
    paragraph._element.addnext(table._element)
    return table


def insert_native_table_before_paragraph(
    document: Document,
    paragraph: Paragraph,
    matrix: list[list[str]],
    total_width_emu: int | None = None,
) -> Any:
    table = build_native_table(document, matrix, total_width_emu)
    paragraph._element.addprevious(table._element)
    return table


def find_qs_peer_table_anchor(document: Document) -> tuple[Paragraph | None, str]:
    peer_section_index: int | None = None
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph_text(paragraph).replace(" ", "")
        if "对标高校整体分析" in text:
            peer_section_index = index
            break

    qs_heading_index: int | None = None
    start_index = (peer_section_index + 1) if peer_section_index is not None else 0
    for index in range(start_index, len(document.paragraphs)):
        paragraph = document.paragraphs[index]
        text = paragraph_text(paragraph).replace(" ", "")
        if "QS" in text and "学科排名" in text and (is_heading_paragraph(text) or "2、" in text[:8]):
            qs_heading_index = index
            break

    if qs_heading_index is None:
        return None, "qs_heading_not_found"

    for index in range(qs_heading_index + 1, len(document.paragraphs)):
        raw_text = paragraph_text(document.paragraphs[index])
        text = re.sub(r"\s+", "", raw_text)
        if not text:
            continue
        if "QS" in raw_text and not is_heading_paragraph(text) and not is_caption_paragraph(text):
            return document.paragraphs[index], "after_qs_peer_paragraph"
        if "News" in raw_text or "News" in text:
            return document.paragraphs[index], "before_next_peer_source"
        if is_heading_paragraph(text):
            return document.paragraphs[index], "before_next_heading"

    for index in range(len(document.paragraphs) - 1, qs_heading_index, -1):
        text = paragraph_text(document.paragraphs[index]).strip()
        if text:
            return document.paragraphs[index], "after_last_qs_paragraph"
    return document.paragraphs[qs_heading_index], "after_qs_heading"


def insert_virtual_native_table(document: Document, row: dict[str, str], matrix: list[list[str]]) -> list[dict[str, Any]]:
    artifact_id = row["artifact_id"]
    if artifact_id != "table_07_qs_mainland_peer_ranking":
        return [
            {
                "slot": row["logical_slot_id"],
                "artifact_id": artifact_id,
                "operation": "insert_virtual_native_table",
                "status": "skipped",
                "reason": "unsupported_virtual_table",
            }
        ]

    anchor, anchor_mode = find_qs_peer_table_anchor(document)
    if anchor is None:
        return [
            {
                "slot": row["logical_slot_id"],
                "artifact_id": artifact_id,
                "operation": "insert_virtual_native_table",
                "status": "skipped",
                "reason": anchor_mode,
            }
        ]

    caption_text = cleaned_reference_caption(row) or "表：QS排名中内地高校位于前列的比较"
    table_width = content_width_emu(document)
    if anchor_mode == "before_next_heading":
        insert_caption_before_paragraph(anchor, caption_text)
        insert_native_table_before_paragraph(document, anchor, matrix, table_width)
    else:
        caption = insert_caption_after_paragraph(anchor, caption_text)
        insert_native_table_after_paragraph(document, caption, matrix, table_width)

    return [
        {
            "slot": row["logical_slot_id"],
            "artifact_id": artifact_id,
            "operation": "insert_virtual_native_table",
            "status": "ok",
            "anchor_mode": anchor_mode,
            "source": str(artifact_matrix_path(artifact_id) or ""),
            "rows": len(matrix),
            "cols": max((len(item) for item in matrix), default=0),
        }
    ]


def insert_virtual_native_tables(document: Document, rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    log_rows: list[dict[str, Any]] = []
    for row in rows:
        if row.get("slot_kind") != "table" or "virtual_table" not in row.get("reference_object", ""):
            continue
        matrix_path = artifact_matrix_path(row.get("artifact_id", ""))
        matrix = read_csv_matrix(matrix_path) if matrix_path else []
        if not matrix:
            log_rows.append(
                {
                    "slot": row.get("logical_slot_id", ""),
                    "artifact_id": row.get("artifact_id", ""),
                    "operation": "insert_virtual_native_table",
                    "status": "skipped",
                    "reason": "missing_matrix",
                }
            )
            continue
        log_rows.extend(insert_virtual_native_table(document, row, matrix))
    return log_rows


def qs_history_matrix() -> list[list[str]]:
    rows = [
        row
        for row in read_artifact_rows("figure_07_all_rankings_history", "chart_data.csv")
        if clean_text(row.get("ranking_source", "")).upper() == "QS"
    ]
    rows = sorted(rows, key=lambda row: to_number(row.get("ranking_year", "")) or to_number(row.get("release_year", "")) or 0)
    matrix = [["排名年度", "排名口径", "世界排名", "区间中位位次"]]
    for row in rows:
        year = first_present(row, ["ranking_year", "release_year"])
        subject = first_present(row, ["ranking_subject_raw", "ranking_subject_alias"]) or resolved_subject_value("qs_subject_name", "")
        matrix.append([year, subject, rank_text(row), format_number(row.get("rank_midpoint", ""), 1)])
    return matrix if len(matrix) > 1 else []


def find_qs_history_table_anchor(document: Document) -> tuple[Paragraph | None, str]:
    history_section_index: int | None = None
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph_text(paragraph).replace(" ", "")
        if "排名历史发展分析" in text:
            history_section_index = index
            break
    if history_section_index is None:
        return None, "history_section_not_found"

    qs_heading_index: int | None = None
    for index in range(history_section_index + 1, len(document.paragraphs)):
        text = re.sub(r"\s+", "", paragraph_text(document.paragraphs[index]))
        if "对标高校整体分析" in text:
            break
        if "QS" in text and "学科排名" in text and (is_heading_paragraph(text) or "2、" in text[:8]):
            qs_heading_index = index
            break
    if qs_heading_index is None:
        return None, "qs_history_heading_not_found"

    for index in range(qs_heading_index + 1, len(document.paragraphs)):
        paragraph = document.paragraphs[index]
        raw_text = paragraph_text(paragraph)
        text = re.sub(r"\s+", "", raw_text)
        if not text:
            continue
        if "News" in raw_text or is_heading_paragraph(text):
            return paragraph, "before_next_history_source"
        if "QS" in raw_text and not is_caption_paragraph(text):
            return paragraph, "after_qs_history_paragraph"
    return document.paragraphs[qs_heading_index], "after_qs_history_heading"


def insert_qs_history_table(document: Document) -> list[dict[str, Any]]:
    caption_text = f"表：{read_target_subject_name()}学科在QS学科排名中的综合位次历史"
    if any(caption_text in paragraph_text(paragraph) for paragraph in document.paragraphs):
        return [
            {
                "slot": "qs_history_table_insert",
                "artifact_id": "figure_07_all_rankings_history",
                "operation": "insert_qs_history_table",
                "status": "skipped",
                "reason": "already_inserted",
            }
        ]

    matrix = qs_history_matrix()
    if not matrix:
        return [
            {
                "slot": "qs_history_table_insert",
                "artifact_id": "figure_07_all_rankings_history",
                "operation": "insert_qs_history_table",
                "status": "skipped",
                "reason": "missing_qs_history_rows",
            }
        ]

    anchor, anchor_mode = find_qs_history_table_anchor(document)
    if anchor is None:
        return [
            {
                "slot": "qs_history_table_insert",
                "artifact_id": "figure_07_all_rankings_history",
                "operation": "insert_qs_history_table",
                "status": "skipped",
                "reason": anchor_mode,
            }
        ]

    if anchor_mode == "before_next_history_source":
        insert_caption_before_paragraph(anchor, caption_text)
        insert_native_table_before_paragraph(document, anchor, matrix, content_width_emu(document))
    else:
        caption = insert_caption_after_paragraph(anchor, caption_text)
        insert_native_table_after_paragraph(document, caption, matrix, content_width_emu(document))

    return [
        {
            "slot": "qs_history_table_insert",
            "artifact_id": "figure_07_all_rankings_history",
            "operation": "insert_qs_history_table",
            "status": "ok",
            "anchor_mode": anchor_mode,
            "rows": len(matrix),
            "cols": len(matrix[0]) if matrix else 0,
        }
    ]


def source_history_matrix_for_word(source: str) -> list[list[str]]:
    rows = sorted(
        source_history_rows(source),
        key=lambda row: to_number(row.get("ranking_year", "")) or to_number(row.get("release_year", "")) or 0,
    )
    matrix = [["排名年度", "排名口径", "世界排名", "区间中位位次"]]
    for row in rows:
        year = first_present(row, ["ranking_year", "release_year"])
        subject = first_present(row, ["ranking_subject_raw", "ranking_subject_alias"]) or read_target_subject_name()
        matrix.append([year, subject, rank_text(row), format_number(row.get("rank_midpoint", ""), 1)])
    return matrix if len(matrix) > 1 else []


def find_source_history_table_anchor(document: Document, source: str) -> tuple[Paragraph | None, str]:
    history_section_index: int | None = None
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph_text(paragraph).replace(" ", "")
        if "排名历史发展分析" in text:
            history_section_index = index
            break
    if history_section_index is None:
        return None, "history_section_not_found"

    heading_index: int | None = None
    for index in range(history_section_index + 1, len(document.paragraphs)):
        text = re.sub(r"\s+", "", paragraph_text(document.paragraphs[index]))
        if "对标高校整体分析" in text:
            break
        if source in text and "学科排名" in text and (is_heading_paragraph(text) or "、" in text[:5]):
            heading_index = index
            break
    if heading_index is None:
        return None, f"{source.lower()}_history_heading_not_found"

    for index in range(heading_index + 1, len(document.paragraphs)):
        paragraph = document.paragraphs[index]
        raw_text = paragraph_text(paragraph)
        text = re.sub(r"\s+", "", raw_text)
        if not text:
            continue
        if is_heading_paragraph(text):
            return paragraph, "before_next_history_source"
        if source in raw_text and not is_caption_paragraph(text):
            return paragraph, f"after_{source.lower()}_history_paragraph"
    return document.paragraphs[heading_index], f"after_{source.lower()}_history_heading"


def insert_the_history_table(document: Document) -> list[dict[str, Any]]:
    source = "THE"
    caption_text = f"表：{read_target_subject_name()}学科在THE学科排名中的综合位次历史"
    if any(caption_text in paragraph_text(paragraph) for paragraph in document.paragraphs):
        return [
            {
                "slot": "the_history_table_insert",
                "artifact_id": "figure_07_all_rankings_history",
                "operation": "insert_the_history_table",
                "status": "skipped",
                "reason": "already_inserted",
            }
        ]

    matrix = source_history_matrix_for_word(source)
    if not matrix:
        return [
            {
                "slot": "the_history_table_insert",
                "artifact_id": "figure_07_all_rankings_history",
                "operation": "insert_the_history_table",
                "status": "skipped",
                "reason": "missing_the_history_rows",
            }
        ]

    anchor, anchor_mode = find_source_history_table_anchor(document, source)
    if anchor is None:
        return [
            {
                "slot": "the_history_table_insert",
                "artifact_id": "figure_07_all_rankings_history",
                "operation": "insert_the_history_table",
                "status": "skipped",
                "reason": anchor_mode,
            }
        ]

    if anchor_mode == "before_next_history_source":
        insert_caption_before_paragraph(anchor, caption_text)
        insert_native_table_before_paragraph(document, anchor, matrix, content_width_emu(document))
    else:
        caption = insert_caption_after_paragraph(anchor, caption_text)
        insert_native_table_after_paragraph(document, caption, matrix, content_width_emu(document))

    return [
        {
            "slot": "the_history_table_insert",
            "artifact_id": "figure_07_all_rankings_history",
            "operation": "insert_the_history_table",
            "status": "ok",
            "anchor_mode": anchor_mode,
            "rows": len(matrix),
            "cols": len(matrix[0]) if matrix else 0,
        }
    ]


def remove_duplicate_the_history_summary_after_table(document: Document) -> int:
    summary_prefix = "THE \u5df2\u5f62\u6210\u53ef\u6838\u9a8c\u7684\u7efc\u5408\u4f4d\u6b21\u5386\u53f2\u5e8f\u5217"
    gap_phrase = "\u5206\u9879\u6307\u6807\u5386\u53f2\u8868"
    caption_phrase = "THE\u5b66\u79d1\u6392\u540d\u4e2d\u7684\u7efc\u5408\u4f4d\u6b21\u5386\u53f2"
    the_heading_phrase = "THE\u5b66\u79d1\u6392\u540d"

    changed = 0
    in_the_section = False
    saw_gap_aware_summary = False
    after_history_caption = False
    for paragraph in list(document.paragraphs):
        text = clean_text(paragraph_text(paragraph))
        compact = re.sub(r"\s+", "", text)
        style_name = getattr(getattr(paragraph, "style", None), "name", "") or ""
        if style_name == "Heading 1":
            in_the_section = False
            saw_gap_aware_summary = False
            after_history_caption = False
            continue
        if style_name == "Heading 2":
            in_the_section = the_heading_phrase in compact
            saw_gap_aware_summary = False
            after_history_caption = False
            continue
        if not in_the_section:
            continue
        if text.startswith(summary_prefix) and gap_phrase in text:
            saw_gap_aware_summary = True
            continue
        if caption_phrase in compact:
            after_history_caption = True
            continue
        if (
            after_history_caption
            and saw_gap_aware_summary
            and text.startswith(summary_prefix)
            and gap_phrase not in text
        ):
            remove_paragraph(paragraph)
            changed += 1
            after_history_caption = False
    return changed


def insert_qs_peer_indicator_figure(document: Document) -> list[dict[str, Any]]:
    artifact_id = "figure_10_qs_peer_indicator_relative_rank"
    image_path = artifact_png_path(artifact_id)
    rows = read_artifact_rows(artifact_id, "chart_data.csv")
    if image_path is None or not rows:
        return [
            {
                "slot": "qs_peer_indicator_insert",
                "artifact_id": artifact_id,
                "operation": "insert_qs_peer_indicator_figure",
                "status": "skipped",
                "reason": "artifact_not_generated",
            }
        ]

    subject_name = read_target_subject_name()
    qs_subject_label = (
        report_selected_subject_label("QS", "")
        or subject_name
    )
    indicator_count = len(rows)
    peer_sizes = [to_number(row.get("peer_scope_size", "")) for row in rows]
    peer_count = int(max([value for value in peer_sizes if value is not None] or [0]))
    summary = (
        f"QS 已具备 2026 年分项指标数据，可对前200位内地高校样本进行最新年对标。"
        f"本图围绕 QS {qs_subject_label}口径的{indicator_count}个分项指标，展示同济大学相对同行高校的位置；"
        "多年趋势仍需待 2023-2025 年 QS 分项指标结构化后再展开。"
    )
    if peer_count:
        summary = summary.replace("前200位内地高校样本", f"前200位内地高校样本（{peer_count}所）")

    target: Paragraph | None = None
    section_index: int | None = None
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph_text(paragraph).replace(" ", "")
        if "对标高校分指标分析" in text:
            section_index = index
            break
    if section_index is None:
        return [
            {
                "slot": "qs_peer_indicator_insert",
                "artifact_id": artifact_id,
                "operation": "insert_qs_peer_indicator_figure",
                "status": "skipped",
                "reason": "peer_indicator_section_not_found",
            }
        ]
    qs_heading_index: int | None = None
    if section_index is not None:
        for index in range(section_index + 1, len(document.paragraphs)):
            text = paragraph_text(document.paragraphs[index]).replace(" ", "")
            if "对标高校论文指标分析" in text:
                break
            if "QS" in text and "学科排名" in text and (is_heading_paragraph(text) or "2、" in text[:8]):
                qs_heading_index = index
                break
    if qs_heading_index is not None:
        for index in range(qs_heading_index + 1, len(document.paragraphs)):
            text = paragraph_text(document.paragraphs[index])
            compact = text.replace(" ", "")
            if "U.S.News" in compact or "U.S.News" in text or "U.S. News" in text or is_heading_paragraph(compact):
                break
            if "QS" in text and "分项" in text and not is_caption_paragraph(compact):
                target = document.paragraphs[index]
                break
    if target is None:
        start_index = (qs_heading_index + 1) if qs_heading_index is not None else ((section_index + 1) if section_index is not None else 0)
        for index in range(start_index, len(document.paragraphs)):
            text = paragraph_text(document.paragraphs[index])
            if "对标高校论文指标分析" in text:
                break
            if "3、U.S.News学科排名" in text or "3、U.S. News学科排名" in text or "3、U.S. News 学科排名" in text:
                target = document.paragraphs[index].insert_paragraph_before()
                break
    if target is None:
        return [
            {
                "slot": "qs_peer_indicator_insert",
                "artifact_id": artifact_id,
                "operation": "insert_qs_peer_indicator_figure",
                "status": "skipped",
                "reason": "target_paragraph_not_found",
            }
        ]

    set_paragraph_text_with_template_style(target, summary, paragraph_run_style(target))
    insert_image_after_paragraph(document, target, image_path)
    image_paragraph = Paragraph(target._p.getnext(), target._parent)
    insert_caption_after_paragraph(image_paragraph, f"图10、QS排名前200位内地高校{qs_subject_label}口径分项指标相对排名")
    return [
        {
            "slot": "qs_peer_indicator_insert",
            "artifact_id": artifact_id,
            "operation": "insert_qs_peer_indicator_figure",
            "status": "ok",
            "source_png": str(image_path),
            "indicator_count": indicator_count,
            "peer_count": peer_count,
        }
    ]


def replace_tables_and_caption_only_figures(rows: list[dict[str, str]], source_docx: Path, target_docx: Path) -> list[dict[str, Any]]:
    document = Document(str(source_docx))
    log_rows: list[dict[str, Any]] = []

    table_replacements: list[tuple[int, str, Path | None, list[list[str]] | None]] = []
    for row in rows:
        if row["slot_kind"] != "table":
            continue
        table_index = parse_slot_index(row["reference_object"], "table_")
        image_path = artifact_png_path(row["artifact_id"])
        matrix_path = artifact_matrix_path(row["artifact_id"])
        matrix = read_csv_matrix(matrix_path) if matrix_path and TABLE_REPLACEMENT_MODE == "native" else None
        if table_index is None and TABLE_REPLACEMENT_MODE == "native" and matrix and "virtual_table" in row.get("reference_object", ""):
            continue
        if table_index is None or (TABLE_REPLACEMENT_MODE == "native" and not matrix) or (
            TABLE_REPLACEMENT_MODE == "image" and image_path is None
        ):
            log_rows.append(
                {
                    "slot": row["logical_slot_id"],
                    "artifact_id": row["artifact_id"],
                    "operation": f"replace_table_as_{TABLE_REPLACEMENT_MODE}",
                    "status": "skipped",
                    "reason": "missing_table_index_or_source",
                }
            )
            continue
        table_replacements.append((table_index, row["artifact_id"], image_path, matrix))

    for table_index, artifact_id, image_path, matrix in sorted(table_replacements, reverse=True):
        if TABLE_REPLACEMENT_MODE == "native":
            table_width = max_section_content_width_emu(document) if artifact_id == "table_09_peer_full_ranking_history_appendix" else content_width_emu(document)
            replace_table_with_native_table(document, table_index, matrix or [], table_width)
            source_path = artifact_matrix_path(artifact_id)
        else:
            insert_image_before_table(document, table_index, image_path)  # type: ignore[arg-type]
            source_path = image_path
        log_rows.append(
            {
                "slot": f"table_{table_index:03d}",
                "artifact_id": artifact_id,
                "operation": f"replace_table_as_{TABLE_REPLACEMENT_MODE}",
                "status": "ok",
                "source": str(source_path),
            }
        )

    log_rows.extend(replace_narrative_placeholders_strict(document, rows))

    if INSERT_CAPTION_ONLY_FIGURES:
        caption_only_rows = [
            row
            for row in rows
            if row["slot_kind"] == "figure_caption_without_embedded_object"
            and artifact_png_path(row["artifact_id"]) is not None
        ]
        for row in caption_only_rows:
            caption = row["reference_caption"]
            image_path = artifact_png_path(row["artifact_id"])
            matched = False
            for paragraph in document.paragraphs:
                text = paragraph_text(paragraph)
                if caption and caption in text:
                    insert_image_after_paragraph(document, paragraph, image_path)
                    matched = True
                    log_rows.append(
                        {
                            "slot": row["logical_slot_id"],
                            "artifact_id": row["artifact_id"],
                            "operation": "insert_caption_only_figure",
                            "status": "ok",
                            "source_png": str(image_path),
                        }
                    )
                    break
            if not matched:
                log_rows.append(
                    {
                        "slot": row["logical_slot_id"],
                        "artifact_id": row["artifact_id"],
                        "operation": "insert_caption_only_figure",
                        "status": "skipped",
                        "reason": "caption_not_found",
                    }
                )

    document.save(str(target_docx))
    return log_rows


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def normalize_table_cell(value: Any) -> str:
    text = re.sub(r"\s+", "", clean_text(value))
    return text.replace("U.S.News", "USNews")


def table04_expected_matrix_for_audit(matrix: list[list[str]]) -> list[list[str]]:
    if not matrix:
        return matrix
    header = matrix[0]
    try:
        source_idx = header.index("\u699c\u5355")
        band_idx = header.index("\u4f4d\u6b21\u533a\u95f4")
        count_idx = header.index("\u9ad8\u6821\u6570\u91cf")
    except ValueError:
        return matrix
    totals: dict[str, float] = {}
    for row in matrix[1:]:
        if len(row) <= max(source_idx, band_idx, count_idx):
            continue
        if clean_text(row[band_idx]) != "1-200\u5408\u8ba1":
            continue
        try:
            totals[clean_text(row[source_idx])] = float(clean_text(row[count_idx]) or 0)
        except ValueError:
            totals[clean_text(row[source_idx])] = 0
    if not totals:
        return matrix
    filtered_rows = [
        row
        for row in matrix[1:]
        if totals.get(clean_text(row[source_idx]), 0) > 0
        and clean_text(row[count_idx]) != "0"
    ]
    return [header] + filtered_rows


def document_has_appendix_section(document: Document) -> bool:
    appendix_text = "\u9644\u8868"
    return any(appendix_text in paragraph_text(paragraph) for paragraph in document.paragraphs)


def table_to_matrix(table: Any) -> list[list[str]]:
    return [[cell.text.strip() for cell in table_row.cells] for table_row in table.rows]


def find_table_matrix_for_audit(document: Document, expected_matrix: list[list[str]], preferred_index: int | None) -> tuple[list[list[str]], str]:
    expected_header = [normalize_table_cell(value) for value in expected_matrix[0]] if expected_matrix else []
    candidates: list[tuple[int, list[list[str]], list[str]]] = []
    for index, table in enumerate(document.tables, start=1):
        matrix = table_to_matrix(table)
        actual_header = [normalize_table_cell(value) for value in (matrix[0][: len(expected_header)] if matrix else [])]
        candidates.append((index, matrix, actual_header))
    for index, matrix, actual_header in candidates:
        if actual_header == expected_header and len(matrix) == len(expected_matrix):
            return matrix, f"header_and_row_count_match_table_{index:03d}"
    for index, matrix, actual_header in candidates:
        if actual_header == expected_header:
            return matrix, f"header_match_table_{index:03d}"
    if preferred_index and 1 <= preferred_index <= len(document.tables):
        return table_to_matrix(document.tables[preferred_index - 1]), f"preferred_index_table_{preferred_index:03d}"
    return [], "not_found"


def audit_replacements(rows: list[dict[str, str]], output_docx: Path) -> list[dict[str, Any]]:
    audit_rows: list[dict[str, Any]] = []
    image_targets = extract_template_image_targets(output_docx)

    with ZipFile(output_docx) as zip_file:
        docx_members = set(zip_file.namelist())
        docx_media_hashes = {
            member: sha256_bytes(zip_file.read(member))
            for member in docx_members
            if member.startswith("word/media/")
        }
    visible_media_hashes = {
        member: docx_media_hashes[member]
        for member in image_targets.values()
        if member in docx_media_hashes
    }

    document = Document(str(output_docx))
    visible_index_by_slot = visible_figure_index_map(rows)

    for row in rows:
        artifact_id = row["artifact_id"]
        if row["slot_kind"] == "figure" and row["reference_media"]:
            source_path = artifact_svg_path(artifact_id) or artifact_png_path(artifact_id)
            source_hash = sha256_bytes(source_path.read_bytes()) if source_path else ""
            visible_index = visible_index_by_slot.get(row["logical_slot_id"], -1)
            target_member = image_targets.get(visible_index, "")
            target_hash = docx_media_hashes.get(target_member, "")
            matching_targets = [member for member, docx_hash in visible_media_hashes.items() if source_hash and source_hash == docx_hash]
            status = "ok" if source_hash and target_hash == source_hash else "mismatch"
            audit_rows.append(
                {
                    "logical_slot_id": row["logical_slot_id"],
                    "slot_kind": row["slot_kind"],
                    "artifact_id": artifact_id,
                    "expected_source": str(source_path) if source_path else "",
                    "docx_target": target_member,
                    "matching_visible_targets": ";".join(matching_targets),
                    "check": "visible_media_hash_equal_at_slot_order",
                    "status": status,
                    "visible_image_index": visible_index,
                }
            )
        elif row["slot_kind"] == "figure_caption_without_embedded_object":
            source_png = artifact_png_path(artifact_id)
            source_hash = sha256_bytes(source_png.read_bytes()) if source_png else ""
            visible_index = visible_index_by_slot.get(row["logical_slot_id"], -1)
            target_member = image_targets.get(visible_index, "")
            target_hash = docx_media_hashes.get(target_member, "")
            status = "ok" if source_hash and target_hash == source_hash else "mismatch"
            audit_rows.append(
                {
                    "logical_slot_id": row["logical_slot_id"],
                    "slot_kind": row["slot_kind"],
                    "artifact_id": artifact_id,
                    "expected_source": str(source_png) if source_png else "",
                    "docx_target": target_member,
                    "check": "visible_media_hash_equal_at_slot_order",
                    "status": status,
                    "visible_image_index": visible_index,
                }
            )
        elif row["slot_kind"] == "table":
            if artifact_id == "table_09_peer_full_ranking_history_appendix" and not document_has_appendix_section(document):
                continue
            table_index = parse_slot_index(row["reference_object"], "table_")
            matrix_path = artifact_matrix_path(artifact_id)
            expected_matrix = read_csv_matrix(matrix_path) if matrix_path else []
            if artifact_id == "table_04_mainland_rank_band_distribution":
                expected_matrix = table04_expected_matrix_for_audit(expected_matrix)
            actual_matrix, match_mode = find_table_matrix_for_audit(document, expected_matrix, table_index)
            expected_header = expected_matrix[0] if expected_matrix else []
            actual_header = actual_matrix[0][: len(expected_header)] if actual_matrix else []
            row_count_ok = len(actual_matrix) == len(expected_matrix)
            header_ok = [normalize_table_cell(value) for value in expected_header] == [normalize_table_cell(value) for value in actual_header]
            status = "ok" if row_count_ok and header_ok else "mismatch"
            audit_rows.append(
                {
                    "logical_slot_id": row["logical_slot_id"],
                    "slot_kind": row["slot_kind"],
                    "artifact_id": artifact_id,
                    "expected_source": str(matrix_path) if matrix_path else "",
                    "docx_target": f"table_{table_index:03d}" if table_index else "",
                    "check": "native_table_header_and_row_count",
                    "status": status,
                    "match_mode": match_mode,
                    "expected_rows": len(expected_matrix),
                    "actual_rows": len(actual_matrix),
                    "expected_cols": len(expected_header),
                    "actual_cols": len(actual_header),
                }
            )

    return audit_rows


def write_audit_files(audit_rows: list[dict[str, Any]]) -> None:
    payload = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "output_docx": str(OUTPUT_DOCX),
        "summary": {
            "ok": sum(1 for item in audit_rows if item["status"] == "ok"),
            "mismatch": sum(1 for item in audit_rows if item["status"] != "ok"),
        },
        "checks": audit_rows,
    }
    REPLACEMENT_AUDIT_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8", newline="\n")

    lines = [
        "# Word 图表替换审计",
        "",
        f"- output_docx: `{OUTPUT_DOCX}`",
        f"- ok: {payload['summary']['ok']}",
        f"- mismatch: {payload['summary']['mismatch']}",
        "",
        "| slot | kind | artifact | check | status | target |",
        "|---|---|---|---|---|---|",
    ]
    for item in audit_rows:
        lines.append(
            "| {slot} | {kind} | {artifact} | {check} | {status} | {target} |".format(
                slot=item["logical_slot_id"],
                kind=item["slot_kind"],
                artifact=item["artifact_id"],
                check=item["check"],
                status=item["status"],
                target=item["docx_target"],
            )
        )
    REPLACEMENT_AUDIT_MD.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")


def replace_embedded_images(rows: list[dict[str, str]], source_docx: Path, target_docx: Path) -> list[dict[str, Any]]:
    image_targets = extract_template_image_targets(source_docx)
    replacements: dict[str, Path] = {}
    log_rows: list[dict[str, Any]] = []

    for row in rows:
        if row["slot_kind"] != "figure" or not row["reference_media"]:
            continue
        image_index = parse_slot_index(row["reference_object"], "image_")
        target_member = image_targets.get(image_index or -1)
        image_path = artifact_png_path(row["artifact_id"])
        svg_path = artifact_svg_path(row["artifact_id"])
        if target_member and target_member.lower().endswith(".svg") and svg_path is not None:
            image_path = svg_path
        if image_index is None or image_path is None or target_member is None:
            log_rows.append(
                {
                    "slot": row["logical_slot_id"],
                    "artifact_id": row["artifact_id"],
                    "operation": "replace_embedded_image",
                    "status": "skipped",
                    "reason": "missing_image_index_or_media_or_target",
                }
            )
            continue
        replacements[target_member] = image_path
        log_rows.append(
            {
                "slot": row["logical_slot_id"],
                "artifact_id": row["artifact_id"],
                "operation": "replace_embedded_image",
                "status": "ok",
                "target_member": target_member,
                "source_media": str(image_path),
            }
        )

    if replacements:
        replace_zip_members(target_docx, target_docx, replacements)
    return log_rows


def inject_available_svg_images(rows: list[dict[str, str]], docx_path: Path) -> list[dict[str, Any]]:
    image_refs = extract_template_image_refs(docx_path)
    visible_index_by_slot = visible_figure_index_map(rows)
    svg_replacements: list[dict[str, str]] = []
    log_rows: list[dict[str, Any]] = []
    for row in rows:
        if row["slot_kind"] != "figure" or not row["reference_media"]:
            continue
        svg_path = artifact_svg_path(row["artifact_id"])
        visible_image_index = visible_index_by_slot.get(row["logical_slot_id"])
        image_ref = image_refs.get(visible_image_index or -1, {})
        if svg_path is None or visible_image_index is None or not image_ref:
            continue
        old_target = image_ref["target"]
        if old_target.lower().endswith(".svg"):
            log_rows.append(
                {
                    "slot": row["logical_slot_id"],
                    "artifact_id": row["artifact_id"],
                    "operation": "prepare_svg_image_relationship",
                    "status": "skipped",
                    "reason": "already_svg",
                    "visible_image_index": visible_image_index,
                    "old_target": old_target,
                    "source_svg": str(svg_path),
                }
            )
            continue
        new_target = re.sub(r"\.[A-Za-z0-9]+$", ".svg", old_target)
        new_member = f"word/{new_target}"
        svg_replacements.append(
            {
                "slot": row["logical_slot_id"],
                "artifact_id": row["artifact_id"],
                "old_target": old_target,
                "new_target": new_target,
                "new_member": new_member,
                "svg_path": str(svg_path),
            }
        )
        log_rows.append(
            {
                "slot": row["logical_slot_id"],
                "artifact_id": row["artifact_id"],
                "operation": "prepare_svg_image_relationship",
                "status": "ok",
                "visible_image_index": visible_image_index,
                "old_target": old_target,
                "new_target": new_target,
                "source_svg": str(svg_path),
            }
        )
    svg_change_count = inject_svg_image_relationships(docx_path, svg_replacements)
    if svg_replacements:
        log_rows.append(
            {
                "operation": "inject_svg_image_relationships",
                "status": "ok" if svg_change_count else "skipped",
                "svg_count": len(svg_replacements),
                "changed_items": svg_change_count,
            }
        )
    return log_rows


def append_replacement_log(extra_operations: list[dict[str, Any]], audit_rows: list[dict[str, Any]]) -> None:
    if REPLACEMENT_LOG_JSON.exists():
        payload = json.loads(REPLACEMENT_LOG_JSON.read_text(encoding="utf-8"))
    else:
        payload = {
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "template_docx": str(TEMPLATE_DOCX),
            "subject_run_dir": str(SUBJECT_RUN_DIR),
            "output_docx": str(OUTPUT_DOCX),
            "operations": [],
        }
    operations = payload.setdefault("operations", [])
    operations.extend(extra_operations)
    payload["summary"] = {
        "ok": sum(1 for item in operations if item.get("status") == "ok"),
        "skipped": sum(1 for item in operations if item.get("status") == "skipped"),
        "audit_ok": sum(1 for item in audit_rows if item["status"] == "ok"),
        "audit_mismatch": sum(1 for item in audit_rows if item["status"] != "ok"),
    }
    REPLACEMENT_LOG_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8", newline="\n")


def replace_visible_figure_media_by_order(rows: list[dict[str, str]], docx_path: Path) -> list[dict[str, Any]]:
    image_targets = extract_template_image_targets(docx_path)
    replacements: dict[str, Path] = {}
    log_rows: list[dict[str, Any]] = []
    for visible_image_index, row in iter_visible_figure_rows(rows):
        artifact_id = row.get("artifact_id", "")
        target_member = image_targets.get(visible_image_index, "")
        image_path = artifact_png_path(artifact_id)
        svg_path = artifact_svg_path(artifact_id)
        if target_member.lower().endswith(".svg") and svg_path is not None:
            image_path = svg_path
        if image_path is None or not target_member:
            log_rows.append(
                {
                    "slot": row.get("logical_slot_id", ""),
                    "artifact_id": artifact_id,
                    "operation": "replace_visible_figure_media_by_order",
                    "status": "skipped",
                    "visible_image_index": visible_image_index,
                    "reason": "missing_media_or_visible_target",
                }
            )
            continue
        replacements[target_member] = image_path
        log_rows.append(
            {
                "slot": row.get("logical_slot_id", ""),
                "artifact_id": artifact_id,
                "operation": "replace_visible_figure_media_by_order",
                "status": "ok",
                "visible_image_index": visible_image_index,
                "target_member": target_member,
                "source_media": str(image_path),
            }
        )
    if replacements:
        replace_zip_members(docx_path, docx_path, replacements)
    return log_rows


def refresh_visible_figure_media(rows: list[dict[str, str]], docx_path: Path) -> list[dict[str, Any]]:
    log_rows = replace_visible_figure_media_by_order(rows, docx_path)
    log_rows.extend(inject_available_svg_images(rows, docx_path))
    caption_changes = renumber_captions_in_docx_xml(docx_path)
    log_rows.append(
        {
            "slot": "caption_refresh",
            "artifact_id": "",
            "operation": "renumber_captions_after_media_refresh",
            "status": "ok",
            **caption_changes,
        }
    )
    return log_rows


def normalize_odd_page_sections(docx_path: Path) -> int:
    temp_docx = docx_path.with_suffix(".tmp.docx")
    changed = 0
    with ZipFile(docx_path, "r") as zin, ZipFile(temp_docx, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                changed = data.count(b' w:val="oddPage"')
                data = data.replace(b' w:val="oddPage"', b' w:val="nextPage"')
            zout.writestr(item, data)
    temp_docx.replace(docx_path)
    return changed


def compact_cover_section_break(document: Document) -> int:
    body = document._body._element
    children = list(body)

    first_section_index = None
    for index, child in enumerate(children):
        if child.tag != qn("w:p"):
            continue
        if child.find(".//" + qn("w:sectPr")) is not None:
            first_section_index = index
            break
    if first_section_index is None:
        return 0

    removed = 0
    index = first_section_index - 1
    while index >= 0 and removed < 8:
        paragraph = children[index]
        if paragraph.tag != qn("w:p") or paragraph.find(".//" + qn("w:sectPr")) is not None:
            break
        has_text = bool("".join(text.text or "" for text in paragraph.findall(".//" + qn("w:t"))).strip())
        has_drawing = paragraph.find(".//" + qn("w:drawing")) is not None or paragraph.find(".//" + qn("w:pict")) is not None
        if has_text or has_drawing:
            break
        body.remove(paragraph)
        removed += 1
        index -= 1

    return removed


def is_empty_paragraph_element(paragraph_element: Any) -> bool:
    has_text = bool("".join(text.text or "" for text in paragraph_element.findall(".//" + qn("w:t"))).strip())
    has_drawing = (
        paragraph_element.find(".//" + qn("w:drawing")) is not None
        or paragraph_element.find(".//" + qn("w:pict")) is not None
    )
    has_section = paragraph_element.find(".//" + qn("w:sectPr")) is not None
    return not has_text and not has_drawing and not has_section


def paragraph_element_text(paragraph_element: Any) -> str:
    return "".join(text.text or "" for text in paragraph_element.findall(".//" + qn("w:t"))).strip()


def compact_cover_main_spacing(document: Document) -> int:
    body = document._body._element
    children = [child for child in list(body) if child.tag == qn("w:p")]
    if not children:
        return 0

    text_by_index = {index: paragraph_element_text(child) for index, child in enumerate(children)}
    subtitle_index = next((index for index, text in text_by_index.items() if text == "（第一版）"), None)
    department_index = next((index for index, text in text_by_index.items() if "发展规划与学科建设部" in text), None)
    if subtitle_index is None or department_index is None or department_index <= subtitle_index:
        return 0

    blank_indices = [
        index
        for index in range(subtitle_index + 1, department_index)
        if is_empty_paragraph_element(children[index])
    ]
    keep_blank_count = 5
    remove_indices = blank_indices[: max(0, len(blank_indices) - keep_blank_count)]

    leading_blank_indices = [
        index
        for index in range(0, subtitle_index)
        if is_empty_paragraph_element(children[index])
    ]
    if len(leading_blank_indices) > 1:
        remove_indices.extend(leading_blank_indices[: len(leading_blank_indices) - 1])

    for index in sorted(remove_indices, reverse=True):
        body.remove(children[index])
    return len(remove_indices)


def paragraph_has_page_break(paragraph: Paragraph) -> bool:
    page_breaks = paragraph._p.findall(".//" + qn("w:br"))
    return any(page_break.get(qn("w:type")) == "page" for page_break in page_breaks)


def remove_empty_page_breaks_before_heading1(document: Document) -> int:
    body = document._body._element
    paragraphs = list(document.paragraphs)
    removed = 0
    for index, paragraph in enumerate(paragraphs[:-1]):
        if paragraph_text(paragraph).strip():
            continue
        if not paragraph_has_page_break(paragraph):
            continue
        next_paragraph = paragraphs[index + 1]
        next_text = paragraph_text(next_paragraph).strip()
        next_style = getattr(next_paragraph.style, "style_id", "")
        if next_style != "1" and not re.match(r"^[一二三四五六七八九十]+、", next_text):
            continue
        body.remove(paragraph._p)
        removed += 1
    return removed


def ensure_page_break_before_first_body_heading(document: Document) -> int:
    body = document._body._element
    for paragraph in document.paragraphs:
        text = paragraph_text(paragraph).strip()
        style_id = getattr(paragraph.style, "style_id", "")
        if style_id != "1" and not text.startswith("一、"):
            continue
        previous = paragraph._p.getprevious()
        if previous is not None and previous.tag == qn("w:p"):
            previous_paragraph = Paragraph(previous, paragraph._parent)
            if paragraph_has_page_break(previous_paragraph):
                return 0
        page_break_paragraph = OxmlElement("w:p")
        run = OxmlElement("w:r")
        br = OxmlElement("w:br")
        br.set(qn("w:type"), "page")
        run.append(br)
        page_break_paragraph.append(run)
        paragraph._p.addprevious(page_break_paragraph)
        return 1
    return 0


def summarize_qs_comprehensive_history_for_residual_text(subject_name: str, qs_subject_name: str) -> str:
    rows = [
        row
        for row in read_artifact_rows("figure_07_all_rankings_history", "chart_data.csv")
        if clean_text(row.get("ranking_source", "")).upper() == "QS"
    ]
    if not rows:
        return ""
    rows = sorted(rows, key=lambda row: to_number(row.get("ranking_year", "")) or to_number(row.get("release_year", "")) or 0)
    first = rows[0]
    latest = rows[-1]
    first_year = int(to_number(first.get("ranking_year", "")) or to_number(first.get("release_year", "")) or 0)
    latest_year = int(to_number(latest.get("ranking_year", "")) or to_number(latest.get("release_year", "")) or 0)
    first_mid = to_number(first.get("rank_midpoint", ""))
    latest_mid = to_number(latest.get("rank_midpoint", ""))
    subject_label = report_selected_subject_label("QS", "") or clean_text(first.get("ranking_subject_raw", "")) or clean_text(qs_subject_name) or subject_name
    mapping_note = mapping_table_scope_note_for_source("QS", subject_label)
    if first_mid is not None and latest_mid is not None and len(rows) >= 2:
        direction = "明显提升" if latest_mid < first_mid else "有所回落" if latest_mid > first_mid else "基本稳定"
        return (
            f"QS 已形成可核验的综合位次历史序列，当前采用“{subject_label}”口径："
            f"{first_year}年为{rank_text(first)}，{latest_year}年为{rank_text(latest)}，"
            f"区间中位位次约由{first_mid:g}变为{latest_mid:g}，整体{direction}。"
            "该变化可作为观察国际声誉和同行竞争位势的重要证据；"
            "但分项归因仍需以声誉、雇主评价、篇均被引和 H 指数等指标为边界，避免把综合位次提升直接等同于所有维度改善。"
            + mapping_note
        )
    latest_year_text = f"{latest_year}年" if latest_year else "最新年份"
    return (
        f"QS 学科排名已有{subject_name}学科综合位次记录，当前采用“{subject_label}”口径，"
        f"{latest_year_text}位次为{rank_text(latest)}。该结果可用于观察综合竞争位置；"
        "分项指标多年历史尚未完全结构化，后续应补齐声誉、雇主评价、篇均被引和 H 指数等口径后再展开归因。"
        + mapping_note
    )


def subject_label_from_artifact_rows(rows: list[dict[str, str]], default: str = "") -> str:
    keys = [
        "ranking_subject_raw",
        "ranking_subject_or_area",
        "ranking_subject_alias",
        "ranking_subject",
        "subject_name",
        "Field",
    ]
    for row in rows:
        value = first_present(row, keys)
        if value and not is_missing_subject_label(value):
            return value
    return default


def source_history_rows(source: str) -> list[dict[str, str]]:
    return [
        row
        for row in read_artifact_rows("figure_07_all_rankings_history", "chart_data.csv")
        if clean_text(row.get("ranking_source", "")) == source
    ]


def artifact_subject_or_area(artifact_id: str) -> str:
    spec_path = SUBJECT_RUN_DIR / "artifacts" / artifact_id / "chart_spec.json"
    if not spec_path.exists():
        return ""
    try:
        payload = json.loads(spec_path.read_text(encoding="utf-8-sig"))
    except json.JSONDecodeError:
        return ""
    return clean_text(payload.get("ranking_subject_or_area", ""))


def report_selected_subject_label(source: str, default: str = "") -> str:
    source_key = "US News" if source in {"U.S. News", "US News"} else source
    override = (
        source_runtime_override_value(source_key, "ranking_subject_or_area")
        or source_runtime_override_value(source_key, "weight_subject_or_area")
    )
    if override:
        return override
    history_label = subject_label_from_artifact_rows(source_history_rows(source_key), "")
    if history_label:
        return history_label
    weight_artifact = {
        "QS": "figure_04_qs_indicator_weight",
        "US News": "figure_05_usnews_indicator_weight",
        "U.S. News": "figure_05_usnews_indicator_weight",
        "THE": "figure_06_the_indicator_weight",
        "软科": "figure_03_arwu_indicator_weight",
    }.get(source, "")
    if weight_artifact:
        artifact_label = artifact_subject_or_area(weight_artifact)
        if artifact_label:
            return artifact_label
    return default


def summarize_source_history_for_residual_text(source: str, subject_name: str) -> str:
    rows = source_history_rows(source)
    if not rows:
        return ""
    rows = sorted(rows, key=lambda row: to_number(row.get("ranking_year", "")) or to_number(row.get("release_year", "")) or 0)
    first = rows[0]
    latest = rows[-1]
    first_year = int(to_number(first.get("ranking_year", "")) or to_number(first.get("release_year", "")) or 0)
    latest_year = int(to_number(latest.get("ranking_year", "")) or to_number(latest.get("release_year", "")) or 0)
    first_mid = to_number(first.get("rank_midpoint", ""))
    latest_mid = to_number(latest.get("rank_midpoint", ""))
    source_label = display_source_name(source)
    subject_label = subject_label_from_artifact_rows(rows, subject_name)
    mapping_note = mapping_table_scope_note_for_source(source, subject_label)
    if first_mid is not None and latest_mid is not None and len(rows) >= 2:
        direction = "明显提升" if latest_mid < first_mid else "有所回落" if latest_mid > first_mid else "基本稳定"
        return (
            f"{source_label} 已形成可核验的综合位次历史序列，当前采用“{subject_label}”口径："
            f"{first_year}年为{rank_text(first)}，{latest_year}年为{rank_text(latest)}，"
            f"区间中位位次约由{first_mid:g}变为{latest_mid:g}，整体{direction}。"
            "该变化可用于观察综合竞争位置；分项归因仍需回到该榜单自己的指标口径，避免跨榜单直接套用。"
            + mapping_note
        )
    latest_year_text = f"{latest_year}年" if latest_year else "最新年份"
    return (
        f"{source_label} 学科排名已有{subject_name}学科综合位次记录，当前采用“{subject_label}”口径，"
        f"{latest_year_text}位次为{rank_text(latest)}。该结果可用于观察综合竞争位置，分项原因需以对应榜单指标为边界。"
        + mapping_note
    )


def source_peer_table_artifact_id(source: str) -> str:
    return {
        "软科": "table_05_arwu_mainland_peer_ranking",
        "US News": "table_06_usnews_mainland_peer_ranking",
        "QS": "table_07_qs_mainland_peer_ranking",
        "THE": "table_08_the_mainland_peer_ranking",
    }.get(source, "")


def source_peer_indicator_artifact_id(source: str) -> str:
    return {
        "软科": "figure_09_arwu_peer_indicator_relative_rank",
        "QS": "figure_10_qs_peer_indicator_relative_rank",
        "US News": "figure_11_usnews_peer_indicator_relative_rank",
        "THE": "figure_12_the_peer_indicator_relative_rank",
    }.get(source, "")


def source_indicator_history_artifact_ids(source: str) -> list[str]:
    return {
        "软科": ["table_01_arwu_tongji_indicator_history"],
        "US News": ["table_02_usnews_tongji_indicator_history", "figure_10_usnews_indicator_history_trend"],
        "QS": [],
        "THE": ["table_03_the_tongji_indicator_history"],
    }.get(source, [])


def source_unavailable_text(source: str, subject_name: str, context: str) -> str:
    source_label = display_source_name(source)
    if context == "history":
        return f"当前{subject_name}学科暂无可核验的 {source_label} 综合位次历史序列，本节不使用其他榜单或代理口径替代；待补充对应数据后再展开趋势分析。"
    if context == "peer_table":
        return f"当前{subject_name}学科暂无可核验的 {source_label} 前200内地高校对标样本，本报告不生成该榜单对标表；相关同行分析以已生成图表为准。"
    if context == "peer_indicator":
        return f"当前{subject_name}学科暂无可核验的 {source_label} 分项指标相对排名图，本报告不借用其他榜单图表；待补充对应分项数据后再纳入对标分析。"
    return f"当前{subject_name}学科暂无可核验的 {source_label} 对应数据，待补充后再纳入分析。"


def summarize_source_peer_table_for_residual_text(source: str, subject_name: str) -> str:
    artifact_id = source_peer_table_artifact_id(source)
    if not artifact_id or not read_artifact_rows(artifact_id, "table_matrix.csv"):
        return source_unavailable_text(source, subject_name, "peer_table")
    return summarize_peer_table(artifact_id)


def summarize_source_peer_indicator_for_residual_text(source: str, subject_name: str) -> str:
    artifact_id = source_peer_indicator_artifact_id(source)
    if not artifact_id or not read_artifact_rows(artifact_id, "chart_data.csv"):
        return source_unavailable_text(source, subject_name, "peer_indicator")
    return summarize_peer_indicator(artifact_id)


def source_has_report_evidence(source: str) -> bool:
    if source_history_rows(source):
        return True
    table_id = source_peer_table_artifact_id(source)
    if table_id and read_artifact_rows(table_id, "table_matrix.csv"):
        return True
    figure_id = source_peer_indicator_artifact_id(source)
    if figure_id and read_artifact_rows(figure_id, "chart_data.csv"):
        return True
    return False


def available_source_phrase() -> str:
    labels = [display_source_name(source) for source in RANKING_SOURCES if source_has_report_evidence(source)]
    if not labels:
        return "当前可核验"
    if len(labels) == 1:
        return labels[0]
    return "、".join(labels)


def paragraph_source_context(document: Document, paragraph_index: int) -> str:
    for index in range(paragraph_index, -1, -1):
        text = re.sub(r"\s+", "", paragraph_text(document.paragraphs[index]))
        if not text:
            continue
        if "对标高校论文指标分析" in text or re.match(r"^[六七八九十]+、", text):
            return ""
        if "THE" in text and "学科排名" in text and (is_heading_paragraph(text) or "、" in text[:5]):
            return "THE"
        if ("U.S.News" in text or "U.S.News" in text.replace(" ", "") or "U.S.News" in text or "U.S.News" in text) and "学科排名" in text:
            return "US News"
        if ("U.S.News" in text.replace(" ", "") or "U.S.News" in text or "U.S.News" in text or "U.S.News" in text) and is_heading_paragraph(text):
            return "US News"
        if ("U.S.News" in text.replace(" ", "") or "U.S.News" in text or "U.S.News" in text or "U.S.News" in text) and "排名" in text:
            return "US News"
        if "U.S.News" in text.replace(" ", "") or "U.S.News" in text or "U.S.News" in text or "U.S.News" in text:
            return "US News"
        if "QS" in text and "学科排名" in text and (is_heading_paragraph(text) or "、" in text[:5]):
            return "QS"
        if "软科" in text and "学科排名" in text and (is_heading_paragraph(text) or "、" in text[:5]):
            return "软科"
        if re.match(r"^[一二三四五六七八九十]+、", text):
            return ""
    return ""


def replace_residual_template_paragraphs(document: Document) -> int:
    subject_name = read_target_subject_name()
    qs_direct_subject_name = resolved_subject_value("qs_subject_name", "")
    qs_subject_name = report_selected_subject_label("QS", "") or ("" if is_missing_subject_label(qs_direct_subject_name) else qs_direct_subject_name) or subject_name
    qs_available = source_has_tongji_latest("QS")
    qs_history_summary = summarize_qs_comprehensive_history_for_residual_text(subject_name, qs_subject_name)
    qs_history_text = (
        qs_history_summary
        or f"当前 QS {qs_subject_name}口径可用于观察{subject_name}学科综合位次和同行梯队，但本地数据暂未形成可与 U.S. News、THE 直接对应的多年分项历史表；因此本节主要把 QS 作为综合趋势参考。"
        if qs_available
        else f"当前{subject_name}学科暂无可核验的 QS 同济最新综合排名记录，本节不使用其他 QS 学科或近邻口径替代；待补充对应排名和分项数据后再纳入综合趋势比较。"
    )
    qs_peer_rows = read_artifact_rows("table_07_qs_mainland_peer_ranking", "table_matrix.csv")
    qs_peer_count = len(qs_peer_rows)
    qs_peer_text = (
        f"QS 最新口径下进入前200位的内地高校共{qs_peer_count or '若干'}所，样本规模小于软科和 U.S. News，竞争更集中在头部高校。当前报告使用 QS {qs_subject_name}口径判断综合位次梯队；第五章再结合最新年分项指标观察声誉、雇主评价、篇均被引和 H 指数等具体差距。"
        if qs_peer_rows
        else f"当前{subject_name}学科暂无 QS 前200内地高校对标样本，本报告不生成 QS 对标表；相关同行分析以其他可核验榜单样本为主。"
    )
    qs_relative_rows = read_artifact_rows("figure_10_qs_peer_indicator_relative_rank", "chart_data.csv")
    qs_relative_text = (
        "QS 已形成最新年分项指标相对排名图，可用于观察同济大学在学术声誉、雇主声誉、篇均被引和 H 指数等维度上的相对位置；但多年分项趋势尚未结构化，暂不对年度变化作强归因。"
        if qs_relative_rows
        else f"当前{subject_name}学科暂无 QS 分项指标相对排名图，本报告不借用其他口径图表；待补充 QS 分项数据后再纳入对标分析。"
    )
    qs_has_report_evidence = bool(qs_history_summary or qs_peer_rows or qs_relative_rows)
    the_history_text = (
        summarize_source_history_for_residual_text("THE", subject_name)
        or source_unavailable_text("THE", subject_name, "history")
    )
    the_peer_text = summarize_source_peer_table_for_residual_text("THE", subject_name)
    the_relative_text = summarize_source_peer_indicator_for_residual_text("THE", subject_name)
    arwu_has_tongji = source_has_tongji_latest("软科")
    arwu_label = arwu_public_subject_label(subject_name)
    source_phrase = available_source_phrase()
    conclusion_text = (
        f"综合{source_phrase}学科排名和论文指标看，同济大学{subject_name}学科已有若干可核验优势，也存在论文质量、声誉评价或国际合作等需要继续拆解的环节；软科仅保留{arwu_label}公开口径样本，暂不能过度归因同济软科位次变化。后续重点提升 CNCI、篇均被引和高被引论文表现。"
        if qs_has_report_evidence and not arwu_has_tongji
        else f"综合{source_phrase}学科排名和论文指标看，同济大学{subject_name}学科的综合位次、同行梯队和论文指标并不完全同步。下一阶段应把综合排名改善转化为论文质量、声誉评价和国际合作的同步提升。"
        if qs_has_report_evidence
        else f"综合当前可核验的学科排名和论文指标看，同济大学{subject_name}学科仍需把已有图表证据和缺失榜单口径分开说明；待补齐直接口径后再纳入跨榜单判断。"
    )
    suggestion_text = (
        f"后续将 QS 声誉、U.S. News 科研影响、THE 研究质量与国际化拆成年度任务；软科{arwu_label}口径待补齐同济位次和分项后再归因。"
        if qs_has_report_evidence and not arwu_has_tongji
        else "针对软科、QS 和 THE，应将高水平成果、声誉评价、研究质量和国际化拆成年度任务，并以 CNCI、篇均被引和高被引论文占比作为共性质量抓手。"
        if qs_has_report_evidence
        else "针对软科和 THE，应关注高水平成果、研究质量与国际展望；QS 待选定口径排名补齐后再单独纳入分析。"
    )
    source_context_replacements = {
        "THE排名在2023年更新指标体系后研究质量、国际展望相对排名比较低": {
            "QS": qs_history_text,
            "THE": the_history_text,
            "": the_history_text,
        },
        "2025年发布的THE世界排名中中国内地高校有13所进入前200位": {
            "QS": qs_peer_text,
            "THE": the_peer_text,
            "": the_peer_text,
        },
        "2025年THE排名前200位的3所内地高校": {
            "QS": qs_relative_text,
            "THE": the_relative_text,
            "": the_relative_text,
        },
    }
    replacements = [
        (
            "QS学科排名指标中，学术声誉和雇主声誉占比较重，达60%。",
            qs_history_text
            if not qs_available
            else "QS 学科排名强调声誉、雇主评价和论文影响等维度，解读时需与当前采用口径下的综合位次和分项指标同步判断。",
        ),
        (
            "综合分析三大世界排名的数据",
            conclusion_text,
        ),
        (
            "针对提升U.S.News排名的建议是提升论文质量",
            "针对 U.S. News 等强调科研影响力的榜单，应重点提升标准化引文影响、高被引论文占比和国际合作论文质量。",
        ),
        (
            "对于软科排名可提升的空间是",
            suggestion_text,
        ),
    ]
    changed = 0
    for paragraph_index, paragraph in enumerate(document.paragraphs):
        text = paragraph_text(paragraph)
        context = paragraph_source_context(document, paragraph_index)
        replaced = False
        for marker, by_context in source_context_replacements.items():
            if marker not in text:
                continue
            replacement = by_context.get(context) or by_context.get("")
            set_paragraph_text_with_template_style(paragraph, replacement, paragraph_run_style(paragraph))
            changed += 1
            replaced = True
            break
        if replaced:
            continue
        for marker, replacement in replacements:
            if marker in text:
                set_paragraph_text_with_template_style(paragraph, replacement, paragraph_run_style(paragraph))
                changed += 1
                break
    return changed


def iter_all_paragraphs(document: Document) -> list[Paragraph]:
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    return paragraphs


def normalize_report_brand_names(document: Document) -> int:
    changed = 0
    for paragraph in iter_all_paragraphs(document):
        text = paragraph_text(paragraph)
        if not text:
            continue
        replacement = normalize_brand_text(text)
        if replacement != text:
            set_paragraph_text_with_template_style(paragraph, replacement, paragraph_run_style(paragraph))
            changed += 1
    return changed


def apply_run_style(run: Any, template_style: dict[str, Any]) -> None:
    run.font.name = template_style.get("font_name") or "仿宋_GB2312"
    run.font.size = template_style.get("font_size") or Pt(14)
    run.bold = bool(template_style.get("bold"))
    run.italic = bool(template_style.get("italic"))
    run.underline = template_style.get("underline")
    east_asia = template_style.get("east_asia") or run.font.name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def apply_usnews_brand_run_font(document: Document) -> int:
    changed = 0
    for paragraph in iter_all_paragraphs(document):
        text = paragraph_text(paragraph)
        normalized = normalize_brand_text(text)
        if "U.S. News" not in normalized:
            continue
        template_style = paragraph_run_style(paragraph)
        for run in list(paragraph.runs):
            paragraph._p.remove(run._element)
        for part in re.split(r"(U\.S\. News)", normalized):
            if not part:
                continue
            run = paragraph.add_run(part)
            apply_run_style(run, template_style)
            if part == "U.S. News":
                run.font.name = "Times New Roman"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        changed += 1
    return changed


def normalize_rank_prefix_markers(document: Document) -> int:
    changed = 0
    for paragraph in iter_all_paragraphs(document):
        text = paragraph_text(paragraph)
        if not text or ("#" not in text and "=" not in text):
            continue
        replacement = clean_rank_display(text)
        if replacement != text:
            set_paragraph_text_with_template_style(paragraph, replacement, paragraph_run_style(paragraph))
            changed += 1
    return changed


def normalize_inline_image_paragraph_alignment(document: Document) -> int:
    changed = 0
    for paragraph in iter_all_paragraphs(document):
        if not paragraph._element.xpath(".//w:drawing"):
            continue
        if paragraph.alignment != WD_ALIGN_PARAGRAPH.CENTER:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            changed += 1
    return changed


def normalize_section_page_geometry(document: Document) -> int:
    if not document.sections:
        return 0
    template = document.sections[0]
    target_values = {
        "page_width": template.page_width,
        "page_height": template.page_height,
        "left_margin": template.left_margin,
        "right_margin": template.right_margin,
        "top_margin": template.top_margin,
        "bottom_margin": template.bottom_margin,
        "header_distance": template.header_distance,
        "footer_distance": template.footer_distance,
    }
    changed = 0
    for section in document.sections[1:]:
        if section.orientation != WD_ORIENT.PORTRAIT:
            section.orientation = WD_ORIENT.PORTRAIT
            changed += 1
        for attr, expected in target_values.items():
            if getattr(section, attr) != expected:
                setattr(section, attr, expected)
                changed += 1
    return changed


def paragraph_has_image(paragraph: Paragraph) -> bool:
    return bool(paragraph._element.xpath(".//w:drawing"))


def paragraph_is_heading1(paragraph: Paragraph) -> bool:
    return getattr(getattr(paragraph, "style", None), "name", "") == "Heading 1"


def find_heading1_index(document: Document, marker: str) -> int | None:
    for index, paragraph in enumerate(document.paragraphs):
        text = clean_text(paragraph_text(paragraph))
        if paragraph_is_heading1(paragraph) and marker in text:
            return index
    return None


def next_heading1_index(document: Document, start_index: int) -> int:
    for index in range(start_index + 1, len(document.paragraphs)):
        if paragraph_is_heading1(document.paragraphs[index]):
            return index
    return len(document.paragraphs)


def remove_image_paragraphs_between(document: Document, start_index: int, end_index: int) -> int:
    changed = 0
    for paragraph in list(document.paragraphs[start_index + 1 : end_index]):
        if paragraph_has_image(paragraph):
            remove_paragraph(paragraph)
            changed += 1
    return changed


def find_caption_by_keywords(document: Document, required_keywords: list[str]) -> Paragraph | None:
    required = [normalize_brand_text(clean_text(item)) for item in required_keywords if clean_text(item)]
    for paragraph in document.paragraphs:
        text = normalize_brand_text(clean_text(paragraph_text(paragraph)))
        if not is_caption_paragraph(text):
            continue
        if all(keyword in text for keyword in required):
            return paragraph
    return None


def repair_peer_indicator_figure_anchors(document: Document) -> int:
    section_start = find_heading1_index(document, "对标高校分指标分析")
    if section_start is None:
        return 0
    section_end = next_heading1_index(document, section_start)
    changed = remove_image_paragraphs_between(document, section_start, section_end)
    repairs = [
        ("figure_09_arwu_peer_indicator_relative_rank", ["软科", "分项指标相对排名"]),
        ("figure_10_qs_peer_indicator_relative_rank", ["QS", "分项指标相对排名"]),
        ("figure_11_usnews_peer_indicator_relative_rank", ["U.S. News", "分项指标相对排名"]),
        ("figure_12_the_peer_indicator_relative_rank", ["THE", "分项指标相对排名"]),
    ]
    for artifact_id, keywords in repairs:
        image_path = artifact_png_path(artifact_id)
        caption = find_caption_by_keywords(document, keywords)
        if image_path is None or caption is None:
            continue
        insert_image_before_paragraph(document, caption, image_path)
        changed += 1
    return changed


def restore_empty_major_section_headings(document: Document) -> int:
    changed = 0
    replacements = [
        ("1、软科学科排名", "五、对标高校分指标分析"),
        ("1、论文整体", "六、对标高校论文指标分析"),
    ]
    for index, paragraph in enumerate(document.paragraphs):
        if not paragraph_is_heading1(paragraph) or clean_text(paragraph_text(paragraph)):
            continue
        next_heading_text = ""
        for probe in range(index + 1, len(document.paragraphs)):
            candidate = document.paragraphs[probe]
            candidate_text = clean_text(paragraph_text(candidate))
            if not candidate_text:
                continue
            if paragraph_is_heading1(candidate):
                break
            if paragraph_is_heading2(candidate):
                next_heading_text = candidate_text
                break
        for marker, replacement in replacements:
            if marker in next_heading_text:
                set_paragraph_text_with_template_style(paragraph, replacement, paragraph_run_style(paragraph))
                changed += 1
                break
    return changed


def paragraph_is_heading2(paragraph: Paragraph) -> bool:
    return getattr(getattr(paragraph, "style", None), "name", "") == "Heading 2"


def source_from_subsection_heading(text: str) -> str:
    normalized = normalize_brand_text(clean_text(text))
    compact = re.sub(r"\s+", "", normalized)
    if not compact:
        return ""
    if "U.S.News" in compact or "USNews" in compact:
        return "US News"
    if "QS" in compact:
        return "QS"
    if "THE" in compact:
        return "THE"
    if "软科" in compact:
        return "软科"
    return ""


def source_subsection_has_evidence(source: str, context: str) -> bool:
    if context == "peer_table":
        artifact_id = source_peer_table_artifact_id(source)
        return bool(artifact_id and read_artifact_rows(artifact_id, "table_matrix.csv"))
    if context == "peer_indicator":
        artifact_id = source_peer_indicator_artifact_id(source)
        return bool(artifact_id and read_artifact_rows(artifact_id, "chart_data.csv"))
    if context == "indicator_history":
        if source_history_rows(source):
            return True
        if not source_has_tongji_latest(source):
            return False
        artifact_ids = source_indicator_history_artifact_ids(source)
        if source == "US News":
            return any(read_artifact_rows(artifact_id, "table_matrix.csv") or read_artifact_rows(artifact_id, "chart_data.csv") for artifact_id in artifact_ids)
        return any(read_artifact_rows(artifact_id, "table_matrix.csv") for artifact_id in artifact_ids)
    return True


def remove_block_elements_until(paragraph: Paragraph, stop_paragraph: Paragraph | None) -> int:
    parent = paragraph._p.getparent()
    if parent is None:
        return 0
    current = paragraph._p
    stop_element = stop_paragraph._p if stop_paragraph is not None else None
    removed = 0
    while current is not None and current is not stop_element:
        next_current = current.getnext()
        parent.remove(current)
        removed += 1
        current = next_current
    return removed


def remove_block_elements_until_after(paragraph: Paragraph, stop_paragraph: Paragraph | None) -> int:
    parent = paragraph._p.getparent()
    if parent is None:
        return 0
    current = paragraph._p.getnext()
    stop_element = stop_paragraph._p if stop_paragraph is not None else None
    removed = 0
    while current is not None and current is not stop_element:
        next_current = current.getnext()
        parent.remove(current)
        removed += 1
        current = next_current
    return removed


def renumber_heading2_between(document: Document, start_index: int, end_index: int) -> int:
    changed = 0
    ordinal = 1
    for paragraph in document.paragraphs[start_index + 1 : end_index]:
        if not paragraph_is_heading2(paragraph):
            continue
        text = paragraph_text(paragraph)
        replacement = re.sub(r"^\s*\d+\s*[、.．]\s*", f"{ordinal}、", text, count=1)
        if replacement != text:
            set_paragraph_text_with_template_style(paragraph, replacement, paragraph_run_style(paragraph))
            changed += 1
        ordinal += 1
    return changed


def remove_empty_source_subsections(document: Document, section_marker: str, context: str) -> int:
    section_start = find_heading1_index(document, section_marker)
    if section_start is None:
        return 0
    changed = 0
    index = section_start + 1
    while True:
        paragraphs = document.paragraphs
        section_end = next_heading1_index(document, section_start)
        if index >= section_end:
            break
        paragraph = paragraphs[index]
        if not paragraph_is_heading2(paragraph):
            index += 1
            continue
        source = source_from_subsection_heading(paragraph_text(paragraph))
        if not source or source_subsection_has_evidence(source, context):
            index += 1
            continue
        next_index = section_end
        for probe in range(index + 1, section_end):
            if paragraph_is_heading2(paragraphs[probe]) or paragraph_is_heading1(paragraphs[probe]):
                next_index = probe
                break
        stop_paragraph = paragraphs[next_index] if next_index < len(paragraphs) else None
        changed += remove_block_elements_until_after(paragraph, stop_paragraph)
        insert_note_after_heading(paragraph, source_unavailable_text(source, read_target_subject_name(), context))
        changed += 1
        index += 2
    section_start = find_heading1_index(document, section_marker)
    if section_start is None:
        return changed
    section_end = next_heading1_index(document, section_start)
    changed += renumber_heading2_between(document, section_start, section_end)
    return changed


def normalize_evidence_backed_source_sections(document: Document) -> int:
    return (
        remove_empty_source_subsections(document, "排名历史发展分析", "indicator_history")
        + remove_empty_source_subsections(document, "对标高校整体分析", "peer_table")
        + remove_empty_source_subsections(
            document,
            "对标高校分指标分析",
            "peer_indicator",
        )
    )


def paper_artifacts_available() -> bool:
    artifact_ids = [
        "figure_13_paper_output_citation_relative_rank",
        "figure_14_paper_cnci_cpp_relative_rank",
        "figure_15_paper_highly_cited_relative_rank",
        "figure_16_paper_international_collab_relative_rank",
    ]
    return any((SUBJECT_RUN_DIR / "artifacts" / artifact_id / "chart_data.csv").exists() for artifact_id in artifact_ids)


def remove_appendix_section(document: Document) -> int:
    section_start = find_heading1_index(document, "附表")
    if section_start is None:
        return 0
    section_end = next_heading1_index(document, section_start)

    body = document._body._element
    heading_element = document.paragraphs[section_start]._element
    heading_pos = list(body).index(heading_element)

    children_to_remove = []
    for child in list(body)[heading_pos:]:
        if child.tag == qn("w:sectPr"):
            break
        children_to_remove.append(child)

    for child in children_to_remove:
        body.remove(child)

    return 1


def normalize_missing_paper_section(document: Document) -> int:
    if paper_artifacts_available():
        return 0
    section_start = find_heading1_index(document, "对标高校论文指标分析")
    if section_start is None:
        return 0
    section_end = next_heading1_index(document, section_start)
    for paragraph in list(document.paragraphs[section_start + 1 : section_end]):
        remove_paragraph(paragraph)
    note = (
        f"当前 Web of Science 37校论文数据中未匹配到{read_target_subject_name()}学科对应记录，"
        "本报告不使用学校整体论文数据或宽口径代理数据替代；待补充该一级学科论文数据后，再生成论文总量、CNCI、"
        "高被引论文和国际合作等对标图表。"
    )
    insert_note_after_heading(document.paragraphs[section_start], note)
    return 1


def proxy_label_replacements(text: str, subject_name: str) -> str:
    arwu_label = arwu_public_subject_label(subject_name)
    result = text
    if arwu_label and arwu_label != subject_name:
        replacements = {
            f"软科世界一流学科{subject_name}指标权重": f"软科世界一流学科{arwu_label}口径指标权重",
            f"软科{subject_name}指标权重": f"软科{arwu_label}口径指标权重",
            f"软科排名中内地高校位于前列的比较": f"软科{arwu_label}排名中内地高校位于前列的比较",
        }
        for old, new in replacements.items():
            result = result.replace(old, new)
    qs_label = (
        report_selected_subject_label("QS", "")
        or resolved_subject_value("qs_subject_name")
    )
    if qs_label and qs_label != subject_name:
        replacements = {
            f"QS学科排名{subject_name}指标权重": f"QS学科排名{qs_label}口径指标权重",
            f"QS{subject_name}指标权重": f"QS{qs_label}口径指标权重",
            f"QS排名前200位内地高校{subject_name}分项指标相对排名": f"QS排名前200位内地高校{qs_label}口径分项指标相对排名",
        }
        for old, new in replacements.items():
            result = result.replace(old, new)
    usnews_label = report_selected_subject_label("US News", "")
    if usnews_label and usnews_label != subject_name:
        replacements = {
            f"U.S. News 学科排名{subject_name}学科指标权重": f"U.S. News 学科排名{usnews_label}口径指标权重",
            f"U.S. News学科排名{subject_name}学科指标权重": f"U.S. News学科排名{usnews_label}口径指标权重",
            f"U.S. News{subject_name}指标权重": f"U.S. News{usnews_label}口径指标权重",
        }
        for old, new in replacements.items():
            result = result.replace(old, new)
    return result


def normalize_proxy_labels(document: Document) -> int:
    subject_name = read_target_subject_name()
    changed = 0
    for paragraph in document.paragraphs:
        text = paragraph_text(paragraph)
        replacement = proxy_label_replacements(text, subject_name)
        if replacement != text:
            set_paragraph_text_with_template_style(paragraph, replacement, paragraph_run_style(paragraph))
            changed += 1
    return changed


def remove_empty_paragraphs_before_media(document: Document) -> int:
    removed = 0
    indices = list(range(len(document.paragraphs) - 1))
    for i in reversed(indices):
        p = document.paragraphs[i]
        if p.text.strip():
            continue
        if p._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline'):
            continue
        if p._element.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing'):
            continue
        next_p = document.paragraphs[i + 1]
        next_text = next_p.text.strip()
        next_has_image = bool(
            next_p._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline')
            or next_p._element.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing')
        )
        next_is_caption = next_text and (next_text.startswith('图') or next_text.startswith('表')) and ('、' in next_text or '：' in next_text)
        if next_has_image or next_is_caption:
            remove_paragraph(p)
            removed += 1
    return removed


def normalize_cover_title(document: Document) -> int:
    subject_name = read_target_subject_name()
    title = f"同济大学{subject_name}学科的国际排名发展分析"
    changed = 0
    for paragraph in document.paragraphs[:12]:
        if paragraph_text(paragraph) != title:
            continue
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            if run.text:
                run.font.size = Pt(20)
                changed += 1
        break
    return changed


def remove_trailing_empty_paragraphs(document: Document) -> int:
    body = document._body._element
    children = list(body)
    removed = 0
    while len(children) >= 2:
        last_content = children[-2] if children[-1].tag == qn("w:sectPr") else children[-1]
        if last_content.tag != qn("w:p") or not is_empty_paragraph_element(last_content):
            break
        body.remove(last_content)
        removed += 1
        children = list(body)
    return removed


def remove_empty_paragraphs_before_section_breaks(document: Document) -> int:
    body = document._body._element
    removed = 0
    changed = True
    while changed:
        changed = False
        children = list(body)
        for index, child in enumerate(children):
            if child.tag != qn("w:p") or child.find(".//" + qn("w:sectPr")) is None:
                continue
            prev_index = index - 1
            while prev_index >= 0:
                prev = list(body)[prev_index]
                if prev.tag != qn("w:p") or not is_empty_paragraph_element(prev):
                    break
                body.remove(prev)
                removed += 1
                changed = True
                prev_index -= 1
            if changed:
                break
    return removed


def remove_trailing_empty_final_section(document: Document) -> int:
    body = document._body._element
    children = list(body)
    if len(children) < 2:
        return 0

    body_section = children[-1]
    final_paragraph = children[-2]
    if body_section.tag != qn("w:sectPr") or final_paragraph.tag != qn("w:p"):
        return 0
    if paragraph_element_text(final_paragraph):
        return 0
    has_drawing = (
        final_paragraph.find(".//" + qn("w:drawing")) is not None
        or final_paragraph.find(".//" + qn("w:pict")) is not None
    )
    if has_drawing:
        return 0
    paragraph_section = final_paragraph.find(".//" + qn("w:sectPr"))
    if paragraph_section is None:
        return 0

    body.remove(body_section)
    body.remove(final_paragraph)
    body.append(deepcopy(paragraph_section))
    return 1


def remove_trailing_table_empty_paragraph_in_docx_xml(docx_path: Path) -> int:
    """Remove the final empty paragraph that Word may push onto a blank page after a wide table."""
    temp_docx = docx_path.with_suffix(".tmp.docx")
    changed = 0
    w_ns = NS["w"]

    def local_name(element: Any) -> str:
        return etree.QName(element).localname

    def is_empty_nonsection_paragraph(element: Any) -> bool:
        if local_name(element) != "p":
            return False
        text = "".join(element.xpath(".//w:t/text()", namespaces=NS)).strip()
        if text:
            return False
        blocked_children = element.xpath(
            ".//w:sectPr | .//w:drawing | .//w:pict | .//w:object | .//w:br",
            namespaces=NS,
        )
        return not blocked_children

    with ZipFile(docx_path, "r") as zin, ZipFile(temp_docx, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                root = etree.fromstring(data)
                body = root.find(f"{{{w_ns}}}body")
                if body is not None:
                    children = list(body)
                    if len(children) >= 3 and local_name(children[-1]) == "sectPr":
                        maybe_empty = children[-2]
                        previous = children[-3]
                        if local_name(previous) == "tbl" and is_empty_nonsection_paragraph(maybe_empty):
                            body.remove(maybe_empty)
                            changed = 1
                            data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
    temp_docx.replace(docx_path)
    return changed


def replace_subject_placeholders_in_xml_parts(docx_path: Path, subject_name: str, the_subject_area: str = "") -> int:
    temp_docx = docx_path.with_suffix(".tmp.docx")
    changed_members = 0
    with ZipFile(docx_path, "r") as zin, ZipFile(temp_docx, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("word/") and item.filename.endswith(".xml"):
                text = data.decode("utf-8", errors="ignore")
                replaced_text = replace_subject_placeholder_text(text, subject_name, the_subject_area)
                replaced_text = proxy_label_replacements(replaced_text, subject_name)
                if replaced_text != text:
                    data = replaced_text.encode("utf-8")
                    changed_members += 1
            elif item.filename.startswith("word/media/") and item.filename.endswith(".svg"):
                text = data.decode("utf-8", errors="ignore")
                replaced_text = proxy_label_replacements(text, subject_name)
                if replaced_text != text:
                    data = replaced_text.encode("utf-8")
                    changed_members += 1
            zout.writestr(item, data)
    temp_docx.replace(docx_path)
    return changed_members


def scale_embedded_image(docx_path: Path, target_member: str, scale: float) -> int:
    with ZipFile(docx_path) as zip_file:
        document_xml = zip_file.read("word/document.xml").decode("utf-8")
        rels_xml = zip_file.read("word/_rels/document.xml.rels")

    rel_root = ET.fromstring(rels_xml)
    target_rel_ids = {
        rel.attrib.get("Id", "")
        for rel in rel_root
        if f"word/{rel.attrib.get('Target', '')}" == target_member
    }
    if not target_rel_ids:
        return 0

    changed = 0
    for rel_id in target_rel_ids:
        pattern = re.compile(r"(<w:drawing\b(?:(?!</w:drawing>).)*?r:embed=\"" + re.escape(rel_id) + r"\"(?:(?!</w:drawing>).)*?</w:drawing>)", re.DOTALL)

        def scale_drawing(match: re.Match[str]) -> str:
            nonlocal changed
            drawing_xml = match.group(1)

            def scale_extent(extent_match: re.Match[str]) -> str:
                nonlocal changed
                cx = str(max(1, int(int(extent_match.group(2)) * scale)))
                cy = str(max(1, int(int(extent_match.group(4)) * scale)))
                changed += 1
                return f'{extent_match.group(1)}{cx}{extent_match.group(3)}{cy}{extent_match.group(5)}'

            return re.sub(r'(<(?:wp:extent|a:ext)\b[^>]*?\bcx=")(\d+)(" [^>]*?\bcy=")(\d+)(")', scale_extent, drawing_xml)

        document_xml = pattern.sub(scale_drawing, document_xml, count=1)

    if changed:
        replace_docx_xml_member(docx_path, "word/document.xml", document_xml.encode("utf-8"))
    return changed


def embedded_image_width_ratio(row: dict[str, str]) -> float:
    return 1.0


def set_embedded_image_extent(docx_path: Path, target_member: str, width_emu: int, height_emu: int) -> int:
    with ZipFile(docx_path) as zip_file:
        document_xml = zip_file.read("word/document.xml").decode("utf-8")
        rels_xml = zip_file.read("word/_rels/document.xml.rels")

    rel_root = ET.fromstring(rels_xml)
    target_rel_ids = {
        rel.attrib.get("Id", "")
        for rel in rel_root
        if f"word/{rel.attrib.get('Target', '')}" == target_member
    }
    if not target_rel_ids:
        return 0

    changed = 0
    for rel_id in target_rel_ids:
        pattern = re.compile(
            r"(<w:drawing\b(?:(?!</w:drawing>).)*?r:embed=\"" + re.escape(rel_id) + r"\"(?:(?!</w:drawing>).)*?</w:drawing>)",
            re.DOTALL,
        )

        def resize_drawing(match: re.Match[str]) -> str:
            nonlocal changed
            drawing_xml = match.group(1)

            def replace_extent(extent_match: re.Match[str]) -> str:
                nonlocal changed
                changed += 1
                return f'{extent_match.group(1)}{width_emu}{extent_match.group(2)}{height_emu}{extent_match.group(3)}'

            return re.sub(
                r'(<(?:wp:extent|a:ext)\b[^>]*?\bcx=")\d+("[^>]*?\bcy=")\d+(")',
                replace_extent,
                drawing_xml,
            )

        document_xml = pattern.sub(resize_drawing, document_xml, count=1)

    if changed:
        replace_docx_xml_member(docx_path, "word/document.xml", document_xml.encode("utf-8"))
    return changed


def normalize_embedded_image_extents(rows: list[dict[str, str]], docx_path: Path, content_width: int) -> list[dict[str, Any]]:
    image_targets = extract_template_image_targets(docx_path)
    log_rows: list[dict[str, Any]] = []
    for row in rows:
        if row.get("slot_kind") != "figure" or not row.get("reference_media"):
            continue
        image_index = parse_slot_index(row.get("reference_object", ""), "image_")
        target_member = image_targets.get(image_index or -1, "")
        image_path = artifact_png_path(row.get("artifact_id", ""))
        if image_index is None or not target_member or image_path is None:
            log_rows.append(
                {
                    "slot": row.get("logical_slot_id", ""),
                    "artifact_id": row.get("artifact_id", ""),
                    "operation": "normalize_embedded_image_extent",
                    "status": "skipped",
                    "reason": "missing_image_index_or_png_or_target",
                }
            )
            continue
        with Image.open(image_path) as image:
            width_px, height_px = image.size
        width_emu = int(content_width * embedded_image_width_ratio(row))
        height_emu = int(width_emu * height_px / max(1, width_px))
        changes = set_embedded_image_extent(docx_path, target_member, width_emu, height_emu)
        log_rows.append(
            {
                "slot": row.get("logical_slot_id", ""),
                "artifact_id": row.get("artifact_id", ""),
                "operation": "normalize_embedded_image_extent",
                "status": "ok" if changes else "skipped",
                "target_member": target_member,
                "source_png": str(image_path),
                "width_emu": width_emu,
                "height_emu": height_emu,
                "changed_extents": changes,
            }
        )
    return log_rows


def is_word_heading1(paragraph: Paragraph) -> bool:
    style_name = getattr(getattr(paragraph, "style", None), "name", "") or ""
    return style_name == "Heading 1"


def compact_final_recommendation_section(document: Document) -> int:
    heading_tokens = [
        "\u53d1\u5c55\u6f5c\u529b",
        "\u5efa\u8bae",
    ]
    fallback_heading_token = "\u53d1\u5c55\u5efa\u8bae"
    start_index: int | None = None
    for index, paragraph in enumerate(document.paragraphs):
        text = re.sub(r"\s+", "", paragraph_text(paragraph))
        if all(token in text for token in heading_tokens) or fallback_heading_token in text:
            start_index = index
            break
    if start_index is None:
        return 0

    changed = 0
    section_end = next_heading1_index(document, start_index)
    body_paragraphs = [
        paragraph
        for paragraph in document.paragraphs[start_index + 1 : section_end]
        if clean_text(paragraph_text(paragraph)) and not paragraph_is_heading1(paragraph)
    ]
    if body_paragraphs:
        compact_text = (
            f"综合当前可核验图表，同济大学{read_target_subject_name()}学科应把已有证据与缺失榜单口径分开判断；"
            "后续重点提升标准化引文影响、高被引论文占比和国际合作论文质量，待补齐直接口径后再开展跨榜单归因。"
        )
        set_paragraph_text_with_template_style(body_paragraphs[0], compact_text, paragraph_run_style(body_paragraphs[0]))
        changed += 1
        for paragraph in body_paragraphs[1:]:
            remove_paragraph(paragraph)
            changed += 1
    section_end = next_heading1_index(document, start_index)
    for paragraph in list(document.paragraphs[start_index + 1 : section_end]):
        if not clean_text(paragraph_text(paragraph)) and not paragraph_has_image(paragraph):
            remove_paragraph(paragraph)
            changed += 1
    for paragraph in document.paragraphs[start_index:]:
        text = re.sub(r"\s+", "", paragraph_text(paragraph))
        if changed > 0 and is_word_heading1(paragraph) and not any(token in text for token in heading_tokens):
            break
        paragraph_format = paragraph.paragraph_format
        if is_word_heading1(paragraph):
            paragraph_format.space_before = Pt(0)
            paragraph_format.space_after = Pt(3)
        else:
            paragraph_format.space_before = Pt(0)
            paragraph_format.space_after = Pt(0)
            paragraph_format.line_spacing = 1
            try:
                paragraph_format.widow_control = False
            except AttributeError:
                pass
        changed += 1
    return changed


def renumber_figure_captions(document: Document) -> int:
    changed = 0
    figure_no = 1
    caption_pattern = re.compile(r"^图\s*\d+\s*、")
    for paragraph in document.paragraphs:
        text = paragraph_text(paragraph).strip()
        if not text or not is_caption_paragraph(text):
            continue
        if not caption_pattern.match(text):
            continue
        expected_prefix = f"图{figure_no}、"
        new_text = caption_pattern.sub(expected_prefix, text, count=1)
        if new_text != text:
            set_paragraph_text_with_template_style(paragraph, new_text, paragraph_run_style(paragraph))
            changed += 1
        figure_no += 1
    return changed


def renumber_table_captions(document: Document) -> int:
    changed = 0
    table_no = 1
    caption_pattern = re.compile(r"^表\s*(?:\d+\s*)?[、：:]")
    for paragraph in document.paragraphs:
        text = paragraph_text(paragraph).strip()
        if not text or text.startswith("附表") or not is_caption_paragraph(text):
            continue
        if not caption_pattern.match(text):
            continue
        expected_prefix = f"表{table_no}、"
        new_text = caption_pattern.sub(expected_prefix, text, count=1)
        if new_text != text:
            set_paragraph_text_with_template_style(paragraph, new_text, paragraph_run_style(paragraph))
            changed += 1
        table_no += 1
    return changed


def renumber_captions_in_docx_xml(docx_path: Path) -> dict[str, int]:
    figure_changes = 0
    table_changes = 0
    with ZipFile(docx_path, "r") as zip_file:
        document_xml = zip_file.read("word/document.xml")

    root = ET.fromstring(document_xml)
    figure_pattern = re.compile(r"^图\s*\d+\s*、")
    table_pattern = re.compile(r"^表\s*(?:\d+\s*)?[、：:]")
    figure_no = 1
    table_no = 1
    for paragraph in root.findall(".//w:p", NS):
        text_nodes = paragraph.findall(".//w:t", NS)
        if not text_nodes:
            continue
        text = "".join(node.text or "" for node in text_nodes).strip()
        new_text = text
        if figure_pattern.match(text):
            new_text = figure_pattern.sub(f"图{figure_no}、", text, count=1)
            figure_no += 1
            if new_text != text:
                figure_changes += 1
        elif text.startswith("表") and not text.startswith("附表") and table_pattern.match(text):
            new_text = table_pattern.sub(f"表{table_no}、", text, count=1)
            table_no += 1
            if new_text != text:
                table_changes += 1
        if new_text != text:
            text_nodes[0].text = new_text
            for node in text_nodes[1:]:
                node.text = ""

    if figure_changes or table_changes:
        xml_bytes = ET.tostring(root, encoding="utf-8", xml_declaration=True)
        replace_docx_xml_member(docx_path, "word/document.xml", xml_bytes)
    return {
        "figure_caption_changes": figure_changes,
        "table_caption_changes": table_changes,
    }


def postprocess_final_document(rows: list[dict[str, str]], excluded_rows: list[dict[str, str]] | None = None) -> list[dict[str, Any]]:
    subject_name = read_target_subject_name()
    the_subject_area = read_target_subject_area()
    excluded_rows = excluded_rows or []
    document = Document(str(OUTPUT_DOCX))
    placeholder_changes = replace_subject_placeholders(document, subject_name, the_subject_area)
    excluded_slot_cleanups = remove_excluded_slot_content(document, excluded_rows)
    excluded_section_note_changes = normalize_excluded_section_notes(document, excluded_rows)
    residual_paragraph_changes = replace_residual_template_paragraphs(document)
    stale_template_paragraph_removals = remove_stale_template_paragraphs(document)
    qs_history_table_insertions = insert_qs_history_table(document)
    the_history_table_insertions = insert_the_history_table(document)
    duplicate_the_history_summary_removals = remove_duplicate_the_history_summary_after_table(document)
    restored_major_headings_before_qs_insert = restore_empty_major_section_headings(document)
    qs_peer_indicator_insertions = insert_qs_peer_indicator_figure(document)
    brand_name_changes = normalize_report_brand_names(document)
    rank_prefix_changes = normalize_rank_prefix_markers(document)
    proxy_label_changes = normalize_proxy_labels(document)
    cover_title_changes = normalize_cover_title(document)
    virtual_table_insertions = insert_virtual_native_tables(document, rows)
    evidence_section_changes = normalize_evidence_backed_source_sections(document)
    restored_major_headings = restore_empty_major_section_headings(document)
    peer_figure_anchor_repairs = repair_peer_indicator_figure_anchors(document)
    table_caption_pair_repairs = keep_table_captions_with_tables(document)
    appendix_removals = remove_appendix_section(document)
    missing_paper_section_changes = normalize_missing_paper_section(document)
    brand_name_changes += normalize_report_brand_names(document)
    rank_prefix_changes += normalize_rank_prefix_markers(document)
    empty_before_media_removals = remove_empty_paragraphs_before_media(document)
    image_alignment_changes = normalize_inline_image_paragraph_alignment(document)
    figure_caption_renumbers = renumber_figure_captions(document)
    table_caption_renumbers = renumber_table_captions(document)
    usnews_brand_run_changes = apply_usnews_brand_run_font(document)
    removed_heading_page_breaks = remove_empty_page_breaks_before_heading1(document)
    inserted_toc_page_breaks = ensure_page_break_before_first_body_heading(document)
    cover_main_spacing_changes = compact_cover_main_spacing(document)
    cover_spacing_changes = compact_cover_section_break(document)
    trailing_empty_changes = remove_trailing_empty_paragraphs(document)
    pre_section_empty_changes = remove_empty_paragraphs_before_section_breaks(document)
    trailing_final_section_changes = remove_trailing_empty_final_section(document)
    table_geometry_changes = normalize_existing_table_geometry(document)
    final_recommendation_compactions = compact_final_recommendation_section(document)
    section_page_geometry_changes = normalize_section_page_geometry(document)
    doc_content_width = content_width_emu(document)
    document.save(str(OUTPUT_DOCX))

    section_changes = normalize_odd_page_sections(OUTPUT_DOCX)
    xml_placeholder_changes = replace_subject_placeholders_in_xml_parts(OUTPUT_DOCX, subject_name, the_subject_area)
    image_extent_changes = normalize_embedded_image_extents(rows, OUTPUT_DOCX, doc_content_width)
    namespace_sanitize_result = sanitize_docx_ooxml_namespaces.sanitize_docx(OUTPUT_DOCX)
    trailing_table_empty_xml_changes = remove_trailing_table_empty_paragraph_in_docx_xml(OUTPUT_DOCX)

    return [
        {
            "slot": "document_subject_placeholders",
            "artifact_id": "",
            "operation": "replace_subject_placeholders",
            "status": "ok",
            "subject_name": subject_name,
            "the_subject_area": the_subject_area,
            "changed_paragraphs": placeholder_changes,
        },
        {
            "slot": "document_sections",
            "artifact_id": "",
            "operation": "normalize_odd_page_sections",
            "status": "ok",
            "changed_sections": section_changes,
        },
        {
            "slot": "section_page_geometry",
            "artifact_id": "",
            "operation": "normalize_section_page_geometry",
            "status": "ok",
            "changed_values": section_page_geometry_changes,
        },
        {
            "slot": "the_history_summary",
            "artifact_id": "figure_07_all_rankings_history",
            "operation": "remove_duplicate_the_history_summary_after_table",
            "status": "ok",
            "removed_paragraphs": duplicate_the_history_summary_removals,
        },
        {
            "slot": "cover_section_break",
            "artifact_id": "",
            "operation": "compact_cover_section_break",
            "status": "ok",
            "removed_blank_paragraphs": cover_spacing_changes,
        },
        {
            "slot": "cover_main_spacing",
            "artifact_id": "",
            "operation": "compact_cover_main_spacing",
            "status": "ok",
            "removed_blank_paragraphs": cover_main_spacing_changes,
        },
        {
            "slot": "table_geometry",
            "artifact_id": "",
            "operation": "normalize_existing_table_geometry",
            "status": "ok",
            "changed_tables": table_geometry_changes,
        },
        {
            "slot": "heading_page_breaks",
            "artifact_id": "",
            "operation": "remove_empty_page_breaks_before_heading1",
            "status": "ok",
            "removed_page_breaks": removed_heading_page_breaks,
        },
        {
            "slot": "toc_page_break",
            "artifact_id": "",
            "operation": "ensure_page_break_before_first_body_heading",
            "status": "ok",
            "inserted_page_breaks": inserted_toc_page_breaks,
        },
        {
            "slot": "document_xml_subject_placeholders",
            "artifact_id": "",
            "operation": "replace_subject_placeholders_in_xml_parts",
            "status": "ok",
            "changed_members": xml_placeholder_changes,
        },
        {
            "slot": "excluded_slot_content",
            "artifact_id": "",
            "operation": "remove_excluded_slot_content",
            "status": "ok",
            "changed_slots": sum(1 for item in excluded_slot_cleanups if item.get("status") == "ok"),
        },
        {
            "slot": "excluded_section_notes",
            "artifact_id": "",
            "operation": "normalize_excluded_section_notes",
            "status": "ok",
            "changed_paragraphs": excluded_section_note_changes,
        },
        {
            "slot": "residual_template_paragraphs",
            "artifact_id": "",
            "operation": "replace_residual_template_paragraphs",
            "status": "ok",
            "changed_paragraphs": residual_paragraph_changes,
        },
        {
            "slot": "stale_template_paragraphs",
            "artifact_id": "",
            "operation": "remove_stale_template_paragraphs",
            "status": "ok",
            "removed_paragraphs": stale_template_paragraph_removals,
        },
        *qs_history_table_insertions,
        *qs_peer_indicator_insertions,
        *virtual_table_insertions,
        {
            "slot": "major_section_headings_before_qs_insert",
            "artifact_id": "",
            "operation": "restore_empty_major_section_headings_before_qs_insert",
            "status": "ok",
            "changed_paragraphs": restored_major_headings_before_qs_insert,
        },
        {
            "slot": "brand_names",
            "artifact_id": "",
            "operation": "normalize_report_brand_names",
            "status": "ok",
            "changed_paragraphs": brand_name_changes,
        },
        {
            "slot": "usnews_brand_runs",
            "artifact_id": "",
            "operation": "apply_usnews_brand_run_font",
            "status": "ok",
            "changed_paragraphs": usnews_brand_run_changes,
        },
        {
            "slot": "rank_prefix_markers",
            "artifact_id": "",
            "operation": "normalize_rank_prefix_markers",
            "status": "ok",
            "changed_paragraphs": rank_prefix_changes,
        },
        {
            "slot": "empty_paragraphs_before_media",
            "artifact_id": "",
            "operation": "remove_empty_paragraphs_before_media",
            "status": "ok",
            "removed_paragraphs": empty_before_media_removals,
        },
        {
            "slot": "embedded_image_alignment",
            "artifact_id": "",
            "operation": "normalize_inline_image_paragraph_alignment",
            "status": "ok",
            "changed_paragraphs": image_alignment_changes,
        },
        {
            "slot": "peer_indicator_figure_anchors",
            "artifact_id": "",
            "operation": "repair_peer_indicator_figure_anchors",
            "status": "ok",
            "changed_paragraphs": peer_figure_anchor_repairs,
        },
        {
            "slot": "major_section_headings",
            "artifact_id": "",
            "operation": "restore_empty_major_section_headings",
            "status": "ok",
            "changed_paragraphs": restored_major_headings,
        },
        {
            "slot": "evidence_backed_source_sections",
            "artifact_id": "",
            "operation": "normalize_evidence_backed_source_sections",
            "status": "ok",
            "changed_paragraphs": evidence_section_changes,
        },
        *the_history_table_insertions,
        {
            "slot": "appendix_section",
            "artifact_id": "",
            "operation": "remove_appendix_section",
            "status": "ok",
            "removed_sections": appendix_removals,
        },
        {
            "slot": "missing_paper_section",
            "artifact_id": "",
            "operation": "normalize_missing_paper_section",
            "status": "ok",
            "changed_sections": missing_paper_section_changes,
        },
        {
            "slot": "proxy_labels",
            "artifact_id": "",
            "operation": "normalize_proxy_labels",
            "status": "ok",
            "changed_paragraphs": proxy_label_changes,
        },
        {
            "slot": "figure_captions",
            "artifact_id": "",
            "operation": "renumber_figure_captions",
            "status": "ok",
            "changed_paragraphs": figure_caption_renumbers,
        },
        {
            "slot": "table_captions",
            "artifact_id": "",
            "operation": "renumber_table_captions",
            "status": "ok",
            "changed_paragraphs": table_caption_renumbers,
        },
        {
            "slot": "table_caption_pairs",
            "artifact_id": "",
            "operation": "keep_table_captions_with_tables",
            "status": "ok",
            "changed_tables": table_caption_pair_repairs,
        },
        {
            "slot": "cover_title",
            "artifact_id": "",
            "operation": "normalize_cover_title",
            "status": "ok",
            "changed_runs": cover_title_changes,
        },
        {
            "slot": "trailing_empty_paragraphs",
            "artifact_id": "",
            "operation": "remove_trailing_empty_paragraphs",
            "status": "ok",
            "removed_paragraphs": trailing_empty_changes,
        },
        {
            "slot": "pre_section_empty_paragraphs",
            "artifact_id": "",
            "operation": "remove_empty_paragraphs_before_section_breaks",
            "status": "ok",
            "removed_paragraphs": pre_section_empty_changes,
        },
        {
            "slot": "trailing_empty_final_section",
            "artifact_id": "",
            "operation": "remove_trailing_empty_final_section",
            "status": "ok",
            "removed_sections": trailing_final_section_changes,
        },
        {
            "slot": "trailing_table_empty_paragraph_xml",
            "artifact_id": "",
            "operation": "remove_trailing_table_empty_paragraph_in_docx_xml",
            "status": "ok",
            "removed_paragraphs": trailing_table_empty_xml_changes,
        },
        {
            "slot": "final_recommendation_section",
            "artifact_id": "",
            "operation": "compact_final_recommendation_section",
            "status": "ok",
            "changed_paragraphs": final_recommendation_compactions,
        },
        {
            "slot": "ooxml_namespace_sanitize",
            "artifact_id": "",
            "operation": "sanitize_docx_ooxml_namespaces",
            "status": namespace_sanitize_result.get("status", "unknown"),
            "changed_members": len(namespace_sanitize_result.get("changed_members", [])),
        },
    ] + excluded_slot_cleanups + image_extent_changes


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    all_rows = read_all_slot_mapping()
    rows, excluded_rows = split_slot_rows_by_manifest(all_rows)

    working_docx = OUTPUT_DOCX.with_name(OUTPUT_DOCX.stem + "_tables_first.docx")
    shutil.copy2(TEMPLATE_DOCX, working_docx)

    log_rows: list[dict[str, Any]] = []
    log_rows.extend(replace_tables_and_caption_only_figures(rows, working_docx, OUTPUT_DOCX))
    working_docx.unlink(missing_ok=True)
    log_rows.extend(replace_embedded_images(rows, TEMPLATE_DOCX, OUTPUT_DOCX))
    log_rows.extend(postprocess_final_document(rows, excluded_rows))
    log_rows.extend(refresh_visible_figure_media(rows, OUTPUT_DOCX))
    audit_rows = audit_replacements(rows, OUTPUT_DOCX)
    write_audit_files(audit_rows)

    payload = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "template_docx": str(TEMPLATE_DOCX),
        "subject_run_dir": str(SUBJECT_RUN_DIR),
        "output_docx": str(OUTPUT_DOCX),
        "operations": log_rows,
        "summary": {
            "ok": sum(1 for item in log_rows if item["status"] == "ok"),
            "skipped": sum(1 for item in log_rows if item["status"] == "skipped"),
            "audit_ok": sum(1 for item in audit_rows if item["status"] == "ok"),
            "audit_mismatch": sum(1 for item in audit_rows if item["status"] != "ok"),
        },
    }
    REPLACEMENT_LOG_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8", newline="\n")
    print(json.dumps(payload["summary"], ensure_ascii=False))
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
