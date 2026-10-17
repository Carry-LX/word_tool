"""Small, auditable stage previews made from actual inputs and completed files."""
from __future__ import annotations

import csv
import hashlib
import sqlite3
import sys
from pathlib import Path

from docx import Document

from demo_runtime import now, read_json


def image_item(path: Path, label: str) -> dict:
    return {"id": hashlib.sha256(str(path.resolve()).encode("utf-8")).hexdigest()[:18],
            "path": str(path.resolve()), "label": label}


class StageResults:
    def __init__(self, state: dict):
        self.state = state
        self.run_dir: Path | None = None
        self.completed: list[Path] = []
        self.labels: dict[str, str] = {}
        self.latest_match: list[dict] = []

    def make(self, stage: str, status: str, metrics: dict) -> dict:
        if metrics.get("run_dir"):
            self.run_dir = Path(metrics["run_dir"])
        base = {"title": "阶段结果", "summary": "", "stats": [], "columns": [], "rows": [], "images": [], "updated_at": now()}
        if stage == "data":
            agent = sys.modules.get("run_subject_agent")
            if not agent:
                return base
            entries = [("排名数据库", "SQLITE_PATH"), ("学科目录", "DIM_SUBJECT_CSV"), ("学科别名", "DIM_ALIAS_CSV"),
                       ("指标权重", "FACT_WEIGHT_CSV"), ("排名指标", "FACT_INDICATOR_CSV"),
                       ("排名学科关联", "FACT_RANKING_SUBJECT_BRIDGE_CSV"), ("指标学科关联", "FACT_INDICATOR_SUBJECT_BRIDGE_CSV")]
            base.update(title="本次输入检查", summary="检查处理后的数据库和关联表。", columns=["输入", "文件状态"],
                        rows=[[name, "已就绪" if getattr(agent, attribute).is_file() else "缺失"] for name, attribute in entries])
            if status == "success":
                with sqlite3.connect(agent.SQLITE_PATH) as connection:
                    for table, label in [("fact_ranking", "全库排名"), ("fact_ranking_indicator", "全库指标"), ("dim_subject", "学科")]:
                        count = connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
                        base["stats"].append({"label": label, "value": f"{count:,}"})
        elif stage in {"resolve", "sources"}:
            if metrics.get("matched_sources"):
                self.latest_match = metrics["matched_sources"]
            base.update(title="匹配到的学科记录", summary="记录覆盖匹配口径下的各校数据，后续再选择本校与对标样本。",
                        columns=["来源", "匹配记录数"], rows=[[row["source"], f"{row['records']:,}"] for row in self.latest_match])
            if stage == "resolve":
                base["stats"] = [{"label": "目标学科", "value": self.state["subject"]["name"]},
                                 {"label": "学科代码", "value": self.state["subject"]["code"]}]
            else:
                base["title"] = "本次采用的榜单口径"
                rows = metrics.get("source_rows", [])
                if rows:
                    base.update(summary="保留实际映射与本校数据缺口。", columns=["来源", "实际口径", f"{self.state['university']}最新位次"],
                                rows=[[row["ranking_source"], row["selected_ranking_subject_or_area"], str(row.get("tongji_latest_rank") or "无本校记录")] for row in rows],
                                stats=[{"label": "可用榜单", "value": metrics["available_sources"]}, {"label": "可生成素材", "value": metrics["ready_artifacts"]}])
        elif stage == "artifacts":
            directory = Path(metrics["artifact_output_dir"]) if metrics.get("artifact_output_dir") else None
            if directory and metrics.get("artifact_status") == "generated" and directory not in self.completed:
                self.completed.append(directory)
            if self.run_dir and not self.labels:
                path = self.run_dir / "artifact_readiness.csv"
                if path.is_file():
                    with path.open(encoding="utf-8-sig", newline="") as handle:
                        self.labels = {row["artifact_id"]: row["label"] for row in csv.DictReader(handle)}
            images = []
            for output in self.completed:
                path = next((output / name for name in ["chart.png", "table.png"] if (output / name).is_file()), None)
                if path:
                    images.append(image_item(path, self.labels.get(output.name, "分析素材")))
            base.update(title=self.labels.get(metrics.get("artifact_id"), "图表与分析素材"), summary="每完成一项素材，就展示其真实输出。",
                        images=images[-8:], stats=[{"label": "已完成素材", "value": len(self.completed)}, {"label": "已生成预览", "value": len(images)}])
            if directory:
                path = next((directory / name for name in ["table_matrix.csv", "table_data.csv", "chart_data.csv"] if (directory / name).is_file()), None)
                if path:
                    with path.open(encoding="utf-8-sig", newline="") as handle:
                        reader = csv.reader(handle)
                        base["columns"] = next(reader, [])[:4]
                        base["rows"] = [row[:4] for _, row in zip(range(3), reader)]
        elif stage in {"word", "qa"}:
            path = Path(metrics["docx_path"]) if metrics.get("docx_path") else None
            base.update(title="正在组装报告" if stage == "word" else "Word 页面与版式检查",
                        summary="使用已生成素材，组装正文、ESI 趋势与三条精简建议。")
            if path and path.is_file():
                document = Document(path)
                headings = [paragraph.text for paragraph in document.paragraphs if paragraph.style.name == "Heading 1" and paragraph.text.strip()]
                base.update(title="Word 报告已组装", summary=path.name, columns=["报告章节"], rows=[[text] for text in headings],
                            stats=[{"label": "图片", "value": len(document.inline_shapes)}, {"label": "表格", "value": len(document.tables)}, {"label": "章节", "value": len(headings)}])
            report = self.state.get("report")
            if stage == "qa" and report:
                base.update(title="页面检查通过", summary="首页、分析页和末页建议均可预览。", columns=[], rows=[],
                            stats=[{"label": "预览页数", "value": report["page_count"]}, {"label": "空白页", "value": 0}, {"label": "稀疏页", "value": 0}],
                            images=[image_item(Path(report["pages"][index]), label) for index, label in [(0, "报告首页"), (min(4, len(report["pages"]) - 1), "分析页面"), (-1, "ESI 与精简建议")]])
        return base
