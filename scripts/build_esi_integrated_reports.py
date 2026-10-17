"""Revise selected published reports and retain all upstream ranking evidence.

Run directly. The report selection and ESI bindings are explicit runtime policy.
"""
from __future__ import annotations

import csv
import hashlib
import json
import os
import shutil
import subprocess
import sys
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw, ImageOps

import audit_word_report
import report_editorial as editorial
import sanitize_docx_ooxml_namespaces

PROJECT_DIR = Path(__file__).resolve().parent.parent
SOURCE_WORD_ROOT = PROJECT_DIR / "outputs/full_word_reports"
SOURCE_RUN_ROOT = PROJECT_DIR / "outputs/subject_agent_runs"
BUILD_REPORTS = True
RUN_RENDER_QA = True
RUN_WORD_COM = True
RENDER_WORKERS = 2
ONLY_REPORT_IDS: list[str] = []


def write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def latest_run(code: str) -> Path:
    return sorted(path for path in SOURCE_RUN_ROOT.glob(f"*_{code}_*")
                  if (path / "agent_run_summary.json").exists())[-1]


def source_word(code: str) -> Path:
    return next(sorted(SOURCE_WORD_ROOT.glob(f"*_{code}_*"))[-1].glob("*.docx"))


def geometry(document: Document) -> list[dict]:
    return [{key: int(getattr(section, key)) for key in ["page_width", "page_height", "top_margin",
             "bottom_margin", "left_margin", "right_margin"]} for section in document.sections]


def preserved_parts(source: Path, revised: Path) -> dict:
    allowed_changes = {"word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"}
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(revised) as after:
        before_names, after_names = set(before.namelist()), set(after.namelist())
        missing = before_names - after_names
        changed = [name for name in before_names & after_names if before.read(name) != after.read(name)]
        unexpected = sorted(name for name in changed if name not in allowed_changes)
        return {"missing_parts": sorted(missing), "changed_parts": sorted(changed),
                "unexpected_changes": unexpected, "added_parts": sorted(after_names - before_names),
                "status": "pass" if not missing and not unexpected else "fail"}


def restore_untouched_parts(source: Path, revised: Path) -> None:
    """Keep original styles, chrome and settings byte-for-byte after object-model edits."""
    editable = {"word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"}
    temporary = revised.with_suffix(".preserved.docx")
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(revised) as modified:
        original_names = set(original.namelist())
        modified_names = set(modified.namelist())
        with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as target:
            for name in sorted(original_names | modified_names):
                if name in original_names and name not in editable:
                    target.writestr(name, original.read(name))
                else:
                    target.writestr(name, modified.read(name))
    temporary.replace(revised)


def build_full_report(code: str, policy: dict, output: Path) -> dict:
    source, run_dir = source_word(code), latest_run(code)
    source_hash = sha256(source)
    document = Document(source)
    original_geometry = geometry(document)
    before_shapes = len(document.inline_shapes)
    before_tables = len(document.tables)
    binding = policy["bindings"][code]
    report_dir = output / code
    report_dir.mkdir(parents=True, exist_ok=True)
    out = report_dir / f"同济大学{binding['report_subject']}_学科国际排名发展分析_ESI修订版.docx"
    revision = editorial.replace_conclusions(document, run_dir)
    esi_result = editorial.insert_esi_before_conclusions(document, run_dir)
    document.save(out)
    restore_untouched_parts(source, out)
    assert geometry(Document(out)) == original_geometry
    assert len(Document(out).inline_shapes) == before_shapes + 1
    assert len(Document(out).tables) == before_tables
    package_audit = preserved_parts(source, out)
    assert package_audit["status"] == "pass", package_audit
    assert sha256(source) == source_hash
    artifact_dir = Path(esi_result["artifact_dir"])
    shutil.copytree(artifact_dir, report_dir / "esi_source_bundle", dirs_exist_ok=True)
    evidence = editorial.recommendation_evidence(run_dir)
    write_json(report_dir / "recommendation_evidence.json", evidence)
    write_json(report_dir / "package_preservation_audit.json", package_audit)
    result = {"report_id": code, "report_type": "full_subject_report", "subject": binding["report_subject"],
              "source_docx": str(source), "source_sha256": source_hash, "source_run_dir": str(run_dir),
              "docx_path": str(out), "esi_section": esi_result, "recommendation_revision": revision,
              "original_image_count": before_shapes, "new_image_count": before_shapes + 1,
              "geometry_preserved": True, "existing_tables_preserved": True,
              "package_preservation": package_audit["status"]}
    write_json(report_dir / "revision_manifest.json", result)
    return result


def build_topic_report(key: str, policy: dict, output: Path) -> dict:
    data = editorial.read_json(PROJECT_DIR / policy["esi_data_path"])
    esi = next(subject for subject in data["subjects"] if subject["subject_key"] == key)
    directory = output / key
    directory.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
    section.left_margin, section.right_margin = Cm(2.2), Cm(2.2)
    title_style = document.styles["Title"]
    for border in title_style.element.xpath("./w:pPr/w:pBdr"):
        border.getparent().remove(border)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    editorial.set_text(title, f"同济大学ESI{esi['subject_name']}发展分析", size=16)
    intro = document.add_paragraph()
    editorial.set_text(intro, f"本报告采用ESI“{esi['subject_name']}”学科口径，观察期为2024年11月至2026年9月。", size=10.5)
    intro.paragraph_format.keep_with_next = True
    heading = document.add_paragraph(style="Heading 1")
    editorial.set_text(heading, "一、前千分之一发展趋势", size=14)
    narrative = document.add_paragraph()
    editorial.set_text(narrative, policy["esi_narratives"][key], highlight=True)
    narrative.paragraph_format.line_spacing = 1.5
    narrative.paragraph_format.first_line_indent = Cm(0.74)
    narrative.paragraph_format.keep_with_next = True
    narrative.paragraph_format.space_after = Pt(6)
    artifact = PROJECT_DIR / policy["esi_artifacts_dir"] / key
    image = document.add_paragraph()
    image.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image.paragraph_format.line_spacing = 1
    image.paragraph_format.keep_with_next = True
    image.add_run().add_picture(str(artifact / "chart.png"), width=Cm(14.8))
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    editorial.set_text(caption, f"图1、同济大学ESI{esi['subject_name']}排名千分位变化趋势", size=10.5)
    caption.paragraph_format.space_after = Pt(8)
    heading = document.add_paragraph(style="Heading 1")
    editorial.set_text(heading, "二、发展潜力及建议", size=14)
    texts = policy["topic_recommendations"][key]
    for text in texts:
        paragraph = document.add_paragraph()
        editorial.set_text(paragraph, text)
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.paragraph_format.space_after = Pt(4)
    document.core_properties.title = f"同济大学ESI{esi['subject_name']}发展分析"
    out = directory / f"同济大学ESI{esi['subject_name']}_发展分析专题报告.docx"
    document.save(out)
    shutil.copytree(artifact, directory / "esi_source_bundle", dirs_exist_ok=True)
    result = {"report_id": key, "report_type": "esi_topic_report", "subject": esi["subject_name"],
              "docx_path": str(out), "recommendation_revision": {"new_characters": sum(map(len, texts)),
              "point_count": len(texts), "paragraphs": texts}, "esi_section": {"esi_key": key,
              "narrative": policy["esi_narratives"][key]}, "scope": "ESI学科专题，不借用其他一级学科的排名图表"}
    write_json(directory / "revision_manifest.json", result)
    return result


def word_worker(report: dict) -> dict:
    path = Path(report["docx_path"])
    environment = dict(os.environ)
    environment["PYTHONIOENCODING"] = "utf-8"
    environment["UPDATE_WORD_FIELDS_DOCX_PATH"] = str(path)
    updated = subprocess.run([sys.executable, "-B", str(PROJECT_DIR / "scripts/update_word_fields.py")],
                             env=environment, capture_output=True, timeout=70)
    try:
        fields = json.loads(updated.stdout.decode("utf-8"))
    except (ValueError, UnicodeError):
        fields = {"status": "failed", "stderr": updated.stderr.decode("utf-8", errors="replace")[-2000:]}
    write_json(path.parent / "qa/word_field_update.json", fields)
    environment["WORD_OPEN_DOCX_PATH"] = str(path)
    opened = subprocess.run([sys.executable, "-B", str(PROJECT_DIR / "scripts/word_openability_worker.py")],
                            env=environment, capture_output=True, timeout=45)
    result = json.loads(opened.stdout.decode("utf-8"))
    write_json(path.parent / "qa/word_openability.json", result)
    if fields.get("status") != "ok" or result.get("word_com_status") != "opened":
        raise RuntimeError(f"Word validation failed for {report['report_id']}: {fields} {result}")
    return {"field_update": fields, "openability": result}


def render_report(report: dict) -> dict:
    path = Path(report["docx_path"])
    qa = audit_word_report.audit_docx(path, path.parent / "qa", render=True)
    pdf_path = Path(qa["render"]["pdf_files"][0])
    pdf = fitz.open(pdf_path)
    esi_pages, recommendation_pages = [], []
    for index, page in enumerate(pdf):
        text = page.get_text()
        if "2024" in text and "千分位" in text and page.get_image_info():
            esi_pages.append(index + 1)
        if "建设建议" in text:
            recommendation_pages.append(index + 1)
    pdf.close()
    for group, page_numbers in [("all_pages", list(range(1, qa["render"]["page_count"] + 1))),
                                ("revised_pages", sorted(set(esi_pages + recommendation_pages)))]:
        columns = 4 if group == "all_pages" else 2
        width, height = (310, 450) if group == "all_pages" else (650, 920)
        rows = max(1, (len(page_numbers) + columns - 1) // columns)
        sheet = Image.new("RGB", (columns * width, rows * height), "#eeeeee")
        draw = ImageDraw.Draw(sheet)
        for cell, page_number in enumerate(page_numbers):
            image_path = Path(qa["render"]["render_dir"]) / f"page-{page_number}.png"
            with Image.open(image_path) as image:
                thumbnail = ImageOps.contain(image.convert("RGB"), (width - 12, height - 27))
            x, y = cell % columns * width + 6, cell // columns * height + 20
            sheet.paste(thumbnail, (x, y))
            draw.text((x, y - 17), f"Page {page_number}", fill="black")
        sheet.save(path.parent / "qa" / f"{group}_contact_sheet.jpg", quality=94)
    result = {"overall_status": qa["overall_status"], "page_count": qa["render"]["page_count"],
              "blank_pages": qa["render"].get("blank_pages", []), "sparse_pages": qa["render"].get("sparse_pages", []),
              "esi_pages": esi_pages, "recommendation_pages": recommendation_pages,
              "warnings": qa.get("warnings", []), "render_dir": qa["render"]["render_dir"]}
    assert result["overall_status"] != "fail" and not result["blank_pages"]
    assert esi_pages and recommendation_pages
    report["qa"] = result
    write_json(path.parent / "revision_manifest.json", report)
    print(json.dumps({"event": "rendered", "report_id": report["report_id"], **result}, ensure_ascii=True), flush=True)
    return report


def main() -> None:
    policy = editorial.load_policy()
    output = PROJECT_DIR / policy["batch_output_dir"]
    output.mkdir(parents=True, exist_ok=True)
    if BUILD_REPORTS:
        reports = [build_full_report(code, policy, output) for code in policy["selected_report_codes"]
                   if not ONLY_REPORT_IDS or code in ONLY_REPORT_IDS]
        if policy["include_unmapped_esi_topics"]:
            reports.extend(build_topic_report(key, policy, output) for key in policy["unmapped_esi_topics"]
                           if not ONLY_REPORT_IDS or key in ONLY_REPORT_IDS)
        write_json(output / "batch_manifest.json", {"generated_at": datetime.now().isoformat(timespec="seconds"),
                   "reports": reports, "status": "built_pending_qa"})
    else:
        reports = editorial.read_json(output / "batch_manifest.json")["reports"]
        if ONLY_REPORT_IDS:
            reports = [report for report in reports if report["report_id"] in ONLY_REPORT_IDS]
        allowed = set(policy["selected_report_codes"])
        if policy["include_unmapped_esi_topics"]:
            allowed.update(policy["unmapped_esi_topics"])
        reports = [report for report in reports if report["report_id"] in allowed]
    if RUN_WORD_COM:
        for report in reports:
            report["word_validation"] = word_worker(report)
            print(json.dumps({"event": "word_checked", "report_id": report["report_id"]}), flush=True)
    if RUN_RENDER_QA:
        with ThreadPoolExecutor(max_workers=RENDER_WORKERS) as pool:
            futures = [pool.submit(render_report, report) for report in reports]
            reports = [future.result() for future in as_completed(futures)]
        reports.sort(key=lambda report: report["report_id"])
    write_json(output / "batch_manifest.json", {"generated_at": datetime.now().isoformat(timespec="seconds"),
               "reports": reports, "status": "qa_completed_pending_visual_review"})
    package = output / f"同济大学{len(reports)}个ESI学科报告_精简建议修订版.zip"
    with zipfile.ZipFile(package, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for report in reports:
            path = Path(report["docx_path"])
            archive.write(path, path.name)
    print(json.dumps({"event": "completed", "report_count": len(reports), "package": str(package)}, ensure_ascii=True))


if __name__ == "__main__":
    main()
