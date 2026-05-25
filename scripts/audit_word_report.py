from __future__ import annotations

import json
import html
import re
import os
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.oxml.ns import qn
from PIL import Image, ImageStat


PROJECT_DIR = Path(__file__).resolve().parent.parent


def resolve_documents_render_script() -> Path:
    candidates = sorted(
        (Path.home() / ".codex" / "plugins" / "cache" / "openai-primary-runtime" / "documents").glob(
            "*" + "/skills/documents/render_docx.py"
        ),
        key=lambda path: path.stat().st_mtime if path.exists() else 0,
        reverse=True,
    )
    candidates.append(
        Path.home()
        / ".codex"
        / "plugins"
        / "cache"
        / "openai-primary-runtime"
        / "documents"
        / "26.430.10722"
        / "skills"
        / "documents"
        / "render_docx.py"
    )
    return next((path for path in candidates if path.exists()), candidates[-1])


DOCUMENTS_RENDER_SCRIPT = resolve_documents_render_script()

# 显式变量，避免只能靠命令行参数。
DOCX_PATH = PROJECT_DIR / "outputs" / "full_word_reports" / "latest.docx"
QA_OUTPUT_DIR = PROJECT_DIR / "outputs" / "word_qa"
ENABLE_RENDER_QA = True
EMIT_PDF = True


def now_text() -> str:
    return datetime.now().isoformat(timespec="seconds")


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8", newline="\n")


def inspect_zip(docx_path: Path) -> dict[str, Any]:
    result: dict[str, Any] = {"status": "ok", "media_count": 0, "warnings": []}
    try:
        with ZipFile(docx_path) as zip_file:
            members = set(zip_file.namelist())
            required = {"[Content_Types].xml", "word/document.xml", "word/_rels/document.xml.rels"}
            missing = sorted(required - members)
            result["member_count"] = len(members)
            media_members = sorted(member for member in members if member.startswith("word/media/"))
            result["media_count"] = len(media_members)
            result["has_document_xml"] = "word/document.xml" in members
            result["missing_required_members"] = missing
            if missing:
                result["status"] = "fail"
            if "word/document.xml" in members:
                document_xml = zip_file.read("word/document.xml").decode("utf-8", errors="ignore")
                result["odd_page_section_break_count"] = document_xml.count('w:val="oddPage"')
                if result["odd_page_section_break_count"]:
                    result["warnings"].append("document_contains_odd_page_section_breaks")
                xml_text = html.unescape(document_xml)
                stale_xml_patterns = {
                    "stale_arwu_alumni_award_template": "\u6821\u53cb\u83b7\u5956",
                    "stale_arwu_teacher_award_template": "\u6559\u5e08\u83b7\u5956",
                    "stale_arwu_ns_paper_template": "N&S\u8bba\u6587",
                    "unresolved_xx_subject": "XX\u5b66\u79d1",
                    "unresolved_xx_subject_spaced": "XX \u5b66\u79d1",
                }
                xml_hits = {
                    name: xml_text.count(pattern)
                    for name, pattern in stale_xml_patterns.items()
                    if pattern in xml_text
                }
                result["bad_text_xml_counts"] = xml_hits
                if xml_hits:
                    result["status"] = "fail"
                    result["warnings"].append("unresolved_template_text_present_in_xml")
            if "word/_rels/document.xml.rels" in members:
                rels_xml = zip_file.read("word/_rels/document.xml.rels").decode("utf-8", errors="ignore")
                media_targets = sorted(set(re.findall(r'Target="(media/[^"]+)"', rels_xml)))
                referenced_members = {f"word/{target}" for target in media_targets}
                stale_media = sorted(set(media_members) - referenced_members)
                missing_media = sorted(referenced_members - set(media_members))
                svg_members = sorted(member for member in referenced_members if member.lower().endswith(".svg"))
                result["referenced_media_count"] = len(referenced_members)
                result["stale_media_count"] = len(stale_media)
                result["missing_referenced_media"] = missing_media
                result["referenced_svg_count"] = len(svg_members)
                if stale_media:
                    result["warnings"].append("stale_unreferenced_media_members_present")
                    result["stale_media_samples"] = stale_media[:8]
                if missing_media:
                    result["status"] = "fail"
                    result["warnings"].append("referenced_media_members_missing")
                if svg_members and "[Content_Types].xml" in members:
                    content_types = zip_file.read("[Content_Types].xml").decode("utf-8", errors="ignore")
                    if 'Extension="svg"' not in content_types:
                        result["status"] = "fail"
                        result["warnings"].append("svg_media_without_content_type")
    except BadZipFile as exc:
        result["status"] = "fail"
        result["error"] = f"BadZipFile: {exc}"
    return result


def table_has_width_risk(table: Any) -> dict[str, Any]:
    row_count = len(table.rows)
    col_count = len(table.columns)
    tbl_xml = table._tbl
    pct_widths = len(tbl_xml.findall(".//" + qn("w:tblW"))) and 'w:type="pct"' in tbl_xml.xml
    fixed_layout = 'w:type="fixed"' in tbl_xml.xml
    return {
        "row_count": row_count,
        "col_count": col_count,
        "pct_widths": bool(pct_widths),
        "fixed_layout": bool(fixed_layout),
        "wide_table": col_count >= 10,
        "long_table": row_count >= 30,
    }


def inspect_docx_structure(docx_path: Path) -> dict[str, Any]:
    result: dict[str, Any] = {"status": "ok", "warnings": []}
    document = Document(str(docx_path))
    result["paragraph_count"] = len(document.paragraphs)
    result["table_count"] = len(document.tables)
    result["section_count"] = len(document.sections)
    result["inline_shape_count"] = len(document.inline_shapes)
    empty_table_count = 0
    texts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    table_risks = []
    for table_index, table in enumerate(document.tables, start=1):
        if not table.rows or not table.columns:
            empty_table_count += 1
        risk = table_has_width_risk(table)
        risk["table_index"] = table_index
        table_risks.append(risk)
        for row in table.rows:
            for cell in row.cells:
                value = cell.text.strip()
                if value:
                    texts.append(value)
    result["empty_table_count"] = empty_table_count
    if empty_table_count:
        result["warnings"].append("empty_tables_present")
    wide_tables = [row for row in table_risks if row["wide_table"] or row["long_table"] or row["pct_widths"]]
    result["table_risks"] = wide_tables
    if any(row["pct_widths"] for row in wide_tables):
        result["status"] = "fail"
        result["warnings"].append("tables_with_percentage_widths_present")
    if any(row["wide_table"] for row in wide_tables):
        result["warnings"].append("wide_tables_need_visual_review")
    if any(row["long_table"] for row in wide_tables):
        result["warnings"].append("long_tables_need_repeated_header_review")
    bad_patterns = {
        "stale_arwu_alumni_award_template": "\u6821\u53cb\u83b7\u5956",
        "stale_arwu_teacher_award_template": "\u6559\u5e08\u83b7\u5956",
        "stale_arwu_ns_paper_template": "N&S\u8bba\u6587",
        "stale_arwu_top_alumni_template": "\u9876\u5c16\u6821\u53cb",
        "debug_narrative_placeholder": "自动说明占位",
        "debug_slot_marker": "slot=",
        "debug_artifact_marker": "artifact=",
        "unresolved_xx_subject": r"(?i)xx\s*学",
        "unresolved_update_marker": "待更新",
        "draft_followup_placeholder": "后续正式正文",
        "draft_followup_reading": "后续解读",
        "draft_followup_generic": "后续可",
        "stale_qs_section_the_text": "THE排名在2023年更新指标体系后",
        "stale_the_peer_text_in_qs": "2025年发布的THE世界排名中中国内地高校",
        "subject_typo": "计算机科学于技术",
    }
    pattern_hits: dict[str, list[str]] = {}
    for name, pattern in bad_patterns.items():
        hits = []
        for text in texts:
            matched = re.search(pattern, text) if pattern.startswith("(?i)") else pattern in text
            if matched:
                hits.append(text[:240])
        if hits:
            pattern_hits[name] = hits[:8]
    result["bad_text_pattern_counts"] = {name: len(items) for name, items in pattern_hits.items()}
    result["bad_text_pattern_samples"] = pattern_hits
    if pattern_hits:
        result["status"] = "fail"
        result["warnings"].append("unresolved_template_or_debug_text_present")
    return result


def image_blankness(path: Path) -> dict[str, Any]:
    with Image.open(path) as image:
        rgb = image.convert("RGB")
        stat = ImageStat.Stat(rgb)
        extrema = rgb.getextrema()
        channel_ranges = [high - low for low, high in extrema]
        gray = rgb.convert("L")
        mask = gray.point(lambda value: 255 if value < 245 else 0)
        bbox = mask.getbbox()
        ink_pixels = sum(1 for value in mask.getdata() if value)
        ink_coverage = ink_pixels / max(1, image.width * image.height)
        return {
            "file": str(path),
            "width": image.width,
            "height": image.height,
            "mean": [round(value, 2) for value in stat.mean],
            "channel_ranges": channel_ranges,
            "looks_blank": max(channel_ranges) < 8,
            "ink_coverage": round(ink_coverage, 5),
            "content_bbox": bbox,
            "looks_sparse": ink_coverage < 0.012,
        }


def numeric_page_key(path: Path) -> int:
    match = re.search(r"page-(\d+)", path.stem)
    return int(match.group(1)) if match else 0


def render_environment() -> dict[str, str]:
    env = os.environ.copy()
    wrapper_dir = r"E:\tools\codex-bin"
    path_key = next((key for key in env if key.lower() == "path"), "Path")
    path_parts = [
        part
        for part in env.get(path_key, "").split(";")
        if part and part.rstrip("\\").lower() != r"e:\tools\libreoffice\program"
    ]
    env[path_key] = ";".join([wrapper_dir, *path_parts])
    env["Path"] = env[path_key]
    env["PATH"] = env[path_key]
    return env


def run_render(docx_path: Path, output_dir: Path) -> dict[str, Any]:
    render_dir = output_dir / "rendered_pages"
    render_dir.mkdir(parents=True, exist_ok=True)
    for stale in list(render_dir.glob("page-*.png")) + list(render_dir.glob("*.pdf")):
        stale.unlink(missing_ok=True)
    if not DOCUMENTS_RENDER_SCRIPT.exists():
        return {"status": "skipped", "reason": "documents_render_script_missing", "script": str(DOCUMENTS_RENDER_SCRIPT)}
    command = [sys.executable, str(DOCUMENTS_RENDER_SCRIPT), str(docx_path), "--output_dir", str(render_dir)]
    if EMIT_PDF:
        command.append("--emit_pdf")
    completed = subprocess.run(
        command,
        cwd=str(PROJECT_DIR),
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        timeout=240,
        env=render_environment(),
    )
    pngs = sorted(render_dir.glob("page-*.png"), key=numeric_page_key)
    page_checks = []
    for png in pngs:
        try:
            page_checks.append(image_blankness(png))
        except Exception as exc:
            page_checks.append({"file": str(png), "error": str(exc), "looks_blank": True})
    pdfs = sorted(render_dir.glob("*.pdf"))
    blank_pages = [row for row in page_checks if row.get("looks_blank")]
    sparse_pages = [row for index, row in enumerate(page_checks) if row.get("looks_sparse") and index != 0]
    pdf_ok = (not EMIT_PDF) or any(path.stat().st_size > 0 for path in pdfs)
    status = "ok" if pngs and completed.returncode == 0 and not blank_pages and pdf_ok else "fail"
    if pngs and completed.returncode != 0:
        status = "warning"
    if status == "ok" and sparse_pages:
        status = "warning"
    return {
        "status": status,
        "returncode": completed.returncode,
        "stdout": completed.stdout[-4000:],
        "stderr": completed.stderr[-4000:],
        "render_script": str(DOCUMENTS_RENDER_SCRIPT),
        "render_dir": str(render_dir),
        "page_count": len(pngs),
        "pdf_files": [str(path) for path in pdfs],
        "blank_pages": [row["file"] for row in blank_pages],
        "sparse_pages": [row["file"] for row in sparse_pages],
        "page_checks": page_checks,
    }


def audit_docx(docx_path: Path, output_dir: Path | None = None, render: bool = ENABLE_RENDER_QA) -> dict[str, Any]:
    output_dir = output_dir or QA_OUTPUT_DIR / docx_path.stem
    output_dir.mkdir(parents=True, exist_ok=True)
    payload: dict[str, Any] = {
        "generated_at": now_text(),
        "docx_path": str(docx_path),
        "qa_output_dir": str(output_dir),
        "exists": docx_path.exists(),
        "python_executable": sys.executable,
    }
    if not docx_path.exists():
        payload["overall_status"] = "fail"
        payload["error"] = "docx_missing"
        write_json(output_dir / "word_qa_report.json", payload)
        return payload

    payload["zip"] = inspect_zip(docx_path)
    payload["structure"] = inspect_docx_structure(docx_path)
    if render:
        payload["render"] = run_render(docx_path, output_dir)
    else:
        payload["render"] = {"status": "skipped", "reason": "render_disabled"}
    warnings = []
    for section in ("zip", "structure", "render"):
        warnings.extend(payload.get(section, {}).get("warnings", []))
    if payload["render"].get("status") == "warning":
        warnings.append("render_completed_with_warnings")
    if payload["render"].get("status") == "skipped":
        warnings.append("render_skipped")
    payload["warnings"] = warnings
    hard_fail = payload["zip"]["status"] == "fail" or payload["structure"]["status"] == "fail" or payload["render"]["status"] == "fail"
    payload["overall_status"] = "fail" if hard_fail else "pass_with_warnings" if warnings else "pass"
    write_json(output_dir / "word_qa_report.json", payload)
    lines = [
        "# Word QA Report",
        "",
        f"- generated_at: {payload['generated_at']}",
        f"- docx_path: `{docx_path}`",
        f"- overall_status: {payload['overall_status']}",
        f"- paragraph_count: {payload['structure'].get('paragraph_count', '')}",
        f"- table_count: {payload['structure'].get('table_count', '')}",
        f"- media_count: {payload['zip'].get('media_count', '')}",
        f"- render_status: {payload['render'].get('status', '')}",
        f"- page_count: {payload['render'].get('page_count', '')}",
        f"- warning_count: {len(warnings)}",
        "",
    ]
    (output_dir / "word_qa_report.md").write_text("\n".join(lines), encoding="utf-8", newline="\n")
    return payload


def main() -> dict[str, Any]:
    payload = audit_docx(DOCX_PATH, QA_OUTPUT_DIR / DOCX_PATH.stem)
    print(json.dumps({"overall_status": payload["overall_status"], "qa_output_dir": payload["qa_output_dir"]}, ensure_ascii=False, indent=2))
    return payload


if __name__ == "__main__":
    main()
