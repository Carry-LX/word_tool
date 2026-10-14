"""Build the independent eight-subject ESI trend report, with no CLI arguments.

The Codex Agent transcribed the reference's visible point labels once. This
script consumes that reviewed static data; it does not call an LLM or OCR API.
"""
from __future__ import annotations

import csv
import json
import math
import shutil
import zipfile
from datetime import datetime
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.ticker import FormatStrFormatter
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import inspect_esi_reference

PROJECT_DIR = Path(__file__).resolve().parent.parent
CONFIG_DIR = PROJECT_DIR / "configs/esi_trend"
CONFIG_PATH = CONFIG_DIR / "source_config.json"
VERIFIED_DATA_PATH = CONFIG_DIR / "verified_chart_data.json"
BLUE = "#5B9BD5"
GRAY = "#595959"
FONT_PATH = Path("C:/Windows/Fonts/msyh.ttc")
CHART_SIZE_INCHES = (10.5, 6.3)
CHART_DPI = 200
FIGURES_PER_PAGE = 2
CHART_WIDTH_CM = 14.8


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, value: object) -> None:
    inspect_esi_reference.write_json(path, value)


def write_csv(path: Path, rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def date_label(date: str) -> str:
    year, month = date.split("-")
    return f"{year}年{int(month)}月"


def validate_data(data: dict, inventory: dict) -> None:
    assert inventory["sha256"] == data["source_sha256"], "Reference source has changed."
    assert len(data["subjects"]) == 8
    assert len(set(data["dates"])) == len(data["dates"]) == 12
    assert data["dates"] == sorted(data["dates"])
    month_ids = [int(date[:4]) * 12 + int(date[5:]) for date in data["dates"]]
    assert all(right - left == 2 for left, right in zip(month_ids, month_ids[1:]))
    assert data["metric_unit"] == "‰" and data["better_direction"] == "lower"
    assert data["top_one_per_mille_threshold"] == 1.0
    assert data["top_one_per_ten_thousand_threshold"] == 0.1
    assert [row["source_figure_number"] for row in data["subjects"]] == list(range(3, 11))
    keys = [row["subject_key"] for row in data["subjects"]]
    assert len(set(keys)) == len(keys)
    appendix = next(block["table_rows"] for block in inventory["blocks"]
                    if block.get("table_rows") and "万分位" in block["table_rows"][0])
    name_column = appendix[0].index("学科领域")
    value_column = appendix[0].index("万分位")
    july_values = {row[name_column]: float(row[value_column].strip()) for row in appendix[1:]}
    july_index = data["dates"].index("2026-07")
    for subject in data["subjects"]:
        assert len(subject["values"]) == len(data["dates"])
        assert all(math.isfinite(value) and value > 0 for value in subject["values"])
        assert "?" not in subject["subject_name"] + subject["narrative"]
        block = inventory["blocks"][subject["source_block_index"]]
        assert subject["source_media"] in block["media_targets"]
        assert subject["subject_name"] in block["yellow_text"]
        assert subject["subject_name"] in subject["narrative"]
        assert subject["values"][-1] < 1.0
        assert round(subject["values"][july_index] * 10, 2) == july_values[subject["subject_name"]], (
            "Chart per-mille value disagrees with the independent July appendix per-ten-thousand value."
        )


def trend_summary(subject: dict, dates: list[str]) -> dict:
    values = subject["values"]
    entry_index = next((index for index, value in enumerate(values) if value < 1.0), None)
    rebounds = [dates[index] for index in range(1, len(values)) if values[index] > values[index - 1]]
    return {
        "subject_key": subject["subject_key"], "subject_name": subject["subject_name"],
        "first_period": dates[0], "first_value_per_mille": f"{values[0]:.3f}",
        "latest_period": dates[-1], "latest_value_per_mille": f"{values[-1]:.3f}",
        "improvement_per_mille": f"{values[0] - values[-1]:.3f}",
        "relative_reduction_pct": round((values[0] - values[-1]) / values[0] * 100, 2),
        "period_count": len(values), "rebound_count": len(rebounds),
        "rebound_periods": ";".join(rebounds),
        "first_observed_top_one_per_mille_period": dates[entry_index] if entry_index is not None else "",
        "already_top_one_per_mille_at_start": entry_index == 0,
        "source_figure_number": subject["source_figure_number"],
    }


def draw_chart(subject: dict, dates: list[str], out: Path) -> None:
    if not FONT_PATH.is_file():
        raise FileNotFoundError(f"Chinese chart font not found: {FONT_PATH}")
    font_manager.fontManager.addfont(str(FONT_PATH))
    plt.rcParams.update({"font.family": font_manager.FontProperties(fname=str(FONT_PATH)).get_name(),
                         "axes.unicode_minus": False, "svg.fonttype": "none"})
    figure, axis = plt.subplots(figsize=CHART_SIZE_INCHES, dpi=CHART_DPI)
    figure.patch.set_facecolor("white")
    figure.patch.set_edgecolor("#D9D9D9")
    figure.patch.set_linewidth(1.5)
    figure.subplots_adjust(left=0.105, right=0.957, bottom=0.255, top=0.815)
    x = list(range(len(dates)))
    axis.plot(x, subject["values"], color=BLUE, linewidth=2.8, marker="o", markersize=6.0)
    axis.set_title(f"{subject['subject_name']}（‰）", fontsize=22, color=GRAY, pad=28)
    axis.set_xticks(x, [date_label(date) for date in dates], rotation=48, ha="right", rotation_mode="anchor")
    axis.tick_params(axis="x", labelsize=13, labelcolor=GRAY, color="#D9D9D9", length=4, pad=6)
    axis.tick_params(axis="y", labelsize=13.5, labelcolor=GRAY, length=0, pad=12)
    ticks = subject["y_ticks"]
    axis.set_yticks(ticks)
    axis.yaxis.set_major_formatter(FormatStrFormatter("%.3f"))
    axis.set_ylim(ticks[0], ticks[-1] + (ticks[-1] - ticks[0]) * 0.065)
    axis.set_xlim(-0.18, len(dates) - 0.55)
    axis.grid(axis="y", color="#D9D9D9", linewidth=1)
    axis.set_axisbelow(True)
    for side in ("top", "left", "right"):
        axis.spines[side].set_visible(False)
    axis.spines["bottom"].set_color("#D9D9D9")
    for index, value in enumerate(subject["values"]):
        below = subject["subject_key"] in {"environment_ecology", "materials_science"} and index >= 10
        axis.annotate(f"{value:.3f}", (index, value), xytext=(0, -15 if below else 10),
                      textcoords="offset points", ha="center", va="top" if below else "bottom",
                      fontsize=13.5, color="#404040", annotation_clip=False)
    figure.savefig(out / "chart.png", dpi=CHART_DPI, facecolor="white")
    figure.savefig(out / "chart.svg", facecolor="white")
    plt.close(figure)


def build_artifact(subject: dict, data: dict, source: dict, output: Path) -> dict:
    artifact = output / "artifacts" / subject["subject_key"]
    artifact.mkdir(parents=True, exist_ok=True)
    block = source["blocks"][subject["source_block_index"]]
    caption_block = source["blocks"][subject["source_block_index"] + 1]
    rows = [{"period": period, "subject_key": subject["subject_key"], "subject_name": subject["subject_name"],
             "value_per_mille": f"{value:.3f}", "unit": "‰"}
            for period, value in zip(data["dates"], subject["values"])]
    write_csv(artifact / "chart_data.csv", rows)
    write_csv(artifact / "source_trace.csv", [dict(row, source_figure_number=subject["source_figure_number"],
              source_media=subject["source_media"], source_block_index=subject["source_block_index"],
              source_render_page=subject["source_render_page"], source_sha256=data["source_sha256"],
              extraction_method=data["transcription_method"], source_numeric_precision=3) for row in rows])
    spec = {"artifact_id": f"esi_trend_{subject['subject_key']}", "chart_type": "line_with_markers",
            "subject_name": subject["subject_name"], "x_field": "period", "y_field": "value_per_mille",
            "metric_name": data["metric_name"], "unit": data["metric_unit"], "better_direction": "lower",
            "periods": data["dates"], "values": subject["values"], "y_ticks": subject["y_ticks"],
            "point_label_decimals": 3, "line_color": BLUE, "x_label_rotation_degrees": 48,
            "width_px": int(CHART_SIZE_INCHES[0] * CHART_DPI), "height_px": int(CHART_SIZE_INCHES[1] * CHART_DPI),
            "source_figure_number": subject["source_figure_number"], "source_media": subject["source_media"],
            "source_sha256": data["source_sha256"], "source_precision": "Visible labels rounded to 3 decimals.",
            "top_one_per_mille_threshold": 1.0, "top_one_per_ten_thousand_threshold": 0.1,
            "intent": "逐期展示同济大学该ESI学科的千分位变化；折线下降表示排名改善，保留全部日期和数值标签。",
            "narrative": subject["narrative"]}
    write_json(artifact / "chart_spec.json", spec)
    write_json(artifact / "source_context.json", {"original_narrative": block["yellow_text"],
               "original_caption": caption_block["text"], "source_document": source["source_file"],
               "source_render_page": subject["source_render_page"]})
    reference_image = PROJECT_DIR / read_json(CONFIG_PATH)["source_snapshot_dir"] / "media" / Path(subject["source_media"]).name
    shutil.copyfile(reference_image, artifact / "reference.png")
    (artifact / "narrative.md").write_text(subject["narrative"] + "\n", encoding="utf-8", newline="\n")
    instructions = (f"# {subject['subject_name']}趋势图重绘说明\n\n"
                    "发送 chart_data.csv、chart_spec.json、source_trace.csv、reference.png 和本文件给绘图AI；"
                    "chart.png 为已生成样例，chart.svg 为可编辑矢量版本。\n\n"
                    "横轴为12个双月发布节点，纵轴为ESI排名千分位（‰），数值越小排名越靠前。"
                    "蓝色折线连接各节点，节点圆点旁显示三位小数，日期标签倾斜，背景为白色和浅灰水平网格。"
                    "保留全部观测值，不插值、不将‰改成%、不把0.1‰的万分之一阈值写成1‰。"
                    "各学科保留各自纵轴范围；跨学科比较请使用数据数值，不按斜率视觉比较。\n\n"
                    f"分析文字：{subject['narrative']}\n")
    (artifact / "replot_request.md").write_text(instructions, encoding="utf-8", newline="\n")
    draw_chart(subject, data["dates"], artifact)
    (artifact / "run_log.md").write_text(f"生成时间：{datetime.now().isoformat(timespec='seconds')}\n"
                                          f"源图：{subject['source_figure_number']}；数据点：12；视觉转录静态数据。\n",
                                          encoding="utf-8", newline="\n")
    return {"subject_key": subject["subject_key"], "artifact_dir": str(artifact), "point_count": len(rows)}


def font_run(run, size: float, chinese: str = "宋体", highlight: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), chinese)
    if highlight:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def append_trend_section(document: Document, subject: dict, index: int, artifact_dir: Path,
                         highlight: bool = True) -> None:
    """Reusable assembly function; ESI-to-ministry mapping is deliberately external."""
    narrative = document.add_paragraph()
    narrative.paragraph_format.first_line_indent = Cm(0.74)
    narrative.paragraph_format.line_spacing = 1.5
    narrative.paragraph_format.space_before = Pt(3)
    narrative.paragraph_format.space_after = Pt(5)
    narrative.paragraph_format.keep_with_next = True
    narrative.paragraph_format.keep_together = True
    font_run(narrative.add_run(f"（{index}）{subject['narrative']}"), 12, highlight=highlight)
    image = document.add_paragraph()
    image.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image.paragraph_format.space_before = Pt(0)
    image.paragraph_format.space_after = Pt(0)
    image.paragraph_format.line_spacing = 1
    image.paragraph_format.keep_with_next = True
    image.add_run().add_picture(str(artifact_dir / "chart.png"), width=Cm(CHART_WIDTH_CM))
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.line_spacing = 1
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_together = True
    font_run(caption.add_run(f"图{subject['source_figure_number']} {subject['subject_name']}近两年排名变化趋势"),
             10.5, highlight=highlight)


def create_document(data: dict, config: dict, output: Path) -> Path:
    document = Document()
    title_style = document.styles["Title"]
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    for border in title_style.element.xpath("./w:pPr/w:pBdr"):
        border.getparent().remove(border)
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
    section.left_margin, section.right_margin = Cm(2.2), Cm(2.2)
    section.footer_distance = Cm(0.8)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.space_after = Pt(0)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(5)
    title.paragraph_format.keep_with_next = True
    font_run(title.add_run(config["report_title"]), 16, "黑体")
    intro = document.add_paragraph()
    intro.paragraph_format.space_after = Pt(4)
    intro.paragraph_format.keep_with_next = True
    font_run(intro.add_run("观察期为2024年11月至2026年9月。千分位越小，排名越靠前。"), 10.5)
    for index, subject in enumerate(data["subjects"], start=1):
        if index > 1 and (index - 1) % FIGURES_PER_PAGE == 0:
            document.add_page_break()
        append_trend_section(document, subject, index, output / "artifacts" / subject["subject_key"],
                             config["highlight_narrative"])
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(footer.add_run("数据来源：同济ESI参考报告图3至图10；图号沿用原报告。"), 9)
    document.core_properties.title = config["report_title"]
    document.core_properties.subject = "八个ESI学科排名千分位发展趋势"
    document.core_properties.author = ""
    document.core_properties.comments = "Source SHA256: " + data["source_sha256"]
    path = output / (config["report_title"] + ".docx")
    document.save(path)
    return path


def audit_document(docx_path: Path, data: dict) -> dict:
    document = Document(docx_path)
    assert len(document.inline_shapes) == 8
    captions = [p.text for p in document.paragraphs if p.text.startswith("图")]
    assert len(captions) == 8
    assert captions[-1] == "图10 药理学与毒理学近两年排名变化趋势"
    for index, subject in enumerate(data["subjects"], start=1):
        paragraph = next(p for p in document.paragraphs if p.text == f"（{index}）{subject['narrative']}")
        assert all(run.font.highlight_color == WD_COLOR_INDEX.YELLOW for run in paragraph.runs if run.text)
    with zipfile.ZipFile(docx_path) as archive:
        assert archive.testzip() is None
    return {"status": "pass", "inline_image_count": 8, "caption_count": 8, "yellow_narrative_count": 8,
            "figure_10_caption_corrected": True, "source_point_count": 96,
            "independent_july_table_crosscheck_subject_count": 8}


def create_replot_package(output: Path, config: dict, docx_path: Path) -> Path:
    package = output / "esi_trends_replot_bundle.zip"
    with zipfile.ZipFile(package, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for file in sorted((output / "artifacts").rglob("*")):
            if file.is_file():
                archive.write(file, file.relative_to(output).as_posix())
        for file in CONFIG_DIR.glob("*.json"):
            archive.write(file, file.relative_to(PROJECT_DIR).as_posix())
        for name in ["build_esi_subject_trends.py", "inspect_esi_reference.py", "render_esi_subject_trends.py"]:
            archive.write(PROJECT_DIR / "scripts" / name, "scripts/" + name)
        source = PROJECT_DIR / config["local_source_docx"]
        archive.write(source, source.relative_to(PROJECT_DIR).as_posix())
        for name in ["all_trend_data.csv", "trend_summary.csv", "README.md", "reference_analysis.md"]:
            archive.write(output / name, name)
        archive.write(docx_path, docx_path.name)
    return package


def main() -> dict:
    config, data = read_json(CONFIG_PATH), read_json(VERIFIED_DATA_PATH)
    inspect_esi_reference.inspect_reference()
    inventory = read_json(PROJECT_DIR / config["source_snapshot_dir"] / "reference_inventory.json")
    validate_data(data, inventory)
    output = PROJECT_DIR / config["output_dir"]
    results = [build_artifact(subject, data, inventory, output) for subject in data["subjects"]]
    all_rows = [{"subject_key": s["subject_key"], "subject_name": s["subject_name"], "period": period,
                 "value_per_mille": f"{value:.3f}", "unit": "‰", "source_figure_number": s["source_figure_number"]}
                for s in data["subjects"] for period, value in zip(data["dates"], s["values"])]
    write_csv(output / "all_trend_data.csv", all_rows)
    summaries = [trend_summary(subject, data["dates"]) for subject in data["subjects"]]
    write_csv(output / "trend_summary.csv", summaries)
    docx_path = create_document(data, config, output)
    audit = audit_document(docx_path, data)
    write_json(output / "qa/structure_audit.json", audit)
    analysis = ("# ESI参考报告解析记录\n\n"
                f"源文件SHA256：{data['source_sha256']}\n\n"
                "原文共84个正文块、7个物理表格、11个图片对象；没有原生图表XML或嵌入工作簿。"
                "完整源报告已通过标准链路渲染为20页。图3至图10位于渲染第9至16页，"
                "每张图都有12个双月节点；从可见标签转录得到96条趋势观测值。\n\n"
                "原文结构依次为总体概况、百分之一学科分布、工程学万分之一对比、千分之一基础表、"
                "八学科趋势、E9高校对比、内地高校对比和2026年7月详细数据附表。"
                "新样例仅选八学科趋势段落。\n\n"
                "## 口径与修正\n\n" + "\n".join("- " + note for note in data["source_context_notes"]) + "\n"
                "- 原图10的图题误写地球科学，已按图片标题及黄色文字修正为药理学与毒理学。\n\n"
                "## 数据处理链路\n\n"
                "原始DOCX → 按正文顺序提取文字、黄色高亮、表格和图片 → Agent逐图核对标题、日期与数值"
                " → 保存UTF-8静态数据 → 计算期初/期末变化和阈值进入节点 → 重绘8图 → 装配独立Word。"
                "图中数值保留源图三位小数，不推算舍入前的值。进入前1‰的月份只按观察节点判定。\n\n"
                "## 复核重点\n\n"
                "环境/生态学2026年7月0.235‰与源附表2.35万分位一致；9月0.231‰来自图3。"
                "社会科学总论首次观察到低于1‰为2026年3月，地球科学为7月，药理学与毒理学为5月。"
                "材料科学回升1期、计算机科学2期、临床医学2期，其余5科序列持续下降。\n")
    (output / "reference_analysis.md").write_text(analysis, encoding="utf-8", newline="\n")
    readme = ("# ESI八学科趋势模块\n\n"
              "直接运行 scripts/build_esi_subject_trends.py；输入与输出在 configs/esi_trend/source_config.json 中显式配置。"
              "数据在 verified_chart_data.json，来源为已逐图核对的可见标签，无需外部模型API。\n\n"
              "每个学科输出一张蓝色折线图、一段黄色分析文字、图题、CSV数据、图表规格、来源追踪及重绘说明。"
              "独立Word每页两个学科，保留源图号3至10。\n\n"
              "重绘单图：发送 artifacts/对应学科/ 内的 chart_data.csv、chart_spec.json、source_trace.csv、"
              "reference.png、replot_request.md；如需全部重绘，发送 esi_trends_replot_bundle.zip。\n\n"
              "all_trend_data.csv 为96条总数据，trend_summary.csv 为8科分析结果。"
              "若接入现有一级学科报告，可复用 append_trend_section；ESI与教育部学科的对应关系需另行明确配置。\n")
    (output / "README.md").write_text(readme, encoding="utf-8", newline="\n")
    package = create_replot_package(output, config, docx_path)
    result = {"generated_at": datetime.now().isoformat(timespec="seconds"), "docx_path": str(docx_path),
              "replot_package": str(package), "artifact_count": len(results), "point_count": len(all_rows),
              "structure_audit": audit, "artifacts": results, "render_visual_qa": "pending"}
    write_json(output / "run_summary.json", result)
    return result


if __name__ == "__main__":
    print(json.dumps(main(), ensure_ascii=True, indent=2))
