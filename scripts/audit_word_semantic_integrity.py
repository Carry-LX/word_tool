from __future__ import annotations

import csv
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document


PROJECT_DIR = Path(__file__).resolve().parent.parent
FULL_WORD_REPORTS_DIR = PROJECT_DIR / "outputs" / "full_word_reports"
COVERAGE_CSV = PROJECT_DIR / "outputs" / "subject_coverage_audit" / "tongji_49_academic_degree_generation_coverage_after_supplement.csv"
OUTPUT_DIR = PROJECT_DIR / "outputs" / "word_semantic_audit" / "latest"

RANKING_SOURCES = ["软科", "QS", "US News", "THE"]
SOURCE_LABELS = {
    "软科": ["软科", "ARWU"],
    "QS": ["QS"],
    "US News": ["U.S. News", "U.S.News", "US News"],
    "THE": ["THE"],
}


def now_text() -> str:
    return datetime.now().isoformat(timespec="seconds")


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and pd.isna(value):
        return ""
    text = str(value).strip()
    return "" if text.lower() == "nan" else text


def normalize_code(value: Any) -> str:
    text = re.sub(r"\D+", "", clean_text(value))
    return text.zfill(4) if text else ""


def normalize_key(value: Any) -> str:
    text = clean_text(value).lower().replace("&", " and ")
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "", text)


def read_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8-sig"))


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8", newline="\n")


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def load_coverage() -> dict[tuple[str, str], dict[str, str]]:
    if not COVERAGE_CSV.exists():
        return {}
    df = pd.read_csv(COVERAGE_CSV, dtype=str, encoding="utf-8-sig").fillna("")
    rows: dict[tuple[str, str], dict[str, str]] = {}
    for _, row in df.iterrows():
        code = normalize_code(row.get("ministry_subject_code"))
        name = clean_text(row.get("ministry_subject_name") or row.get("source_subject_name"))
        if not code or not name:
            continue
        rows[(code, name)] = {key: clean_text(value) for key, value in row.items()}
    return rows


def code_from_dir(path: Path) -> str:
    match = re.search(r"_(\d{4})_", path.name)
    return match.group(1) if match else ""


def subject_name_from_dir(path: Path) -> str:
    match = re.match(r"^\d{8}_\d{6}_(\d{4})_(.+)$", path.name)
    return match.group(2).strip() if match else ""


def latest_report_dirs() -> list[Path]:
    latest: dict[tuple[str, str], Path] = {}
    if not FULL_WORD_REPORTS_DIR.exists():
        return []
    for path in FULL_WORD_REPORTS_DIR.iterdir():
        if not path.is_dir():
            continue
        code = code_from_dir(path)
        name = subject_name_from_dir(path)
        if not code or not name:
            continue
        key = (code, name)
        if key not in latest or path.stat().st_mtime > latest[key].stat().st_mtime:
            latest[key] = path
    return [latest[key] for key in sorted(latest)]


def source_mentions(source: str, text: str) -> bool:
    return any(label in text for label in SOURCE_LABELS[source])


def mapping_for_source(coverage_row: dict[str, str], source: str) -> str:
    field = {
        "软科": "arwu_mapping",
        "QS": "qs_mapping",
        "US News": "usnews_mapping",
        "THE": "the_mapping",
    }[source]
    value = clean_text(coverage_row.get(field))
    if normalize_key(value) in {"", "无", "none", "nan", "暂无"}:
        return ""
    return value


def override_mapping_for_source(run_dir: Path, source: str) -> str:
    payload = read_json(run_dir / "agent_run_summary.json")
    resolved = payload.get("resolved_subject") or {}
    overrides = resolved.get("source_runtime_overrides") or {}
    source_override = overrides.get(source) or {}
    if isinstance(source_override, dict):
        value = clean_text(source_override.get("ranking_subject_or_area")) or clean_text(source_override.get("weight_subject_or_area"))
        if value:
            return value
    field = {
        "软科": "arwu_subject_name",
        "QS": "qs_subject_name",
        "US News": "usnews_subject_name",
        "THE": "the_subject_area",
    }.get(source, "")
    return clean_text(resolved.get(field)) if field else ""


def mapping_parts(value: str) -> list[str]:
    parts = re.split(r"[;；、,\n\r]+", clean_text(value))
    rows = []
    seen = set()
    for part in parts:
        text = clean_text(part)
        key = normalize_key(text)
        if key and key not in seen:
            rows.append(text)
            seen.add(key)
    return rows


def mapping_appears_in_text(mapping: str, all_text: str, subject: str) -> bool:
    mapping_key = normalize_key(mapping)
    text_key = normalize_key(all_text)
    if mapping_key and mapping_key in text_key:
        return True
    parts = mapping_parts(mapping)
    subject_key = normalize_key(subject)
    non_subject_parts = [part for part in parts if normalize_key(part) != subject_key]
    candidates = non_subject_parts or parts
    return any(normalize_key(part) in text_key for part in candidates if len(normalize_key(part)) >= 2)


def sources_from_semicolon(value: str) -> set[str]:
    return {clean_text(item) for item in clean_text(value).split(";") if clean_text(item)}


def read_docx_text(docx_path: Path) -> tuple[list[str], str]:
    document = Document(str(docx_path))
    texts: list[str] = []
    for paragraph in document.paragraphs:
        value = clean_text(paragraph.text)
        if value:
            texts.append(value)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                value = clean_text(cell.text)
                if value:
                    texts.append(value)
    return texts, "\n".join(texts)


def replacement_audit_summary(report_dir: Path) -> tuple[int | None, int | None]:
    payload = read_json(report_dir / "replacement_audit.json")
    summary = payload.get("summary") or {}
    ok = summary.get("ok")
    mismatch = summary.get("mismatch")
    if ok is None and "audit_rows" in payload:
        rows = payload.get("audit_rows") or []
        ok = sum(1 for row in rows if row.get("status") == "ok")
        mismatch = sum(1 for row in rows if row.get("status") != "ok")
    return (int(ok) if ok is not None else None, int(mismatch) if mismatch is not None else None)


def artifact_rows(summary: dict[str, Any], artifact_id: str) -> int:
    for row in (summary.get("artifact_summary") or {}).get("results") or []:
        if row.get("artifact_id") == artifact_id:
            try:
                return int(float(row.get("rows") or 0))
            except (TypeError, ValueError):
                return 0
    return 0


def read_artifact_csv(run_dir: Path, artifact_id: str, filename: str) -> list[dict[str, str]]:
    path = run_dir / "artifacts" / artifact_id / filename
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        return [{key: clean_text(value) for key, value in row.items()} for row in csv.DictReader(file)]


def unique_values(rows: list[dict[str, str]], key: str) -> set[str]:
    return {clean_text(row.get(key)) for row in rows if clean_text(row.get(key))}


def page_number_from_path(value: str) -> int:
    match = re.search(r"page-(\d+)", Path(value).stem)
    return int(match.group(1)) if match else 0


def split_sparse_pages(render: dict[str, Any]) -> tuple[list[str], list[str]]:
    page_count = int(render.get("page_count") or 0)
    frontmatter_sparse: list[str] = []
    blocking_sparse: list[str] = []
    for page in render.get("sparse_pages") or []:
        page_number = page_number_from_path(page)
        if page_number and page_number <= 3:
            frontmatter_sparse.append(page)
        elif page_count and page_number == page_count:
            blocking_sparse.append(page)
        else:
            blocking_sparse.append(page)
    return frontmatter_sparse, blocking_sparse


def is_broad_unavailable_note(text: str) -> bool:
    scoped_terms = [
        "该位置",
        "多年",
        "历史",
        "对标样本",
        "相对排名图",
        "权重",
        "分项指标相对排名",
    ]
    return not any(term in text for term in scoped_terms)


def add_issue(issues: list[dict[str, Any]], report: dict[str, Any], severity: str, area: str, item: str, message: str, evidence: str = "") -> None:
    issues.append(
        {
            "severity": severity,
            "area": area,
            "item": item,
            "code": report.get("code", ""),
            "subject": report.get("subject", ""),
            "report_dir": report.get("report_dir", ""),
            "message": message,
            "evidence": evidence,
        }
    )


def audit_report(report_dir: Path, coverage: dict[tuple[str, str], dict[str, str]]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    code = code_from_dir(report_dir)
    subject = subject_name_from_dir(report_dir)
    report = {
        "code": code,
        "subject": subject,
        "report_dir": str(report_dir),
        "docx": "",
        "issue_count": 0,
        "p0_count": 0,
        "p1_count": 0,
        "p2_count": 0,
    }
    issues: list[dict[str, Any]] = []
    docx_files = sorted(report_dir.glob("*.docx"), key=lambda item: item.stat().st_mtime, reverse=True)
    if not docx_files:
        add_issue(issues, report, "P0", "word", "missing_docx", "报告目录中没有 docx 文件。")
        report["issue_count"] = len(issues)
        return report, issues
    docx_path = docx_files[0]
    report["docx"] = str(docx_path)
    try:
        texts, all_text = read_docx_text(docx_path)
    except Exception as exc:
        add_issue(issues, report, "P0", "word", "docx_read_failed", f"无法读取 Word 正文：{exc}")
        report["issue_count"] = len(issues)
        return report, issues

    summary = read_json(report_dir / "subject_word_report_summary.json")
    qa = read_json(report_dir / "qa" / "word_qa_report.json") or (summary.get("word_qa") or {})
    qa_status = clean_text(qa.get("overall_status")) or "missing"
    report["qa_status"] = qa_status
    if qa_status not in {"pass", "pass_with_warnings"}:
        add_issue(issues, report, "P0", "word_qa", "qa_not_pass", f"Word QA 状态为 {qa_status}。")
    render = qa.get("render") or {}
    if render.get("blank_pages"):
        add_issue(issues, report, "P0", "word_qa", "blank_pages", "渲染 QA 发现空白页。", ";".join(render.get("blank_pages") or []))
    frontmatter_sparse_pages, blocking_sparse_pages = split_sparse_pages(render)
    report["frontmatter_sparse_pages"] = ";".join(Path(item).name for item in frontmatter_sparse_pages)
    report["blocking_sparse_pages"] = ";".join(Path(item).name for item in blocking_sparse_pages)
    if blocking_sparse_pages:
        add_issue(issues, report, "P1", "word_qa", "blocking_sparse_pages", "Sparse rendered pages outside front matter need visual/layout review.", ";".join(Path(item).name for item in blocking_sparse_pages))

    ok, mismatch = replacement_audit_summary(report_dir)
    report["replacement_ok"] = ok if ok is not None else ""
    report["replacement_mismatch"] = mismatch if mismatch is not None else ""
    if mismatch is None:
        add_issue(issues, report, "P1", "replacement", "missing_replacement_audit", "缺少 Word 图表替换审计。")
    elif mismatch:
        add_issue(issues, report, "P0", "replacement", "replacement_mismatch", f"Word 图表替换审计存在 {mismatch} 个 mismatch。")

    stale_patterns = {
        "unresolved_xx_subject": r"(?i)xx\s*学科",
        "unresolved_update_marker": "待更新",
        "stale_the_history_template": "THE排名在2023年更新指标体系后研究质量、国际展望相对排名比较低",
        "stale_the_peer_template": "2025年发布的THE世界排名中中国内地高校有13所进入前200位",
        "stale_the_indicator_template": "2025年THE排名前200位的3所内地高校",
        "generic_followup_placeholder": "后续正式正文",
    }
    for item, pattern in stale_patterns.items():
        hits = []
        for text in texts:
            matched = re.search(pattern, text) if pattern.startswith("(?i)") else pattern in text
            if matched:
                hits.append(text[:180])
        if hits:
            add_issue(issues, report, "P1", "template_text", item, "正文中仍有模板残留或占位文本。", " | ".join(hits[:3]))

    run_dir_text = clean_text(summary.get("run_dir"))
    run_dir = Path(run_dir_text) if run_dir_text else Path()
    coverage_row = coverage.get((code, subject), {})
    if coverage_row:
        ready_sources = sources_from_semicolon(coverage_row.get("tongji_sources") or coverage_row.get("ranking_sources"))
        for source in RANKING_SOURCES:
            mapping = override_mapping_for_source(run_dir, source) if run_dir.exists() else ""
            mapping = mapping or mapping_for_source(coverage_row, source)
            if not mapping or source not in ready_sources:
                continue
            if source_mentions(source, all_text) and not mapping_appears_in_text(mapping, all_text, subject):
                add_issue(
                    issues,
                    report,
                    "P1",
                    "source_scope",
                    f"missing_{source}_scope_label",
                    f"{source} 有可用口径映射“{mapping}”，但 Word 正文/图题未明确出现该口径。",
                    f"mapping={mapping}",
                )
        the_mapping = mapping_for_source(coverage_row, "THE")
        if the_mapping and normalize_key(the_mapping) != normalize_key("自然科学") and "THE世界学科排名自然科学领域" in all_text:
            add_issue(
                issues,
                report,
                "P1",
                "source_scope",
                "wrong_default_the_area",
                f"THE 图题仍使用默认自然科学领域，但该学科 THE 映射为“{the_mapping}”。",
            )

    for source in RANKING_SOURCES:
        unavailable_texts = [
            text
            for text in texts
            if source_mentions(source, text)
            and (
                "暂无可直接用于本报告" in text
                or "暂无可核验" in text
                or "未覆盖该榜单" in text
                or "不生成对应图表" in text
            )
            and "仅保留" not in text
            and is_broad_unavailable_note(text)
        ]
        if not unavailable_texts:
            continue
        data_like = [
            text
            for text in texts
            if source_mentions(source, text)
            and re.search(r"^(图|表)\s*\d|第\s*\d|#\s*\d|\d+\s*[-–]\s*\d+|进入前200|指标权重", text)
            and text not in unavailable_texts
        ]
        if data_like:
            add_issue(
                issues,
                report,
                "P1",
                "source_availability",
                f"{source}_unavailable_but_data_text",
                f"正文一处说 {source} 对应数据/图表不可用，但同一报告又出现该来源的数据化描述或图表标题。",
                f"unavailable={unavailable_texts[0][:160]} | data={data_like[0][:160]}",
            )

    if run_dir_text and not run_dir.exists():
        add_issue(issues, report, "P2", "traceability", "missing_run_dir", "Word 摘要指向的 subject_agent_run 已不存在，无法追溯图表原始 chart_data。", run_dir_text)
    if run_dir.exists():
        overview_rows = read_artifact_csv(run_dir, "figure_01_tongji_subject_ranking_overview", "chart_data.csv")
        history_rows = read_artifact_csv(run_dir, "figure_07_all_rankings_history", "chart_data.csv")
        overview_sources = unique_values(overview_rows, "ranking_source")
        history_sources = unique_values(history_rows, "ranking_source")
        if history_sources and not overview_sources.issuperset(history_sources):
            missing = "、".join(sorted(history_sources - overview_sources))
            add_issue(issues, report, "P1", "artifact_consistency", "overview_history_source_mismatch", f"历史趋势图包含 {missing}，但总览图没有同步显示。")
        usnews_rows = read_artifact_csv(run_dir, "figure_10_usnews_indicator_history_trend", "chart_data.csv")
        usnews_spec = read_json(run_dir / "artifacts" / "figure_10_usnews_indicator_history_trend" / "chart_spec.json")
        indicator_count = len(unique_values(usnews_rows, "indicator_name"))
        rendering_strategy = clean_text(usnews_spec.get("rendering_strategy"))
        numeric_value_source = clean_text(usnews_spec.get("numeric_value_source"))
        if indicator_count > 8 and not rendering_strategy.startswith("all_indicator"):
            add_issue(
                issues,
                report,
                "P1",
                "chart_rendering",
                "usnews_indicator_trend_had_more_than_8_series",
                f"U.S. News 分项趋势数据含 {indicator_count} 个指标；旧渲染逻辑只展示前 8 个，需使用全量小多图或明确截断说明。",
            )
        if usnews_rows and any(clean_text(row.get("indicator_value")).startswith("#") for row in usnews_rows) and "cleaned_indicator_value" not in numeric_value_source:
            add_issue(
                issues,
                report,
                "P1",
                "chart_rendering",
                "usnews_indicator_value_hash_prefix",
                "U.S. News 分项趋势存在 #203 这类排名文本，渲染时必须使用 indicator_value_num 或清洗 # 后再转数值。",
            )
    else:
        usnews_rows = artifact_rows(summary, "figure_10_usnews_indicator_history_trend")
        if usnews_rows > 24:
            add_issue(
                issues,
                report,
                "P1",
                "chart_rendering",
                "usnews_indicator_trend_probable_truncation",
                f"摘要显示 U.S. News 分项趋势有 {usnews_rows} 行；按 3 年估算超过 8 个指标，旧渲染逻辑可能只展示前 8 个。",
            )

    for severity in ["P0", "P1", "P2"]:
        report[f"{severity.lower()}_count"] = sum(1 for issue in issues if issue["severity"] == severity)
    report["issue_count"] = len(issues)
    return report, issues


def main() -> dict[str, Any]:
    coverage = load_coverage()
    report_rows: list[dict[str, Any]] = []
    issue_rows: list[dict[str, Any]] = []
    for report_dir in latest_report_dirs():
        report, issues = audit_report(report_dir, coverage)
        report_rows.append(report)
        issue_rows.extend(issues)

    summary = {
        "generated_at": now_text(),
        "report_count": len(report_rows),
        "issue_count": len(issue_rows),
        "p0_count": sum(1 for row in issue_rows if row["severity"] == "P0"),
        "p1_count": sum(1 for row in issue_rows if row["severity"] == "P1"),
        "p2_count": sum(1 for row in issue_rows if row["severity"] == "P2"),
        "issue_counts_by_item": {},
        "outputs": {
            "summary_json": str(OUTPUT_DIR / "word_semantic_audit_summary.json"),
            "issues_csv": str(OUTPUT_DIR / "word_semantic_issues.csv"),
            "reports_csv": str(OUTPUT_DIR / "word_semantic_reports.csv"),
            "markdown": str(OUTPUT_DIR / "word_semantic_audit.md"),
        },
    }
    for row in issue_rows:
        key = f"{row['severity']}|{row['area']}|{row['item']}"
        summary["issue_counts_by_item"][key] = summary["issue_counts_by_item"].get(key, 0) + 1

    report_fields = [
        "code",
        "subject",
        "report_dir",
        "docx",
        "qa_status",
        "frontmatter_sparse_pages",
        "blocking_sparse_pages",
        "replacement_ok",
        "replacement_mismatch",
        "issue_count",
        "p0_count",
        "p1_count",
        "p2_count",
    ]
    issue_fields = ["severity", "area", "item", "code", "subject", "report_dir", "message", "evidence"]
    write_json(OUTPUT_DIR / "word_semantic_audit_summary.json", summary)
    write_csv(OUTPUT_DIR / "word_semantic_reports.csv", report_rows, report_fields)
    write_csv(OUTPUT_DIR / "word_semantic_issues.csv", issue_rows, issue_fields)

    lines = [
        "# Word Semantic Audit",
        "",
        f"- generated_at: {summary['generated_at']}",
        f"- report_count: {summary['report_count']}",
        f"- issue_count: {summary['issue_count']}",
        f"- P0: {summary['p0_count']}",
        f"- P1: {summary['p1_count']}",
        f"- P2: {summary['p2_count']}",
        "",
        "## Issue Types",
        "",
    ]
    for key, count in sorted(summary["issue_counts_by_item"].items()):
        lines.append(f"- {key}: {count}")
    lines.extend(["", "## P0/P1 Issues", ""])
    for row in issue_rows:
        if row["severity"] in {"P0", "P1"}:
            lines.append(f"- {row['severity']} {row['code']} {row['subject']} [{row['area']}/{row['item']}]: {row['message']}")
    (OUTPUT_DIR / "word_semantic_audit.md").write_text("\n".join(lines), encoding="utf-8", newline="\n")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return summary


if __name__ == "__main__":
    main()
