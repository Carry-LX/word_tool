from __future__ import annotations

import json
import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook


PROJECT_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_DIR / "data"
OUTPUT_DIR = PROJECT_DIR / "configs" / "stage1_inventory"

TARGET_TEMPLATE = DATA_DIR / "【标黄】同济大学学科分析报告——理科（2026年）.docx"
REFERENCE_REPORT = DATA_DIR / "同济大学XX学科的国际排名发展分析-2026年5月.docx"

MAX_SAMPLE_ROWS = 4
MAX_SAMPLE_COLS = 18


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def safe_name(path: Path) -> str:
    return path.name


def json_default(value: Any) -> str:
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value)


def write_json(path: Path, data: Any) -> None:
    path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2, default=json_default),
        encoding="utf-8",
        newline="\n",
    )


def yaml_scalar(value: Any) -> str:
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)):
        return str(value)
    text = str(value)
    escaped = text.replace("\\", "\\\\").replace('"', '\\"')
    return f'"{escaped}"'


def to_yaml(data: Any, indent: int = 0) -> str:
    prefix = " " * indent
    if isinstance(data, dict):
        lines: list[str] = []
        for key, value in data.items():
            if isinstance(value, (dict, list)):
                lines.append(f"{prefix}{key}:")
                lines.append(to_yaml(value, indent + 2))
            else:
                lines.append(f"{prefix}{key}: {yaml_scalar(value)}")
        return "\n".join(lines)
    if isinstance(data, list):
        lines = []
        if not data:
            return f"{prefix}[]"
        for item in data:
            if isinstance(item, dict):
                lines.append(f"{prefix}-")
                lines.append(to_yaml(item, indent + 2))
            elif isinstance(item, list):
                lines.append(f"{prefix}-")
                lines.append(to_yaml(item, indent + 2))
            else:
                lines.append(f"{prefix}- {yaml_scalar(item)}")
        return "\n".join(lines)
    return f"{prefix}{yaml_scalar(data)}"


def write_yaml(path: Path, data: Any) -> None:
    path.write_text(to_yaml(data) + "\n", encoding="utf-8", newline="\n")


def infer_workbook_role(file_name: str, sheet_name: str, headers: list[str]) -> dict[str, Any]:
    joined = " ".join([file_name, sheet_name, *headers]).lower()
    source = []
    if "us news" in joined or "u.s.news" in joined or "usnews" in joined:
        source.append("US News")
    if "qs" in joined:
        source.append("QS")
    if "软科" in joined or "arwu" in joined:
        source.append("软科")
    if "the" in joined or "泰晤士" in joined:
        source.append("THE")

    role = "supporting_reference"
    if any(k in joined for k in ["国际排名", "rank", "排名", "总得分", "overall"]):
        role = "ranking_detail"
    if any(k in joined for k in ["指标权重", "权重", "indicator"]):
        role = "indicator_weight"
    if any(k in joined for k in ["映射", "对应", "学科代码", "一级学科"]):
        role = "subject_mapping"
    if any(k in joined for k in ["评估结果", "双一流", "高峰高原"]):
        role = "domestic_evaluation"
    if any(k in joined for k in ["学院", "学术学位点", "专业学位点"]):
        role = "tongji_subject_mapping"

    available_fields = {
        "university": any(k in joined for k in ["学校", "高校", "大学", "中文名称", "university", "name", "单位名称"]),
        "subject": any(k in joined for k in ["学科", "subject", "领域"]),
        "rank": any(k in joined for k in ["排名", "rank", "位次"]),
        "score": any(k in joined for k in ["得分", "总分", "score"]),
        "year": any(k in joined for k in ["年度", "年份", "year", "发布时间"]),
        "region": any(k in joined for k in ["地区", "省份", "国家", "country", "location"]),
        "indicator": any(k in joined for k in ["指标", "声誉", "论文", "引用", "教学", "研究", "国际", "权重"]),
    }

    return {
        "role": role,
        "ranking_sources": source,
        "available_fields": available_fields,
    }


def first_nonempty_header_and_samples(ws: Any) -> tuple[int | None, list[str], list[list[str]]]:
    header_row_index = None
    headers: list[str] = []
    samples: list[list[str]] = []

    max_scan_row = min(ws.max_row or 0, 12)
    for row_index, row in enumerate(
        ws.iter_rows(min_row=1, max_row=max_scan_row, values_only=True),
        start=1,
    ):
        values = [clean_text(value) for value in row[:MAX_SAMPLE_COLS]]
        nonempty = [value for value in values if value]
        if len(nonempty) >= 2:
            header_row_index = row_index
            headers = values
            break

    if header_row_index is not None:
        for row in ws.iter_rows(
            min_row=header_row_index + 1,
            max_row=min(ws.max_row or 0, header_row_index + MAX_SAMPLE_ROWS),
            values_only=True,
        ):
            values = [clean_text(value) for value in row[:MAX_SAMPLE_COLS]]
            if any(values):
                samples.append(values)

    return header_row_index, headers, samples


def build_data_catalog() -> dict[str, Any]:
    files = sorted(DATA_DIR.glob("*.xlsx"))
    catalog: dict[str, Any] = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "data_dir": str(DATA_DIR),
        "excel_count": len(files),
        "workbooks": [],
    }

    for workbook_path in files:
        workbook_entry: dict[str, Any] = {
            "file": safe_name(workbook_path),
            "path": str(workbook_path),
            "size_bytes": workbook_path.stat().st_size,
            "sheets": [],
        }
        try:
            wb = load_workbook(workbook_path, read_only=True, data_only=True)
            for ws in wb.worksheets:
                header_row, headers, samples = first_nonempty_header_and_samples(ws)
                inferred = infer_workbook_role(workbook_path.name, ws.title, headers)
                workbook_entry["sheets"].append(
                    {
                        "sheet": ws.title,
                        "max_row": ws.max_row,
                        "max_column": ws.max_column,
                        "header_row": header_row,
                        "headers": headers,
                        "sample_rows": samples,
                        **inferred,
                    }
                )
            wb.close()
        except Exception as exc:
            workbook_entry["error"] = f"{type(exc).__name__}: {exc}"
        catalog["workbooks"].append(workbook_entry)

    return catalog


def paragraph_kind(text: str) -> str:
    if re.match(r"^[一二三四五六七八九十]+、", text):
        return "section"
    if re.match(r"^\d+、", text):
        return "subsection"
    if re.match(r"^图\d+", text):
        return "figure_caption"
    if re.match(r"^表\d+", text):
        return "table_caption"
    if text:
        return "body"
    return "blank"


def extract_doc_structure(path: Path) -> dict[str, Any]:
    doc = Document(path)
    paragraphs = []
    highlighted = []
    for index, paragraph in enumerate(doc.paragraphs, start=1):
        text = clean_text(paragraph.text)
        if not text:
            continue
        highlight_runs = []
        for run in paragraph.runs:
            if run.font.highlight_color is not None and clean_text(run.text):
                highlight_runs.append(
                    {
                        "text": clean_text(run.text),
                        "highlight": str(run.font.highlight_color),
                    }
                )
        entry = {
            "paragraph_index": index,
            "kind": paragraph_kind(text),
            "text": text,
            "highlight_runs": highlight_runs,
        }
        paragraphs.append(entry)
        if highlight_runs:
            highlighted.append(entry)

    tables = []
    for table_index, table in enumerate(doc.tables, start=1):
        row_count = len(table.rows)
        col_count = max((len(row.cells) for row in table.rows), default=0)
        sample_rows = []
        for row in table.rows[:4]:
            sample_rows.append([clean_text(cell.text) for cell in row.cells[:10]])
        tables.append(
            {
                "table_index": table_index,
                "rows": row_count,
                "columns": col_count,
                "sample_rows": sample_rows,
            }
        )

    images = []
    for image_index, shape in enumerate(doc.inline_shapes, start=1):
        images.append(
            {
                "image_index": image_index,
                "width_emu": shape.width,
                "height_emu": shape.height,
                "width_cm": round(shape.width / 360000, 2),
                "height_cm": round(shape.height / 360000, 2),
            }
        )

    return {
        "file": safe_name(path),
        "path": str(path),
        "paragraph_count": len(doc.paragraphs),
        "nonempty_paragraph_count": len(paragraphs),
        "table_count": len(doc.tables),
        "image_count": len(doc.inline_shapes),
        "paragraphs": paragraphs,
        "highlighted_paragraphs": highlighted,
        "tables": tables,
        "images": images,
    }


def build_reference_report_markdown(reference: dict[str, Any]) -> str:
    lines = [
        "# 参考报告结构摘录",
        "",
        f"- 文件：{reference['file']}",
        f"- 非空段落：{reference['nonempty_paragraph_count']}",
        f"- 表格数量：{reference['table_count']}",
        f"- 图片数量：{reference['image_count']}",
        "",
        "## 章节、图表与表格标题",
        "",
    ]
    for paragraph in reference["paragraphs"]:
        if paragraph["kind"] in {"section", "subsection", "figure_caption", "table_caption"}:
            lines.append(f"- P{paragraph['paragraph_index']} `{paragraph['kind']}`：{paragraph['text']}")

    lines.extend(["", "## 表格结构样例", ""])
    for table in reference["tables"]:
        sample = table["sample_rows"][0] if table["sample_rows"] else []
        lines.append(
            f"- 表{table['table_index']}：{table['rows']}行 x {table['columns']}列；首行：{' | '.join(sample)}"
        )

    lines.extend(["", "## 图片尺寸样例", ""])
    for image in reference["images"]:
        lines.append(
            f"- 图像{image['image_index']}：约 {image['width_cm']}cm x {image['height_cm']}cm"
        )

    return "\n".join(lines) + "\n"


def build_report_slots() -> dict[str, Any]:
    slots = {
        "version": 1,
        "purpose": "把理科学科分析报告模板拆成可由数据分析 agent 执行的报告位置。",
        "target_template": safe_name(TARGET_TEMPLATE),
        "reference_report": safe_name(REFERENCE_REPORT),
        "global_context": {
            "target_university": "同济大学",
            "default_subject_family": "理科",
            "default_year": 2026,
            "reference_subject": "数学",
        },
        "slots": [
            {
                "id": "overview_competitive_summary",
                "section": "一、总述",
                "output_types": ["text"],
                "goal": "概述同济理科学科的竞争格局、优势、挑战与机会。",
                "reference_style": "参考报告“一、概述”的长段落叙述。",
                "verification": "文字中的排名、数量、年份必须可追溯到后续明细数据。",
            },
            {
                "id": "ranking_indicator_system",
                "section": "二、国际排名/1、四大学科排名指标体系分析",
                "output_types": ["text", "chart"],
                "goal": "比较 QS、THE、US News、软科四大排名指标体系及权重共性。",
                "reference_style": "参考报告图2-图4及对应说明。",
                "verification": "各指标权重加总与权重表一致。",
            },
            {
                "id": "tongji_ranking_history",
                "section": "二、国际排名/2、同济理科在四大国际学科排名历史分析",
                "output_types": ["text", "table", "chart"],
                "goal": "展示同济对应学科近年四大国际排名变化趋势。",
                "reference_style": "参考报告图5、表1-表3的历史趋势表达。",
                "verification": "每条趋势至少保留来源榜单、发布年度、排名年度、原始排名字段。",
            },
            {
                "id": "indicator_detail_comparison",
                "section": "二、国际排名/3、国际学科排名的细分指标数据对比",
                "output_types": ["text", "table", "chart"],
                "goal": "比较同济与对标高校在细分指标上的得分和相对排名。",
                "reference_style": "参考报告图7-图9的分项指标相对排名图。",
                "verification": "分项指标名称需与对应榜单原始字段及指标权重表对齐。",
            },
            {
                "id": "double_first_class_peer_comparison",
                "section": "二、国际排名/4、双一流高校学科排名对比分析",
                "output_types": ["text", "table", "chart"],
                "goal": "筛选双一流或国内主要对标高校，比较整体排名和分项表现。",
                "reference_style": "参考报告表4-表8及对标高校整体分析段落。",
                "verification": "对标高校范围需由双一流表、国内评估表或用户确认列表确定。",
            },
            {
                "id": "domestic_evaluation_comparison",
                "section": "三、国内评估/1、双一流高校第五轮学科评估结果对比",
                "output_types": ["text", "table"],
                "goal": "比较同济与双一流高校在第五轮学科评估、双一流、高峰高原等国内评价中的位置。",
                "reference_style": "采用简洁表格与结论段落。",
                "verification": "第五轮评估数据标记“不完全”时，结论需说明数据边界。",
            },
            {
                "id": "paper_strength_analysis",
                "section": "四、论文实力",
                "output_types": ["text", "table", "chart"],
                "goal": "从论文数量、被引、高质量论文、国际合作等角度分析论文实力。",
                "reference_style": "参考报告图10-图13的论文指标分析。",
                "verification": "若当前数据目录缺少独立论文库，优先使用 US News/软科/THE 中可代理的论文指标，并标记口径。",
            },
            {
                "id": "research_frontier_interdiscipline",
                "section": "五、国内外高校研究前沿及学科交叉",
                "output_types": ["text", "table"],
                "goal": "根据顶尖教师方向、排名指标或外部补充数据判断研究前沿和交叉方向。",
                "reference_style": "目标模板中该段落被标黄，说明需要重点补全。",
                "verification": "如果现有 Excel 不含教师方向数据，需标记为待补数据或转为基于榜单指标的弱分析。",
            },
            {
                "id": "opportunities_challenges",
                "section": "六、机会挑战",
                "output_types": ["text"],
                "goal": "综合前述排名、评估、论文与对标分析，给出机会、挑战和建议。",
                "reference_style": "参考报告“同济发展潜力及建议”。",
                "verification": "建议必须由前文数据结论支撑，不单独凭空生成。",
            },
        ],
    }
    return slots


def build_slot_data_mapping(catalog: dict[str, Any]) -> dict[str, Any]:
    def source(file_contains: str, role: str | None = None) -> list[dict[str, Any]]:
        matches = []
        for workbook in catalog["workbooks"]:
            if file_contains not in workbook["file"]:
                continue
            for sheet in workbook.get("sheets", []):
                if role is None or sheet.get("role") == role:
                    matches.append(
                        {
                            "file": workbook["file"],
                            "sheet": sheet["sheet"],
                            "role": sheet.get("role"),
                            "headers": sheet.get("headers", [])[:14],
                            "available_fields": sheet.get("available_fields", {}),
                        }
                    )
        return matches

    mappings = {
        "version": 1,
        "mapping_status": "draft_for_review",
        "slots": {
            "overview_competitive_summary": {
                "primary_sources": source("同济大学学院+学科映射关系")
                + source("同济大学第五轮学科评估结果")
                + source("教育部学科与四大国际学科排名映射关系"),
                "secondary_sources": source("2026年QS") + source("2026年THE") + source("2021-2026年软科") + source("2025-2026年度US News"),
                "join_keys": ["一级学科名称/学科名称", "榜单学科名称", "学校名称"],
                "confidence": 0.78,
                "risks": ["总述需要综合判断，部分优势挑战需要依赖后续 slot 的统计结果。"],
            },
            "ranking_indicator_system": {
                "primary_sources": source("四大世界学科2025指标权重") + source("世界大学学科排名2025指标权重"),
                "join_keys": ["榜单名称", "学科领域/学科名称"],
                "confidence": 0.9,
                "risks": ["QS/THE 权重表存在多级表头，需要专门清洗。"],
            },
            "tongji_ranking_history": {
                "primary_sources": source("世界大学学科排名THE（2022-2026年）")
                + source("2026年QS")
                + source("2025年QS")
                + source("2024年QS")
                + source("2023年QS")
                + source("2021-2026年软科")
                + source("2024年软科")
                + source("2023年软科")
                + source("2025-2026年度US News")
                + source("2024-2025年US News")
                + source("2023US News")
                + source("2022US News")
                + source("2021US News"),
                "supporting_sources": source("教育部学科与四大国际学科排名映射关系")
                + source("同济大学学院+学科映射关系"),
                "join_keys": ["同济大学", "学科中文/学科英文/榜单学科", "发布年度/排名年度"],
                "confidence": 0.86,
                "risks": ["不同榜单年份口径不同，US News 文件名和排名年度需统一。"],
            },
            "indicator_detail_comparison": {
                "primary_sources": source("2026年THE")
                + source("2026年QS")
                + source("2021-2026年软科")
                + source("2025-2026年度US News")
                + source("世界大学学科排名THE（2022-2026年）"),
                "supporting_sources": source("四大世界学科2025指标权重")
                + source("教育部学科与四大国际学科排名映射关系"),
                "join_keys": ["学校名称", "学科名称", "指标字段"],
                "confidence": 0.82,
                "risks": ["部分榜单只有得分没有指标相对排名，需要运行时按同学科同年度重算相对排名。"],
            },
            "double_first_class_peer_comparison": {
                "primary_sources": source("全国双一流学科明细表")
                + source("同济大学第五轮学科评估结果")
                + source("2026年THE")
                + source("2026年QS")
                + source("2021-2026年软科")
                + source("2025-2026年度US News"),
                "supporting_sources": source("教育部学科与四大国际学科排名映射关系"),
                "join_keys": ["一级学科", "学校名称", "中国内地/所在地区"],
                "confidence": 0.8,
                "risks": ["对标高校口径需要用户确认：双一流、第五轮同档、还是榜单前200内地高校。"],
            },
            "domestic_evaluation_comparison": {
                "primary_sources": source("第五轮评估结果")
                + source("副本第五轮评估结果")
                + source("全国高校第四轮&第三轮学科评估结果")
                + source("全国双一流学科明细表")
                + source("同济大学第五轮学科评估结果"),
                "join_keys": ["一级学科", "学校名称/单位名称"],
                "confidence": 0.88,
                "risks": ["第五轮数据文件注明不完全，生成报告时必须保留数据边界说明。"],
            },
            "paper_strength_analysis": {
                "primary_sources": source("2025-2026年度US News")
                + source("2021-2026年软科")
                + source("2026年THE")
                + source("世界大学学科排名THE（2022-2026年）"),
                "join_keys": ["学校名称", "学科名称", "论文/引用/高被引/国际合作相关指标"],
                "confidence": 0.68,
                "risks": ["当前目录未发现独立 WOS/Scopus 论文库，可能只能使用排名表中的论文代理指标。"],
            },
            "research_frontier_interdiscipline": {
                "primary_sources": [],
                "supporting_sources": source("同济大学学院+学科映射关系")
                + source("教育部学科与四大国际学科排名映射关系"),
                "join_keys": ["学科名称", "学院", "交叉方向"],
                "confidence": 0.35,
                "risks": ["现有表头盘点未发现顶尖教师方向或研究前沿专门数据，此 slot 需要补数据或改成弱分析。"],
            },
            "opportunities_challenges": {
                "primary_sources": ["依赖前序所有 slot 的结构化结论"],
                "join_keys": ["slot_id"],
                "confidence": 0.72,
                "risks": ["必须在前序统计完成后生成，不能独立生成。"],
            },
        },
    }
    return mappings


def build_entity_aliases() -> dict[str, Any]:
    return {
        "version": 1,
        "universities": {
            "同济大学": ["Tongji University", "同济"],
            "清华大学": ["Tsinghua University"],
            "北京大学": ["Peking University"],
            "上海交通大学": ["Shanghai Jiao Tong University"],
            "复旦大学": ["Fudan University"],
            "浙江大学": ["Zhejiang University"],
            "南京大学": ["Nanjing University"],
            "中国科学技术大学": ["University of Science and Technology of China", "USTC"],
        },
        "ranking_sources": {
            "软科": ["ARWU", "ShanghaiRanking", "软科世界一流学科排名"],
            "QS": ["Quacquarelli Symonds", "QS世界学科排名"],
            "THE": ["Times Higher Education", "泰晤士", "THE世界学科排名"],
            "US News": ["U.S.News", "U.S. News", "USNews", "美国新闻"],
        },
        "field_aliases": {
            "university": ["学校名称", "高校名称", "中文名称", "大学", "Name", "university", "单位名称"],
            "subject": ["学科", "学科中文", "学科名称", "学科领域", "Subject", "一级学科"],
            "rank": ["国际排名", "世界排名", "排名", "Rank", "subjectrank"],
            "score": ["总得分", "总分", "综合得分", "OVERALL_SCORE", "overallscore", "globalscore"],
            "year": ["发布年度", "排名年度", "年份", "发布时间", "pubyear", "Time"],
            "region": ["所在地区", "地区", "国家", "省份", "Country", "location"],
        },
    }


def summarize_catalog(catalog: dict[str, Any]) -> dict[str, Any]:
    sheet_counter = Counter()
    source_counter = Counter()
    role_counter = Counter()
    for workbook in catalog["workbooks"]:
        for sheet in workbook.get("sheets", []):
            role_counter[sheet.get("role", "unknown")] += 1
            for source_name in sheet.get("ranking_sources", []):
                source_counter[source_name] += 1
            sheet_counter["total"] += 1

    return {
        "excel_count": catalog["excel_count"],
        "sheet_count": sheet_counter["total"],
        "role_counts": dict(role_counter),
        "ranking_source_counts": dict(source_counter),
    }


def build_summary_markdown(
    catalog: dict[str, Any],
    template: dict[str, Any],
    reference: dict[str, Any],
    slots: dict[str, Any],
    mapping: dict[str, Any],
) -> str:
    catalog_summary = summarize_catalog(catalog)
    low_confidence_slots = [
        slot_id
        for slot_id, item in mapping["slots"].items()
        if isinstance(item, dict) and item.get("confidence", 1) < 0.7
    ]

    lines = [
        "# 阶段一盘点核验摘要",
        "",
        "## 本阶段目标",
        "",
        "把原始 Word/Excel 资产转换成后续分析 agent 可读取、可审查、可修改的配置基础。",
        "",
        "## 已生成文件",
        "",
        "- `data_catalog.json`：Excel 工作簿、sheet、字段、样例值与角色推断。",
        "- `target_template_outline.json`：标黄理科模板的段落结构。",
        "- `reference_report_outline.json`：数学参考报告的段落、表格、图片结构。",
        "- `reference_report_structure.md`：参考报告中章节、图表、表格标题的人工可读摘录。",
        "- `report_slots.yaml`：报告待生成位置与任务定义。",
        "- `slot_data_mapping.yaml`：每个任务建议使用的数据源、联接键、风险。",
        "- `entity_aliases.yaml`：学校、榜单、字段别名初稿。",
        "",
        "## 核验结果",
        "",
        f"- Excel 文件数：{catalog_summary['excel_count']}",
        f"- Excel sheet 数：{catalog_summary['sheet_count']}",
        f"- 目标模板非空段落数：{template['nonempty_paragraph_count']}",
        f"- 目标模板标黄段落数：{len(template['highlighted_paragraphs'])}",
        f"- 参考报告表格数：{reference['table_count']}",
        f"- 参考报告图片数：{reference['image_count']}",
        f"- 报告 slot 数：{len(slots['slots'])}",
        f"- 低置信度 slot：{', '.join(low_confidence_slots) if low_confidence_slots else '无'}",
        "",
        "## 数据角色统计",
        "",
    ]
    for role, count in catalog_summary["role_counts"].items():
        lines.append(f"- {role}: {count}")

    lines.extend(["", "## 需要人工确认的点", ""])
    lines.extend(
        [
            "1. 对标高校范围：双一流高校、第五轮同档高校、还是各榜单前200内地高校？",
            "2. 目标学科：当前参考报告是数学，后续生成其他学科时需要给出一级学科名称或学院名称。",
            "3. “研究前沿及学科交叉”目前没有明显专门数据源，需要补教师方向/研究方向数据，或改为基于排名指标的弱分析。",
            "4. 论文实力若要达到参考报告图10-图13的颗粒度，可能需要独立 WOS/Scopus 论文数据；当前可先使用排名明细中的论文代理指标。",
        ]
    )

    return "\n".join(lines) + "\n"


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    catalog = build_data_catalog()
    target_template = extract_doc_structure(TARGET_TEMPLATE)
    reference_report = extract_doc_structure(REFERENCE_REPORT)
    report_slots = build_report_slots()
    slot_mapping = build_slot_data_mapping(catalog)
    entity_aliases = build_entity_aliases()

    write_json(OUTPUT_DIR / "data_catalog.json", catalog)
    write_json(OUTPUT_DIR / "target_template_outline.json", target_template)
    write_json(OUTPUT_DIR / "reference_report_outline.json", reference_report)
    write_yaml(OUTPUT_DIR / "report_slots.yaml", report_slots)
    write_yaml(OUTPUT_DIR / "slot_data_mapping.yaml", slot_mapping)
    write_yaml(OUTPUT_DIR / "entity_aliases.yaml", entity_aliases)

    (OUTPUT_DIR / "reference_report_structure.md").write_text(
        build_reference_report_markdown(reference_report),
        encoding="utf-8",
        newline="\n",
    )

    summary = build_summary_markdown(
        catalog=catalog,
        template=target_template,
        reference=reference_report,
        slots=report_slots,
        mapping=slot_mapping,
    )
    (OUTPUT_DIR / "stage1_summary.md").write_text(summary, encoding="utf-8", newline="\n")

    print(summary)


if __name__ == "__main__":
    main()
