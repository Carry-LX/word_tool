"""Evidence-bound short conclusions and explicitly bound ESI report sections."""
from __future__ import annotations

import csv
import json
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import agent_runtime

PROJECT_DIR = Path(__file__).resolve().parent.parent


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def load_policy() -> dict:
    profile = agent_runtime.load_runtime_profile()
    path = agent_runtime.deep_get(profile, "word_pipeline.editorial_policy_path",
                                  "configs/runtime/report_editorial_policy.json")
    return read_json(agent_runtime.resolve_project_path(path))


def request_for_run(run_dir: Path) -> dict:
    path = run_dir / "agent_run_summary.json"
    return read_json(path).get("request", {}) if path.exists() else {}


def csv_rows(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def tongji_row(rows: list[dict]) -> dict:
    return next((row for row in rows if str(row.get("is_tongji", "")).lower() in {"1", "true"}
                 or "tongji" in row.get("university_name_en", "").lower()
                 or row.get("university_name_cn") == "同济大学"), {})


def esi_for_code(code: str, policy: dict | None = None) -> dict | None:
    policy = policy or load_policy()
    binding = policy["bindings"].get(code)
    if not binding:
        return None
    data = read_json(PROJECT_DIR / policy["esi_data_path"])
    return next(subject for subject in data["subjects"] if subject["subject_key"] == binding["esi_key"])


def recommendation_evidence(run_dir: Path) -> dict:
    policy = load_policy()
    request = request_for_run(run_dir)
    code = str(request.get("target_subject_code", "")).zfill(4)
    root = run_dir / "artifacts"
    quality_path = root / "figure_14_paper_cnci_cpp_relative_rank/chart_data.csv"
    quantity_path = root / "figure_13_paper_output_citation_relative_rank/chart_data.csv"
    collaboration_path = root / "figure_16_paper_international_collab_relative_rank/chart_data.csv"
    overview_path = root / "figure_01_tongji_subject_ranking_overview/chart_data.csv"
    return {"subject_code": code, "request": request, "esi": esi_for_code(code, policy),
            "quality": tongji_row(csv_rows(quality_path)), "quantity": tongji_row(csv_rows(quantity_path)),
            "collaboration": tongji_row(csv_rows(collaboration_path)), "overview": csv_rows(overview_path),
            "evidence_paths": [str(path) for path in [quality_path, quantity_path, collaboration_path, overview_path]
                               if path.exists()]}


def concise_recommendations(run_dir: Path) -> list[str]:
    policy, evidence = load_policy(), recommendation_evidence(run_dir)
    code, esi, quality = evidence["subject_code"], evidence["esi"], evidence["quality"]
    paragraphs = []
    if esi:
        first, last = esi["values"][0], esi["values"][-1]
        paragraphs.append(f"1. 发展基础：ESI{esi['subject_name']}千分位由{first:.3f}‰优化至{last:.3f}‰，位次持续前移。")
    else:
        overview = evidence["overview"]
        row = next((row for row in overview if row.get("ranking_source") == "软科"), overview[0] if overview else {})
        if row:
            year = row.get("ranking_year") or row.get("release_year", "")
            source = row.get("ranking_source", "")
            subject = row.get("ranking_subject_display") or row.get("ranking_subject_raw", "")
            rank = re.sub(r"^[#=]+", "", row.get("rank_raw", ""))
            paragraphs.append(f"1. 发展基础：{source}{year}年{subject}位列第{rank}，据此确定下一阶段对标目标。")
    cnci_rank, cpp_rank = quality.get("CNCI_relative_rank"), quality.get("citation_impact_relative_rank")
    if cnci_rank and cpp_rank:
        paragraphs.append(f"{len(paragraphs) + 1}. 提升重点：37校同学科样本中，CNCI和篇均被引分别列第{int(float(cnci_rank))}、{int(float(cpp_rank))}位。")
    action = policy["bindings"].get(code, {}).get("action")
    if not action:
        action = "围绕优势研究方向培育代表性成果，持续跟踪论文质量指标。" if quality else "围绕已有排名指标开展对标，优先补齐当前缺失的同学科数据。"
    paragraphs.append(f"{len(paragraphs) + 1}. 建设建议：{action}")
    assert sum(len(text) for text in paragraphs) <= policy["max_recommendation_characters"]
    return paragraphs


def is_conclusion_heading(paragraph) -> bool:
    text = re.sub(r"\s+", "", paragraph.text)
    return paragraph.style.name == "Heading 1" and "发展潜力" in text and "建议" in text


def find_conclusion(document: Document):
    return next((paragraph for paragraph in document.paragraphs if is_conclusion_heading(paragraph)), None)


def set_text(paragraph, text: str, template=None, size: float = 12, highlight: bool = False) -> None:
    properties = deepcopy(template._p.pPr) if template is not None and template._p.pPr is not None else None
    paragraph.clear()
    if properties is not None:
        old = paragraph._p.pPr
        if old is not None:
            paragraph._p.remove(old)
        paragraph._p.insert(0, properties)
    properties = paragraph._p.get_or_add_pPr()
    snap = properties.find(qn("w:snapToGrid"))
    if snap is None:
        snap = OxmlElement("w:snapToGrid")
        properties.append(snap)
    snap.set(qn("w:val"), "0")
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    if highlight:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def insert_paragraph_before(document: Document, anchor, text: str = "", style=None):
    paragraph = document.add_paragraph(style=style)
    anchor._p.addprevious(paragraph._p)
    if text:
        set_text(paragraph, text)
    return paragraph


def replace_conclusions(document: Document, run_dir: Path) -> dict:
    heading = find_conclusion(document)
    if heading is None:
        raise ValueError("Conclusion chapter was not found")
    body = document._element.body
    elements = list(body)
    start = elements.index(heading._p) + 1
    old_text = []
    for element in elements[start:]:
        if element.tag == qn("w:sectPr"):
            break
        if element.tag == qn("w:p"):
            style = element.find("w:pPr/w:pStyle", namespaces=element.nsmap)
            if style is not None and style.get(qn("w:val")) == "Heading1":
                break
            old_text.append("".join(node.text or "" for node in element.xpath(".//w:t")))
        body.remove(element)
    texts = concise_recommendations(run_dir)
    anchor = heading._p
    for text in texts:
        paragraph = document.add_paragraph()
        anchor.addnext(paragraph._p)
        set_text(paragraph, text)
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.keep_together = True
        anchor = paragraph._p
    return {"old_characters": sum(len(text) for text in old_text),
            "new_characters": sum(len(text) for text in texts), "point_count": len(texts), "paragraphs": texts}


def insert_esi_before_conclusions(document: Document, run_dir: Path) -> dict:
    policy = load_policy()
    request = request_for_run(run_dir)
    code = str(request.get("target_subject_code", "")).zfill(4)
    esi = esi_for_code(code, policy)
    if esi is None:
        return {"inserted": False, "reason": "no_explicit_esi_binding"}
    title = "七、ESI前千分之一学科发展趋势"
    if any(paragraph.text == title for paragraph in document.paragraphs):
        return {"inserted": False, "reason": "already_present"}
    heading = find_conclusion(document)
    if heading is None:
        raise ValueError("Conclusion chapter was not found")
    heading.text = "八、同济发展潜力及建议"
    heading.paragraph_format.page_break_before = False
    esi_heading = insert_paragraph_before(document, heading, style="Heading 1")
    set_text(esi_heading, title, size=14)
    esi_heading.paragraph_format.page_break_before = True
    esi_heading.paragraph_format.keep_with_next = True
    scope = insert_paragraph_before(document, heading)
    set_text(scope, f"本节采用ESI“{esi['subject_name']}”口径，观察期为2024年11月至2026年9月；千分位越小，排名越靠前。", size=10.5)
    scope.paragraph_format.space_after = Pt(5)
    scope.paragraph_format.keep_with_next = True
    narrative = insert_paragraph_before(document, heading)
    set_text(narrative, policy["esi_narratives"][esi["subject_key"]], highlight=True)
    narrative.paragraph_format.first_line_indent = Cm(0.74)
    narrative.paragraph_format.line_spacing = 1.5
    narrative.paragraph_format.space_after = Pt(5)
    narrative.paragraph_format.keep_with_next = True
    narrative.paragraph_format.keep_together = True
    artifact = PROJECT_DIR / policy["esi_artifacts_dir"] / esi["subject_key"]
    if not (artifact / "chart.png").is_file():
        raise FileNotFoundError(f"ESI artifact has not been built: {artifact}")
    image = insert_paragraph_before(document, heading)
    image.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image.paragraph_format.line_spacing = 1
    image.paragraph_format.space_after = Pt(0)
    image.paragraph_format.keep_with_next = True
    image.add_run().add_picture(str(artifact / "chart.png"), width=Cm(14.8))
    max_caption = max((int(match.group(1)) for paragraph in document.paragraphs
                       if (match := re.match(r"^图\s*(\d+)\s*[、：:]", paragraph.text.strip()))), default=0)
    caption = insert_paragraph_before(document, heading)
    set_text(caption, f"图{max_caption + 1}、同济大学ESI{esi['subject_name']}排名千分位变化趋势", size=10.5)
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = False
    return {"inserted": True, "esi_key": esi["subject_key"], "esi_subject": esi["subject_name"],
            "figure_number": max_caption + 1, "artifact_dir": str(artifact),
            "narrative": policy["esi_narratives"][esi["subject_key"]]}
